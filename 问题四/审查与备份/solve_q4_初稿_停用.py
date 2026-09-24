"""Problem 4 exact partition solver.

The input is a frozen, independently verified Q3 polishing snapshot.  Q4
does not change routes, times, vehicles, batteries, relay positions or
communication providers.  It only partitions service areas and recomputes
the resources required when each group is operated independently.
"""
from __future__ import annotations

import csv
import hashlib
import itertools
import json
import math
import shutil
from collections import defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
SNAP = HERE / "输入快照"
OUT = HERE / "结果"
DATA = ROOT / "D题" / "数据" / "无人机应急物资运输基础数据"
TEMPLATE = ROOT / "D题" / "结果提交模板.xlsx"
SOURCE = SNAP / "精修方案_源文件.json"
TOL = 2e-4

RESOURCE_ORDER = ("U_A", "U_B", "U_C", "P_A", "P_B", "P_C", "U_R", "P_R")
INVENTORY = {"U_A": 4, "U_B": 2, "U_C": 2,
             "P_A": 6, "P_B": 4, "P_C": 4,
             "U_R": 2, "P_R": 6}
FULL_CHARGE = {"A": 1800.0, "B": 2400.0, "C": 3000.0, "R": 1800.0}


def dump(path: Path, obj) -> None:
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def read_inventory() -> Dict[str, int]:
    """Read inventories from the original workbooks, rather than hardcoding them."""
    result = {}
    wb = load_workbook(DATA / "运输无人机数据.xlsx", read_only=True, data_only=True)
    ws = wb["数据"]
    rows = list(ws.values)
    wb.close()
    # Rows 19--21 in the attachment are A/B/C battery inventories.
    for row in rows:
        if row and row[0] in {"A", "B", "C"} and row[1] is not None and row[2] is not None:
            # Parameter rows also have A/B/C, but column 1 is the model name.
            if isinstance(row[1], (int, float)):
                result[f"P_{row[0]}"] = int(row[1])
    wb = load_workbook(DATA / "运输无人机数据.xlsx", read_only=True, data_only=True)
    ws = wb["数据"]
    rows = list(ws.values)
    wb.close()
    for row in rows:
        if row and row[0] in {"A", "B", "C"} and isinstance(row[3], (int, float)):
            # The model parameter rows contain maximum payload in column 3.
            pass
    # Entity lists provide the transportation UAV inventory.
    for row in rows:
        if row and isinstance(row[0], str) and row[0].startswith("U") and row[1] in {"A", "B", "C"}:
            result[f"U_{row[1]}"] = result.get(f"U_{row[1]}", 0) + 1

    wb = load_workbook(DATA / "中继无人机数据.xlsx", read_only=True, data_only=True)
    ws = wb["数据"]
    rows = list(ws.values)
    wb.close()
    for row in rows:
        # The shared-energy inventory row is ('R', 6, 1800, ...).
        if row and row[0] == "R" and isinstance(row[1], (int, float)):
            result["P_R"] = int(row[1])
        if row and isinstance(row[0], str) and row[0].startswith("R0"):
            result["U_R"] = result.get("U_R", 0) + 1
    if result != INVENTORY:
        raise ValueError(f"workbook inventory mismatch: {result} != {INVENTORY}")
    return result


class UnionFind:
    def __init__(self, items: Iterable[str]):
        self.parent = {x: x for x in items}

    def find(self, x: str) -> str:
        p = self.parent[x]
        if p != x:
            self.parent[x] = self.find(p)
        return self.parent[x]

    def union(self, a: str, b: str) -> None:
        ra, rb = self.find(a), self.find(b)
        if ra != rb:
            self.parent[rb] = ra


def peak(intervals: Sequence[Tuple[float, float]]) -> int:
    """Peak occupancy of half-open intervals; releases occur before starts."""
    events = []
    for a, b in intervals:
        if b <= a + 1e-10:
            continue
        events.append((float(a), 1))
        events.append((float(b), -1))
    events.sort(key=lambda x: (x[0], x[1]))
    cur = best = 0
    for _, delta in events:
        cur += delta
        best = max(best, cur)
    return best


def canonical_partition(groups: Sequence[Sequence[int]]) -> Tuple[Tuple[int, ...], ...]:
    return tuple(sorted((tuple(sorted(g)) for g in groups), key=lambda g: (g[0], len(g), g)))


def set_partitions(n: int, k: int) -> List[Tuple[Tuple[int, ...], ...]]:
    """All unlabeled partitions of range(n) into k nonempty groups."""
    out = set()

    def rec(i: int, groups: List[List[int]]) -> None:
        if i == n:
            if len(groups) == k:
                out.add(canonical_partition(groups))
            return
        if len(groups) < k:
            rec(i + 1, groups + [[i]])
        for j in range(len(groups)):
            groups[j].append(i)
            rec(i + 1, groups)
            groups[j].pop()

    rec(0, [])
    return sorted(out)


def lcm_many(values: Iterable[int]) -> int:
    out = 1
    for v in values:
        out = math.lcm(out, int(v))
    return out


def normalize_row(row: Dict) -> Dict:
    return {k: row[k] for k in row}


def load_snapshot() -> Tuple[Dict, Dict, Dict, Dict]:
    plan = json.loads(SOURCE.read_text(encoding="utf-8"))
    routes = {str(r["route_id"]): normalize_row(r) for r in plan["routes"]}
    relays = {str(r["relay_id"]): normalize_row(r) for r in plan["relays"]}
    slots = {}
    for rid, relay in relays.items():
        slot = int(relay["slot"])
        if slot in slots:
            raise ValueError(f"duplicate relay slot {slot}")
        slots[slot] = rid
    zones = sorted({str(z) for r in routes.values() for z in r["zones"]})
    expected = {f"S{i:03d}" for i in range(1, 16)}
    if set(zones) != expected:
        raise ValueError(f"service-zone set mismatch: {zones}")

    # link timestamps in the frozen JSON are route-relative; the verifier adds
    # route['start'] before checking the relay service window.
    relay_routes: Dict[str, set] = defaultdict(set)
    link_rows = []
    for rid, route in routes.items():
        for number, link in enumerate(route.get("links", []), 1):
            if "relay_slot" not in link:
                continue
            slot = int(link["relay_slot"])
            if slot not in slots:
                raise ValueError(f"route {rid} refers to unknown relay slot {slot}")
            relay_id = slots[slot]
            t0 = float(link["t0"]) + float(route["start"])
            t1 = float(link["t1"]) + float(route["start"])
            relay = relays[relay_id]
            if t0 < float(relay["alpha"]) - TOL or t1 > float(relay["beta"]) + TOL:
                raise ValueError(f"relay interval does not contain link: {rid}/{relay_id}/{t0}-{t1}")
            relay_routes[relay_id].add(rid)
            link_rows.append({"route_id": rid, "link_number": number, "relay_slot": slot,
                              "relay_id": relay_id, "t0": t0, "t1": t1,
                              "mask": int(link.get("mask", 0))})
    if set(relay_routes) != set(relays):
        missing = sorted(set(relays) - set(relay_routes))
        raise ValueError(f"executed relay task has no traceable provider intervals: {missing}")

    # Build the task-association graph over service zones.
    uf = UnionFind(zones)
    for route in routes.values():
        rz = sorted(set(route["zones"]))
        for z in rz[1:]:
            uf.union(rz[0], z)
    relay_zones = {}
    for relay_id, route_ids in relay_routes.items():
        rz = sorted({z for rid in route_ids for z in routes[rid]["zones"]})
        relay_zones[relay_id] = rz
        for z in rz[1:]:
            uf.union(rz[0], z)
    components_dict = defaultdict(list)
    for z in zones:
        components_dict[uf.find(z)].append(z)
    components = sorted((tuple(sorted(v)) for v in components_dict.values()), key=lambda x: x[0])
    zone_component = {z: i for i, comp in enumerate(components) for z in comp}
    route_component = {}
    for rid, route in routes.items():
        ci = {zone_component[z] for z in route["zones"]}
        if len(ci) != 1:
            raise ValueError(f"transport route spans multiple components: {rid}")
        route_component[rid] = next(iter(ci))
    relay_component = {}
    for relay_id, rz in relay_zones.items():
        ci = {zone_component[z] for z in rz}
        if len(ci) != 1:
            raise ValueError(f"relay task spans multiple components: {relay_id}")
        relay_component[relay_id] = next(iter(ci))
    mappings = {"slot_to_relay": [{"slot": s, "relay_id": slots[s]} for s in sorted(slots)],
                "route_relay_links": link_rows,
                "relay_routes": {k: sorted(v) for k, v in sorted(relay_routes.items())},
                "relay_zones": relay_zones}
    structure = {"zones": zones, "components": [list(x) for x in components],
                 "zone_component": zone_component, "route_component": route_component,
                 "relay_component": relay_component, "M_T": len(zones), "M": len(components)}
    return plan, routes, relays, {"slots": slots, "relay_routes": relay_routes,
                                  "relay_zones": relay_zones, "components": components,
                                  "zone_component": zone_component, "route_component": route_component,
                                  "relay_component": relay_component, "mappings": mappings,
                                  "structure": structure}


def interval_data(routes: Dict, relays: Dict, assoc: Dict) -> Dict[str, List[Tuple[float, float]]]:
    data = {r: [] for r in RESOURCE_ORDER}
    for route in routes.values():
        g = str(route["vehicle"])
        data[f"U_{g}"].append((float(route["start"]), float(route["finish"])))
        data[f"P_{g}"].append((float(route["start"]), float(route["recharge"])))
    for relay in relays.values():
        data["U_R"].append((float(relay["start"]), float(relay["body_ready"])))
        data["P_R"].append((float(relay["start"]), float(relay["recharge"])))
    return data


def tasks_for_group(partition: Tuple[Tuple[int, ...], ...], group: Tuple[int, ...], routes: Dict,
                    relays: Dict, assoc: Dict) -> Tuple[List[Dict], List[Dict]]:
    selected = set(group)
    route_tasks = [r for rid, r in routes.items() if assoc["route_component"][rid] in selected]
    relay_tasks = [r for rid, r in relays.items() if assoc["relay_component"][rid] in selected]
    return sorted(route_tasks, key=lambda r: (float(r["start"]), r["route_id"])), sorted(relay_tasks,
                                                                                         key=lambda r: (float(r["start"]), r["relay_id"]))


def work_imbalance(values: Sequence[float]) -> float:
    total = sum(values)
    if total <= 1e-12:
        return 0.0
    g = len(values)
    return g / (2 * (g - 1)) * sum(abs(v / total - 1 / g) for v in values)


def evaluate_partition(partition: Tuple[Tuple[int, ...], ...], routes: Dict, relays: Dict, assoc: Dict,
                       inventory: Dict[str, int]) -> Dict:
    groups = []
    total_routes = []
    total_relays = []
    transport_work = []
    relay_work = []
    D = {r: 0 for r in RESOURCE_ORDER}
    group_resources = []
    for group in partition:
        rt, rr = tasks_for_group(partition, group, routes, relays, assoc)
        total_routes.extend(rt)
        total_relays.extend(rr)
        transport_work.append(sum(float(x["finish"]) - float(x["start"]) for x in rt))
        relay_work.append(sum(float(x["finish"]) - float(x["start"]) for x in rr))
        intervals = {r: [] for r in RESOURCE_ORDER}
        for route in rt:
            g = str(route["vehicle"])
            intervals[f"U_{g}"].append((float(route["start"]), float(route["finish"])))
            intervals[f"P_{g}"].append((float(route["start"]), float(route["recharge"])))
        for relay in rr:
            intervals["U_R"].append((float(relay["start"]), float(relay["body_ready"])))
            intervals["P_R"].append((float(relay["start"]), float(relay["recharge"])))
        gr = {r: peak(intervals[r]) for r in RESOURCE_ORDER}
        group_resources.append(gr)
        for r in RESOURCE_ORDER:
            D[r] += gr[r]
        groups.append({"units": list(group), "services": sorted({z for x in group for z in assoc["components"][x]}),
                       "route_ids": [x["route_id"] for x in rt], "relay_ids": [x["relay_id"] for x in rr],
                       "resources": gr, "transport_work_s": transport_work[-1], "relay_work_s": relay_work[-1]})
    A = sum(D[r] / inventory[r] for r in RESOURCE_ORDER) / len(RESOURCE_ORDER)
    BT = work_imbalance(transport_work)
    BR = work_imbalance(relay_work)
    active = [x for x, w in [("T", sum(transport_work)), ("R", sum(relay_work))] if w > 1e-12]
    B = (BT + BR) / 2 if len(active) == 2 else (BT if active == ["T"] else BR)
    gaps = {r: max(D[r] - inventory[r], 0) for r in RESOURCE_ORDER}
    stock = {r: max(inventory[r] - D[r], 0) for r in RESOURCE_ORDER}
    concentrated = {r: peak(interval_data(routes, relays, assoc)[r]) for r in RESOURCE_ORDER}
    split_loss = {r: D[r] - concentrated[r] for r in RESOURCE_ORDER}
    lcm = lcm_many(inventory.values())
    Q = sum((lcm // inventory[r]) * D[r] for r in RESOURCE_ORDER)
    H_stock = sum(gaps[r] / inventory[r] for r in RESOURCE_ORDER)
    code = "|".join("+".join(str(x) for x in group) for group in partition)
    return {"partition": [list(x) for x in partition], "groups": groups,
            "D": D, "A": A, "BT": BT, "BR": BR, "B": B, "Q": Q,
            "transport_work_s": transport_work, "relay_work_s": relay_work,
            "transport_work_total_s": sum(transport_work), "relay_work_total_s": sum(relay_work),
            "gaps": gaps, "stock": stock, "concentrated_P": concentrated,
            "split_loss": split_loss, "H_stock": H_stock, "partition_code": code}


def dominates(a: Dict, b: Dict) -> bool:
    return ((a["A"] <= b["A"] + 1e-12 and a["B"] <= b["B"] + 1e-12)
            and (a["A"] < b["A"] - 1e-12 or a["B"] < b["B"] - 1e-12))


def nondominated(records: List[Dict]) -> List[Dict]:
    return [x for x in records if not any(dominates(y, x) for y in records if y is not x)]


def epsilon_records(records: List[Dict]) -> List[Dict]:
    eps = max(r["Q"] for r in records)
    out = []
    while True:
        feasible = [r for r in records if r["Q"] <= eps]
        if not feasible:
            break
        best = min(feasible, key=lambda r: (r["B"], r["Q"], r["partition_code"]))
        out.append({"epsilon": eps, "selected_Q": best["Q"], "selected_B": best["B"],
                    "partition_code": best["partition_code"]})
        nxt = best["Q"] - 1
        if nxt < min(r["Q"] for r in records):
            break
        eps = nxt
    return out


def minimax(records: List[Dict], delta: float = 0.3) -> Tuple[Dict, Dict]:
    lo, hi = max(0.0, 0.5 - delta / 2), min(1.0, 0.5 + delta / 2)
    values = {}
    for theta in (lo, hi):
        values[theta] = min(theta * r["A"] + (1 - theta) * r["B"] for r in records)
    for r in records:
        r["regret_lo"] = lo * r["A"] + (1 - lo) * r["B"] - values[lo]
        r["regret_hi"] = hi * r["A"] + (1 - hi) * r["B"] - values[hi]
        r["regret_max"] = max(r["regret_lo"], r["regret_hi"], 0.0)
    candidates = nondominated(records)
    selected = min(candidates, key=lambda r: (r["regret_max"], r["H_stock"], r["B"], r["Q"], r["partition_code"]))
    return selected, {"delta": delta, "theta_endpoints": [lo, hi], "v": {str(k): v for k, v in values.items()}}


def write_csv(path: Path, rows: List[Dict], fields: Sequence[str] | None = None) -> None:
    if not rows:
        path.write_text("", encoding="utf-8-sig")
        return
    fields = list(fields or rows[0].keys())
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def export_outputs(plan: Dict, routes: Dict, relays: Dict, assoc: Dict,
                   all_records: Dict[int, List[Dict]], selected: Dict[int, Dict],
                   minimax_meta: Dict[int, Dict], inventory: Dict[str, int]) -> None:
    OUT.mkdir(exist_ok=True)
    # Flatten partition-level summary.
    flat = []
    for G, records in all_records.items():
        for r in records:
            flat.append({"G": G, "partition_code": r["partition_code"], "A": r["A"], "BT": r["BT"],
                         "BR": r["BR"], "B": r["B"], "Q": r["Q"], "regret_max": r.get("regret_max", ""),
                         "H_stock": r["H_stock"], "D": json.dumps(r["D"], ensure_ascii=False),
                         "gaps": json.dumps(r["gaps"], ensure_ascii=False)})
    write_csv(OUT / "分区枚举明细.csv", flat)
    nd = []
    for G, records in all_records.items():
        for r in nondominated(records):
            nd.append({"G": G, "partition_code": r["partition_code"], "A": r["A"], "BT": r["BT"],
                       "BR": r["BR"], "B": r["B"], "Q": r["Q"], "regret_max": r.get("regret_max", ""),
                       "H_stock": r["H_stock"]})
    write_csv(OUT / "非支配档案.csv", nd)
    eps_rows = []
    for G, records in all_records.items():
        for row in epsilon_records(records):
            eps_rows.append({"G": G, **row})
    write_csv(OUT / "epsilon搜索记录.csv", eps_rows)

    result_rows = []
    group_detail_rows = []
    gap_rows = []
    final_json = {"source": str(SOURCE), "source_sha256": sha256(SOURCE), "inventory": inventory,
                  "objective": plan.get("objective"), "structure": assoc["structure"],
                  "selection": {}, "solutions": {}}
    for G in (2, 3):
        r = selected[G]
        meta = minimax_meta[G]
        final_json["selection"][str(G)] = {"partition_code": r["partition_code"], "A": r["A"], "BT": r["BT"],
                                             "BR": r["BR"], "B": r["B"], "Q": r["Q"],
                                             "regret_max": r["regret_max"], "regret_lo": r["regret_lo"],
                                             "regret_hi": r["regret_hi"], "H_stock": r["H_stock"],
                                             "minimax_meta": meta}
        final_json["solutions"][str(G)] = r
        for h, group in enumerate(r["groups"], 1):
            row = {"K（2或3）": G, "任务组编号": f"G{G}-{h:02d}",
                   "服务区列表": "、".join(group["services"]),
                   "A型运输无人机数": group["resources"]["U_A"],
                   "B型运输无人机数": group["resources"]["U_B"],
                   "C型运输无人机数": group["resources"]["U_C"],
                   "A型电池组数": group["resources"]["P_A"],
                   "B型电池组数": group["resources"]["P_B"],
                   "C型电池组数": group["resources"]["P_C"],
                   "中继无人机数": group["resources"]["U_R"],
                   "中继能源组件数": group["resources"]["P_R"]}
            result_rows.append(row)
            for rid in group["route_ids"]:
                group_detail_rows.append({"G": G, "任务组编号": f"G{G}-{h:02d}", "类型": "运输", "任务编号": rid})
            for rid in group["relay_ids"]:
                group_detail_rows.append({"G": G, "任务组编号": f"G{G}-{h:02d}", "类型": "中继", "任务编号": rid})
        for resource in RESOURCE_ORDER:
            gap_rows.append({"G": G, "资源类型": resource, "集中最低需求P_r": r["concentrated_P"][resource],
                             "独立需求D_r": r["D"][resource], "原库存I_r": inventory[resource],
                             "分区附加需求R_split": r["split_loss"][resource],
                             "库存缺口Delta": r["gaps"][resource], "库存剩余": r["stock"][resource]})
    write_csv(OUT / "Q4_分区配置.csv", result_rows)
    write_csv(OUT / "分组任务明细.csv", group_detail_rows)
    write_csv(OUT / "库存缺口分析.csv", gap_rows)
    dump(OUT / "最终代表方案.json", final_json)

    # Preserve the official template and fill only the Q4 sheet.
    out_xlsx = HERE / "问题四_结果提交.xlsx"
    shutil.copy2(TEMPLATE, out_xlsx)
    wb = load_workbook(out_xlsx)
    ws = wb["Q4_分区配置"]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.value = None
    for i, row in enumerate(result_rows, 2):
        values = [row[h] for h in ["K（2或3）", "任务组编号", "服务区列表", "A型运输无人机数", "B型运输无人机数",
                                   "C型运输无人机数", "A型电池组数", "B型电池组数", "C型电池组数", "中继无人机数",
                                   "中继能源组件数"]]
        for j, value in enumerate(values, 1):
            ws.cell(i, j).value = value
    wb.save(out_xlsx)

    # Frozen input manifest and structural audit.
    manifest = {"source": str(SOURCE), "source_sha256": sha256(SOURCE),
                "verification_files": {p.name: sha256(p) for p in SNAP.glob("*.json")},
                "q3_objective": plan.get("objective"), "inventory": inventory,
                "M_T": assoc["structure"]["M_T"], "M": assoc["structure"]["M"],
                "partition_counts": {"2": len(all_records[2]), "3": len(all_records[3])},
                "scope": "exact partition enumeration conditional on the frozen Q3 polishing snapshot; no Q3 route or communication reoptimization"}
    dump(HERE / "输入快照冻结清单.json", manifest)
    audit = {"status": "PASS", "routes": len(routes), "relays": len(relays), "zones": len(assoc["structure"]["zones"]),
             "M_T": assoc["structure"]["M_T"], "M": assoc["structure"]["M"],
             "partition_counts": {"G2": len(all_records[2]), "G3": len(all_records[3])},
             "all_zone_unique": True, "all_groups_nonempty": True, "route_same_group_checked": True,
             "relay_same_group_checked": True, "slot_mapping_checked": True,
             "resource_interval_peak_checked": True, "source_version_frozen": True,
             "note": "D_r<=I_r was not imposed as a main constraint; gaps are reported separately."}
    dump(OUT / "核验报告.json", audit)
    dump(HERE / "输入快照" / "结构与字段映射.json", {"structure": assoc["structure"], "mappings": assoc["mappings"]})
    write_csv(HERE / "输入快照" / "slot_中继架次映射.csv", assoc["mappings"]["slot_to_relay"])
    write_csv(HERE / "输入快照" / "通信关联.csv", assoc["mappings"]["route_relay_links"])


def main() -> None:
    inventory = read_inventory()
    plan, routes, relays, assoc = load_snapshot()
    if plan.get("verification_status") != "PASS" or plan.get("status") != "independently_verified_fixed_structure_polish":
        raise ValueError("frozen snapshot is not the expected independently verified polishing plan")
    all_records = {}
    selected = {}
    metas = {}
    for G in (2, 3):
        if assoc["structure"]["M"] < G:
            raise ValueError(f"G={G} is structurally infeasible because M={assoc['structure']['M']}")
        parts = set_partitions(assoc["structure"]["M"], G)
        records = [evaluate_partition(p, routes, relays, {**assoc, "components": assoc["components"]}, inventory)
                   for p in parts]
        chosen, meta = minimax(records)
        all_records[G] = records
        selected[G] = chosen
        metas[G] = meta
    export_outputs(plan, routes, relays, assoc, all_records, selected, metas, inventory)
    print(json.dumps({"status": "PASS", "M_T": assoc["structure"]["M_T"], "M": assoc["structure"]["M"],
                      "partitions": {str(G): len(all_records[G]) for G in all_records},
                      "selected": {str(G): {k: selected[G][k] for k in ["partition_code", "A", "BT", "BR", "B", "Q", "regret_max"]}
                                   for G in selected}}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
