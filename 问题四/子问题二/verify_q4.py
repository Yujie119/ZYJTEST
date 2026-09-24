"""Independent Q4 verifier. Does not import the solver or its evaluator.

Rebuilds components with a graph traversal, enumerates labeled assignments,
and computes interval peaks with event sweeps instead of the solver's scan.
Existing Q3 physical certificates are inherited, not rebranded as new physics.
"""
import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from q4_paths import output
from fractions import Fraction as F
from collections import defaultdict,Counter
from itertools import product,combinations
import csv,json,hashlib,traceback
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
R=('U_A','U_B','U_C','P_A','P_B','P_C','U_R','P_R')
F0=F(0);checks=[]
def f(v):return F(str(v))
def read(p):return json.loads(Path(p).read_text(encoding='utf-8'))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def csvread(p):
    with Path(p).open(encoding='utf-8-sig',newline='') as ff:return list(csv.DictReader(ff))
def check(ok,label):
    checks.append({'check':label,'pass':bool(ok)})
    if not ok:raise AssertionError(label)
def near(a,b,tol=1e-10):return abs(float(a)-float(b))<=tol

def parameters():
    inv={r:0 for r in R};cap={};charge={};turn={}
    for fn in ['运输无人机数据.xlsx','中继无人机数据.xlsx']:
        w=load_workbook(ROOT/'输入快照/原始附件'/fn,read_only=True,data_only=True)
        section=''
        for row in w['数据'].values:
            if row[0] in ['三类机型参数','中继机型参数','逐架无人机清单','逐架中继无人机清单','共享电池库存','共享能源组件库存']:
                section=row[0];continue
            if section in ['三类机型参数','中继机型参数'] and row[0] in ['A','B','C','R']:
                relay=row[0]=='R';cap[row[0]]=f(row[7 if relay else 8]);turn[row[0]]=f(row[11]) if relay else F0
            if section.startswith('逐架') and row[1] in ['A','B','C','R']:
                check(row[2]=='O01','initial_location_'+str(row[0]));inv['U_'+row[1]]+=1
            if section.startswith('共享') and row[0] in ['A','B','C','R']:
                inv['P_'+row[0]]=int(row[1]);charge[row[0]]=f(row[2])
        w.close()
    return inv,cap,charge,turn

def sweep(events):
    deltas=[]
    for a,b,k in events:
        deltas.extend([(a,1),(b,-1)])
    active=0;maximum=0
    for t,v in sorted(deltas):active+=v;maximum=max(maximum,active)
    assert active==0
    return maximum

def rebuild(directory,params):
    inv,cap,charge,turn=params;p=read(directory/'方案.json')
    rt={r['route_id']:r for r in p['routes']};rr={r['relay_id']:r for r in p['relays']}
    zones=sorted({z for r in rt.values() for z in r['zones']});graph={z:set() for z in zones}
    def connect(zs):
        for a in zs:graph[a].update(set(zs)-{a})
    def comps():
        unseen=set(zones);out=[]
        while unseen:
            stack=[min(unseen)];part=set()
            while stack:
                z=stack.pop()
                if z in part:continue
                part.add(z);stack.extend(graph[z]-part)
            unseen-=part;out.append(sorted(part))
        return sorted(out)
    for r in rt.values():connect(r['zones'])
    mt=len(comps());relations=defaultdict(set)
    slots={str(r['slot']):j for j,r in rr.items()}
    check(len(slots)==len(rr),directory.name+'_unique_slots')
    slotrows=csvread(directory/'slot_中继任务映射.csv')
    check({r['slot']:r['relay_id'] for r in slotrows}==slots,directory.name+'_slot_table')
    intervals=read(directory/'实际通信分段.json');points=csvread(directory/'通信事件端点核验.csv')
    for row in intervals:
        k=row['route_id'];check(k in rt,directory.name+'_interval_route')
        if row['mode']=='中继':
            j=row['relay_id'];check(j in rr,directory.name+'_interval_relay')
            check(rr[j]['alpha']-2e-4<=row['t0']<=row['t1']<=rr[j]['beta']+2e-4,directory.name+'_relay_window')
            relations[j].add(k)
    for row in points:
        k=row['route_id'];j=row['provider'];check(k in rt,directory.name+'_endpoint_route')
        if j!='G01':
            check(j in rr,directory.name+'_endpoint_relay')
            check(rr[j]['alpha']-2e-4<=float(row['time_s'])<=rr[j]['beta']+2e-4,directory.name+'_endpoint_window')
            relations[j].add(k)
    check(set(relations)==set(rr),directory.name+'_all_relays_have_actual_tasks')
    for j,ks in relations.items():connect({z for k in ks for z in rt[k]['zones']})
    cc=comps();zi={z:i for i,zs in enumerate(cc) for z in zs};taskunit={}
    events={};work={'T':[F0]*len(cc),'R':[F0]*len(cc)}
    for mode,items in [('T',rt),('R',rr)]:
        for k,r in items.items():
            u=zi[r['zones'][0]] if mode=='T' else zi[rt[next(iter(relations[k]))]['zones'][0]]
            taskunit[k]=u;typ=r['vehicle'] if mode=='T' else 'R'
            a=f(r['start']);b=f(r['finish']);soc=1-f(r['energy'])/cap[typ]
            check(F(1,5)-F(1,1000000)<=soc<=1,directory.name+'_safe_soc_'+k)
            ch=charge[typ]*(F(65,100)*max(F(9,10)-soc,F0)/F(9,10)+F(35,100)*min(1-soc,F(1,10))/F(1,10))
            check(near(f(r['recharge'])-b,ch,2e-4),directory.name+'_charge_'+k)
            endu=b if mode=='T' else f(r['body_ready'])
            check(near(endu-b,turn[typ],2e-4),directory.name+'_body_turn_'+k)
            work[mode][u]+=b-a
            events[k,'U_'+typ]=(a,endu);events[k,'P_'+typ]=(a,f(r['recharge']))
    return {'plan':p,'routes':rt,'relays':rr,'components':cc,'M_T':mt,'M':len(cc),'taskunit':taskunit,
            'events':events,'work':work,'relations':relations,'slots':slots}

def calc(c,groups,inv):
    G=len(groups);d={r:0 for r in R};groupres=[];work={'T':[],'R':[]}
    for group in groups:
        nr={r:sweep([(a,b,k) for (k,t),(a,b) in c['events'].items() if t==r and c['taskunit'][k] in group]) for r in R}
        groupres.append(nr)
        for r in R:d[r]+=nr[r]
        for mode in ['T','R']:work[mode].append(sum((c['work'][mode][i] for i in group),F0))
    bc={mode:F(G,2*(G-1))*sum(abs(w/sum(ws)-F(1,G)) for w in ws) for mode,ws in work.items() if sum(ws)}
    a=sum(F(d[r],inv[r]) for r in R)/len(R);b=sum(bc.values())/len(bc)
    return {'groups':groups,'D':d,'groupres':groupres,'A':a,'BT':bc.get('T',F0),'BR':bc.get('R',F0),'B':b,'Q':int(a*96),
            'H_stock':sum(F(max(0,d[r]-inv[r]),inv[r]) for r in R),
            'partition_code':'|'.join('+'.join(map(str,g)) for g in groups),'inventory_feasible':all(d[r]<=inv[r] for r in R)}

def enumerate_independent(c,G,inv):
    # Brute-force all labeled assignments, canonicalize, and deduplicate.
    parts=set()
    for labels in product(range(G),repeat=c['M']):
        if set(labels)!=set(range(G)):continue
        groups=tuple(sorted([tuple(i for i,v in enumerate(labels) if v==h) for h in range(G)]))
        parts.add(groups)
    return [calc(c,p,inv) for p in sorted(parts)]

def decision(rows,delta=F(3,10),weights=None,inv=None):
    if weights:
        rows=[{**r,'A':sum(weights[k]*F(r['D'][k],inv[k]) for k in R)} for r in rows]
    theta=[max(F0,F(1,2)-delta/2),min(F(1),F(1,2)+delta/2)]
    v={t:min(t*r['A']+(1-t)*r['B'] for r in rows) for t in theta}
    for r in rows:
        r['regret_lo']=theta[0]*r['A']+(1-theta[0])*r['B']-v[theta[0]]
        r['regret_hi']=theta[1]*r['A']+(1-theta[1])*r['B']-v[theta[1]]
        r['regret_max']=max(r['regret_lo'],r['regret_hi'])
    nd=[r for r in rows if not any(s['A']<=r['A'] and s['B']<=r['B'] and (s['A']<r['A'] or s['B']<r['B']) for s in rows)]
    best=min(nd,key=lambda r:(r['regret_max'],r['H_stock'],r['B'],r['Q'],r['partition_code']))
    return best,nd

def verify():
    params=parameters();inv=params[0]
    manifest=read(ROOT/'输入快照冻结清单.json')
    for rel,h in {**manifest['immutable_inputs'],**manifest['helpers']}.items():check(sha(ROOT/rel)==h,'frozen_hash_'+rel)
    for row in read(ROOT/'输入快照/源文件清单.json'):
        check(sha(row['original'])==row.get('original_sha256',row['sha256']),'original_unchanged_'+row['original'])
        check(sha(row['copy'])==row['sha256'],'source_copy_'+row['copy'])
    main=ROOT/'输入快照/主方案';c=rebuild(main,params)
    result=read(output('最终代表方案.json'));all_saved=read(output('完整分区档案.json'))
    exact_csv=csvread(output('分区枚举明细.csv'));expected_rows={}
    check(c['components']==result['components'],'component_graph_matches')
    check(c['M_T']==result['M_T'] and c['M']==result['M'],'component_counts_match')
    for G in [2,3]:
        rows=enumerate_independent(c,G,inv);best,nd=decision(rows);expected_rows[G]=rows
        check(len(rows)==len(all_saved[str(G)]),'enumeration_count_G'+str(G))
        check(best['partition_code']==result['selection'][str(G)]['partition_code'],'minimax_choice_G'+str(G))
        for row in rows:
            saved=next(x for x in exact_csv if int(x['G'])==G and x['partition_code']==row['partition_code'])
            for k,v in json.loads(saved['exact_rational']).items():check(row[k]==F(v),'rational_'+str(G)+'_'+row['partition_code']+'_'+k)
            check(row['D']==json.loads(saved['D']),'resource_vector_'+str(G)+'_'+row['partition_code'])
        check(not any(x['inventory_feasible'] for x in rows),'no_stock_feasible_partition_G'+str(G))
    # Raw box identity and deadlines; inherited spatial paths remain immutable.
    w=load_workbook(ROOT/'输入快照/原始附件/物资需求与配送时限.xlsx',read_only=True,data_only=True)
    raw={r[0]:r for r in list(w['逐箱货箱清单'].values)[1:] if r[0]};w.close()
    bp={b['box']:b for b in c['plan']['boxes']};check(set(bp)==set(raw),'all_raw_boxes_unique')
    hard=0
    for k,b in bp.items():
        r=raw[k];check((b['zone'],b['kind'],b['mass'],b['volume'])==(r[1],r[2],r[3],r[4]),'box_identity_'+k)
        ishard=r[2]=='医疗物资' or r[5]=='是';deadline=min([r[7]]+([r[6]] if r[6] is not None else [])) if ishard else r[7]
        check(b['hard']==ishard,'box_hard_flag_'+k)
        if ishard:hard+=1;check(b['delivery']<=deadline+2e-4,'box_deadline_'+k)
        route=c['routes'][b['route_id']]
        check(near(b['delivery'],route['start']+route['offsets'][b['category']],2e-4),'box_delivery_'+k)
    check(len(bp)==80 and hard==31,'80_boxes_31_hard')
    ET=sum(r['energy'] for r in c['routes'].values());ER=sum(r['energy'] for r in c['relays'].values())
    late=[b for b in bp.values() if not b['hard']]
    # The soft objective uses the expected-delivery time; ``deadline`` is
    # intentionally +inf for non-first (soft) boxes in the frozen JSON.
    L=sum(b['priority']*max(0,b['delivery']-b['expected']) for b in late)/sum(b['priority'] for b in late)
    obj=[len(c['routes'])+len(c['relays']),ET+ER,max(r['finish'] for r in [*c['routes'].values(),*c['relays'].values()]),L]
    check(all(near(a,b,2e-4) for a,b in zip(obj,result['objective'])),'q3_objectives_recomputed')
    witness=csvread(output('组内资源分配见证.csv'));taskrows=csvread(output('分组任务明细.csv'))
    certs=read(output('资源峰值证书.json'));boxrows=csvread(output('逐箱继承核验.csv'))
    for G in [2,3]:
        sol=result['solutions'][str(G)];groupzone={};counts={};by=defaultdict(list)
        for h,group in enumerate(sol['groups'],1):
            gid=f'G{G}-{h:02}';check(bool(group['services']),'nonempty_'+gid)
            for z in group['services']:check(z not in groupzone,'unique_zone_'+gid+z);groupzone[z]=gid
            counts[gid]=group['resources']
        check(set(groupzone)==set(z for cc in c['components'] for z in cc),'all_zones_G'+str(G))
        ws=[r for r in witness if int(r['G'])==G]
        check(Counter((r['task'],r['resource_type']) for r in ws)==Counter({k:1 for k in c['events']}),'event_completeness_G'+str(G))
        for row in ws:
            k=row['task'];typ=row['resource_type'];a,b=c['events'][k,typ]
            check((f(row['start']),f(row['ready']))==(a,b),'unchanged_resource_interval_'+str(G)+k+typ)
            units=c['components'][c['taskunit'][k]];gid=groupzone[units[0]]
            check(row['group']==gid and row['new_resource'].startswith(gid+'-'),'group_exclusive_'+row['new_resource'])
            by[row['new_resource']].append((a,b,k))
        for identity,ee in by.items():check(sweep(ee)==1,'nonoverlap_'+identity)
        for gid,nr in counts.items():
            for r in R:
                ids={v['new_resource'] for v in ws if v['group']==gid and v['resource_type']==r}
                check(len(ids)==nr[r],'minimal_assignment_'+gid+r)
                cert=next(x for x in certs if x['G']==G and x['group']==gid and x['resource_type']==r)
                check(cert['minimum_count']==nr[r],'peak_count_'+gid+r)
                for interval in cert['peak_intervals']:
                    a=f(interval['start']);b=f(interval['end'])
                    active={v['task'] for v in ws if v['group']==gid and v['resource_type']==r and f(v['start'])<=a<f(v['ready'])}
                    check(a<b and active==set(interval['tasks']) and len(active)==nr[r],'peak_witness_'+gid+r)
        tr=[r for r in taskrows if int(r['G'])==G]
        check(Counter(r['task'] for r in tr)==Counter({k:1 for k in c['taskunit']}),'each_task_once_G'+str(G))
        for row in tr:
            source=c['routes'].get(row['task'],c['relays'].get(row['task']))
            check(all(f(row[key])==f(source[other]) for key,other in [('start','start'),('finish','finish'),('energy_kwh','energy')]),'task_frozen_'+str(G)+row['task'])
        taskgroup={r['task']:r['group'] for r in tr}
        for j,ks in c['relations'].items():check(all(taskgroup[j]==taskgroup[k] for k in ks),'actual_relay_same_group_'+str(G)+j)
        br=[r for r in boxrows if int(r['G'])==G];check(Counter(r['box'] for r in br)==Counter({k:1 for k in bp}),'box_output_unique_G'+str(G))
        for row in br:check(f(row['delivery'])==f(bp[row['box']]['delivery']) and row['group']==taskgroup[row['task']],'box_output_'+str(G)+row['box'])
    # Independently audit the subproblem-2 accounting and its merge witnesses.
    gaprows=csvread(output('库存缺口分析.csv'));mergerows=csvread(output('合并复用收益.csv'))
    for G in [2,3]:
        selected=result['solutions'][str(G)]
        for r in R:
            central=sweep([(a,b,k) for (k,t),(a,b) in c['events'].items() if t==r])
            row=next(x for x in gaprows if int(x['G'])==G and x['resource']==r)
            demand=selected['D'][r]
            check([int(row[k]) for k in ['central_P','demand','split_loss','gap','stock_left']]==
                  [central,demand,demand-central,max(demand-inv[r],0),max(inv[r]-demand,0)],'gap_accounting_'+str(G)+r)
            gain=0
            for x in (v for v in mergerows if int(v['G'])==G and v['resource']==r):
                left=set(json.loads(x['left_units']));right=set(json.loads(x['right_units']))
                check(not left&right,'merge_disjoint')
                def peak_units(units):
                    return sweep([(a,b,k) for (k,t),(a,b) in c['events'].items() if t==r and c['taskunit'][k] in units])
                val=peak_units(left)+peak_units(right)-peak_units(left|right)
                check(val==int(x['reuse_gain']) and val>=0,'merge_gain_'+str(G)+r)
                gain+=val
            check(gain==demand-central,'merge_telescope_'+str(G)+r)
    # Check epsilon records and reconstruct the exact LP bound certificate.
    for log in read(output('MILP与LP核验.json')):
        G=log['G'];rows=expected_rows[G]
        cols=[u for k in range(1,c['M']-G+2) for u in combinations(range(c['M']),k)]
        bcol=[];qcol=[]
        for u in cols:
            b=F0
            for mode in ['T','R']:
                total=sum(c['work'][mode]);share=sum(c['work'][mode][i] for i in u)/total
                b+=abs(share-F(1,G))*F(G,4*(G-1))
            bcol.append(b)
            qcol.append(sum(F(12,inv[r])*sweep([(a,z,k) for (k,t),(a,z) in c['events'].items() if t==r and c['taskunit'][k] in u]) for r in R))
        for record in log['epsilon']:
            eligible=[r for r in rows if r['Q']<=record['epsilon']]
            if not eligible:check(record['status']=='INTEGER_INFEASIBLE','epsilon_infeasible');continue
            best=min(eligible,key=lambda r:(r['B'],r['Q'],r['H_stock'],r['partition_code']))
            check(best['partition_code']==record['partition_code'],'epsilon_choice')
            y=list(map(F,record['dual_y_exact']));t=F(record['dual_t_exact']);check(t<=0,'LP_dual_sign')
            lower=sum(y[:-1])+G*y[-1]+record['epsilon']*t
            lower+=sum(min(F0,bcol[j]-sum(y[i] for i in u)-y[-1]-qcol[j]*t) for j,u in enumerate(cols))
            check(lower==F(record['LP_bound_exact']) and lower<=best['B'],'LP_rational_lower_bound')
    # Alternative snapshots independently repeated, precision holds never edited.
    for alt in read(output('备选快照比较.json')):
        d=main if alt['snapshot']=='主方案' else ROOT/'输入快照/备选方案'/alt['snapshot']
        if alt['status']=='PRECISION_PENDING':
            p=read(d/'方案.json');items={r.get('route_id',r.get('relay_id')):r for r in p['routes']+p['relays']}
            for overlap in alt['overlap_records']:
                aa=items[overlap['previous']];bb=items[overlap['next']]
                end=aa['finish'] if overlap['type']=='U' and 'route_id' in aa else aa['body_ready'] if overlap['type']=='U' else aa['recharge']
                check(f(end)-f(bb['start'])==F(overlap['exact_overlap'])>0,'alternative_precision_record')
            continue
        ac=rebuild(d,params)
        for G in [2,3]:
            rows=enumerate_independent(ac,G,inv);best,_=decision(rows)
            check(best['partition_code']==alt['cases'][str(G)]['chosen']['partition_code'],'alternative_choice_'+alt['snapshot']+str(G))
            for row in alt['cases'][str(G)]['rows']:
                found=next(x for x in rows if x['partition_code']==row['partition_code'])
                check(found['D']==row['D'] and all(found[k]==F(v) for k,v in row['exact_rational'].items()),'alternative_exact_row')
            for record in (r for r in alt['beta_rows'] if r['G']==G):
                eligible=[r for r in rows if r['B']<=f(record['beta'])]
                if not eligible:check(record['status']=='BALANCE_INFEASIBLE','beta_infeasible')
                else:check(min(eligible,key=lambda r:(r['A'],r['B'],r['H_stock'],r['partition_code']))['partition_code']==record['partition_code'],'beta_choice')
    for rec in read(output('偏好与类别权重敏感性.json')):
        rows=[dict(r) for r in expected_rows[rec['G']]]
        if rec['experiment']=='preference_delta':best,_=decision(rows,F(rec['value']))
        else:
            types={'transport_uav':['U_A','U_B','U_C'],'transport_battery':['P_A','P_B','P_C'],'relay':['U_R','P_R']}[rec['value']]
            raw={r:F(3,2) if r in types else F(1) for r in R};weights={r:v/sum(raw.values()) for r,v in raw.items()}
            best,_=decision(rows,weights=weights,inv=inv)
        check(best['partition_code']==rec['choice']['partition_code'],'sensitivity_choice')
    original=load_workbook(ROOT/'输入快照/原始附件/结果提交模板.xlsx');export=load_workbook(output('问题四_结果提交.xlsx'))
    check(original.sheetnames==export.sheetnames,'workbook_sheets')
    for name in original.sheetnames:
        if name!='Q4_分区配置':check(list(original[name].values)==list(export[name].values),'template_preserved_'+name)
    tab=export['Q4_分区配置'];actual=[list(row[:11]) for row in list(tab.values)[1:] if row[0] is not None]
    expect=[]
    for G in [2,3]:
        for h,group in enumerate(result['solutions'][str(G)]['groups'],1):
            expect.append([G,f'G{G}-{h:02}','、'.join(group['services']),*[group['resources'][r] for r in R]])
    check(actual==expect and len(actual)==5,'submission_exactly_five_selected_groups')
    check(all(isinstance(x,int) and x>=0 for row in actual for x in row[3:]),'submission_integer_counts')
    original.close();export.close()
    return {'partition_counts':{str(g):len(rows) for g,rows in expected_rows.items()},'main_components':c['components'],
            'resource_assignment_rows':len(witness),'physical_scope':'Inherited Q3 physical/DEM certificates and same-snapshot reconstructed communication; no new independent propagation algorithm.',
            'stock_feasible_partitions':{str(g):sum(r['inventory_feasible'] for r in rows) for g,rows in expected_rows.items()},
            'boxes':len(bp),'hard_boxes':hard,'objective_recomputed':obj}

if __name__=='__main__':
    report={}
    try:
        report=verify();report['status']='PASS'
    except Exception as exc:
        report={'status':'FAIL','error':repr(exc),'traceback':traceback.format_exc()}
    report['checks']=len(checks);report['failed_checks']=[r for r in checks if not r['pass']]
    (output('核验报告.json')).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    (output('逐项核验记录.json')).write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))
    if report['status']!='PASS':raise SystemExit(1)
