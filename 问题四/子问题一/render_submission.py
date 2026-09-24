"""Render the actual five exported Q4 workbook rows for visual inspection."""
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from q4_paths import output
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image,ImageDraw,ImageFont
import re

def main():
    w=load_workbook(output('问题四_结果提交.xlsx'));s=w['Q4_分区配置']
    scale=2;pad=44
    widths=[round(s.column_dimensions[get_column_letter(j)].width*7*scale) for j in range(1,12)]
    heights=[round(s.row_dimensions[i].height*96/72*scale) for i in range(1,7)]
    im=Image.new('RGB',(sum(widths)+2*pad,sum(heights)+220),'white');d=ImageDraw.Draw(im)
    regular=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',26)
    bold=ImageFont.truetype('C:/Windows/Fonts/msyhbd.ttc',26)
    title=ImageFont.truetype('C:/Windows/Fonts/msyhbd.ttc',36)
    d.text((pad,24),'问题四 · 分区配置提交表',font=title,fill='#203E57')
    d.text((pad,78),'已选定的两组与三组代表方案｜数量为独立执行需求，含需补配资源',font=regular,fill='#566574')
    def wrap(text,font,width):
        lines=[];line=''
        for token in re.findall(r'S\d{3}、?|[A-Za-z0-9]+|.',text):
            if d.textlength(line+token,font=font)>width and line:lines.append(line);line=token
            else:line+=token
        return lines+[line]
    y=130
    for i,h in enumerate(heights,1):
        x=pad
        for j,ww in enumerate(widths,1):
            c=s.cell(i,j);font=bold if i==1 else regular
            bg='#'+c.fill.fgColor.rgb[-6:];fg='#'+c.font.color.rgb[-6:]
            d.rectangle([x,y,x+ww,y+h],fill=bg,outline='#D9E1E8',width=1)
            lines=wrap(str(c.value),font,ww-24);lh=36
            assert len(lines)*lh<=h-12,(i,j,lines,h)
            yy=y+(h-lh*len(lines))/2
            for line in lines:
                xx=x+12 if j==3 and i>1 else x+(ww-d.textlength(line,font=font))/2
                d.text((xx,yy),line,font=font,fill=fg);yy+=lh
            x+=ww
        y+=h
    d.text((pad,y+20),'来源：问题四_结果提交.xlsx / Q4_分区配置；按工作簿实际单元格读出。',font=regular,fill='#566574')
    im.save(output('提交表预览.png'));w.close()
    print(output('提交表预览.png'))

if __name__=='__main__':main()
