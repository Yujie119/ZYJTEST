"""Export one selected Q2 solution plus a separate archive using openpyxl.

Run with the configured Python.  Sources are read-only; the original submission
workbook is never overwritten.  Formula caches are evaluated from workbook cells
by a deliberately small parser, then checked using openpyxl(data_only=True).
"""
from __future__ import annotations

import argparse
import ast
import copy
import hashlib
import json
import math
import re
import textwrap
import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.workbook.properties import CalcProperties
from PIL import Image, ImageDraw, ImageFont


HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NUM = "#,##0.000000"
TOL = 1e-6


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def close(a, b):
    return math.isclose(float(a), float(b), rel_tol=1e-10, abs_tol=TOL)


def objective(solution):
    routes, boxes = solution["routes"], solution["boxes"]
    soft = [b for b in boxes if not b["hard"]]
    den = sum(b["priority"] for b in soft)
    if den <= 0:
        raise ValueError("The selected instance has no positive soft-priority denominator")
    return [len(routes), sum(r["energy"] for r in routes),
            max(r["finish"] for r in routes),
            sum(b["priority"] * max(0, b["delivery"] - b["expected"])
                for b in soft) / den]


def fingerprint(solution):
    return tuple(sorted((r["candidate_id"], r["uav"], r["battery_id"],
                         round(r["start"], 8)) for r in solution["routes"]))


def inventory():
    paths = list((ROOT / "D题").rglob("运输无人机数据.xlsx"))
    if len(paths) != 1:
        raise ValueError(f"Expected one raw transport workbook: {paths}")
    rows = list(openpyxl.load_workbook(paths[0], data_only=True).active.values)
    uavs = {r[0]: r[1] for r in rows if r[0] and re.fullmatch(r"U\d+", str(r[0]))}
    start = next(i for i, r in enumerate(rows) if r[0] == "共享电池库存")
    batteries = {f"{r[0]}-BAT-{i:02d}": r[0]
                 for r in rows[start + 2:] if r[0]
                 for i in range(1, int(r[1]) + 1)}
    return uavs, batteries, paths[0]


def validate_solution(solution, uavs, batteries):
    routes, boxes = solution["routes"], solution["boxes"]
    if not routes or not boxes:
        raise ValueError("Empty selected solution")
    if len({r["route_id"] for r in routes}) != len(routes):
        raise ValueError("Duplicate route IDs")
    if len({b["box"] for b in boxes}) != len(boxes):
        raise ValueError("Duplicate box IDs")
    route_map = {r["route_id"]: r for r in routes}
    for r in routes:
        if uavs.get(r["uav"]) != r["vehicle"] or batteries.get(r["battery_id"]) != r["vehicle"]:
            raise ValueError("Resource type mismatch")
        for field in ("start", "finish", "energy", "duration", "charge_s", "soc"):
            if not isinstance(r[field], (float, int)) or not math.isfinite(r[field]):
                raise ValueError(f"Nonfinite route field {field}")
        if not close(r["finish"], r["start"] + r["duration"]):
            raise ValueError("Inconsistent route finish")
        if not close(r["recharge"], r["finish"] + r["charge_s"]):
            raise ValueError("Inconsistent recharge completion")
    for b in boxes:
        r = route_map[b["route_id"]]
        if b["box"] not in r["box_ids"]:
            raise ValueError("Box not on its assigned route")
        if not close(b["delivery"], r["start"] + r["offsets"][b["box"]]):
            raise ValueError("Inconsistent delivery time")
        if b["hard"] and b["delivery"] > b["deadline"] + TOL:
            raise ValueError("Hard deadline violated")
    for field, end in (("uav", "finish"), ("battery_id", "recharge")):
        groups = defaultdict(list)
        for r in routes:
            groups[r[field]].append(r)
        for seq in groups.values():
            seq.sort(key=lambda r: r["start"])
            for a, b in zip(seq, seq[1:]):
                if a[end] > b["start"] + TOL:
                    raise ValueError(f"Overlapping {field} intervals")
    result = objective(solution)
    if not all(close(x, y) for x, y in zip(result, solution["objective"])):
        raise ValueError("Source objective inconsistent with source rows")
    return result


class FormulaEngine:
    """Only the Excel syntax emitted below; no arbitrary code evaluation."""
    TOKEN = re.compile(
        r"\s*(?:(?P<ref>(?:'(?:[^']|'')+'!)?\$?[A-Z]{1,3}\$?[0-9]+"
        r"(?::\$?[A-Z]{1,3}\$?[0-9]+)?)|(?P<num>\d+(?:\.\d*)?(?:[eE][+-]?\d+)?)"
        r'|(?P<str>"(?:[^"]|"")*")|(?P<name>[A-Za-z_][A-Za-z_0-9]*)'
        r"|(?P<op><=|>=|<>|[+*/(),=<>-]))"
    )

    def __init__(self, workbook):
        self.wb, self.cache, self.active = workbook, {}, set()

    def cell(self, sheet, address):
        key = (sheet, address)
        if key in self.cache:
            return self.cache[key]
        value = self.wb[sheet][address].value
        if not isinstance(value, str) or not value.startswith("="):
            return value
        if key in self.active:
            raise ValueError(f"Circular formula: {key}")
        self.active.add(key)
        val = self.parse(value[1:], sheet)
        self.active.remove(key)
        if isinstance(val, float) and not math.isfinite(val):
            raise ValueError(f"Nonfinite formula: {key}")
        self.cache[key] = val
        return val

    def parse(self, source, sheet):
        tokens, at = [], 0
        while at < len(source):
            m = self.TOKEN.match(source, at)
            if not m:
                raise ValueError(f"Unsupported formula syntax: {source[at:]}")
            tokens.append((m.lastgroup, m.group(m.lastgroup)))
            at = m.end()
        cursor = 0

        def flatten(args):
            return [v for a in args for v in (a if isinstance(a, list) else [a])]

        def expr(minimum=0):
            nonlocal cursor
            kind, token = tokens[cursor]
            cursor += 1
            if token in ("+", "-"):
                left = expr(30)
                if token == "-": left = -left
            elif token == "(":
                left = expr()
                if tokens[cursor][1] != ")": raise ValueError("Missing closing parenthesis")
                cursor += 1
            elif kind == "num":
                left = float(token)
            elif kind == "str":
                left = token[1:-1].replace('""', '"')
            elif kind == "ref":
                target, address = sheet, token
                if "!" in token:
                    target, address = token.rsplit("!", 1)
                    target = target[1:-1].replace("''", "'")
                address = address.replace("$", "")
                if ":" in address:
                    c1, r1, c2, r2 = range_boundaries(address)
                    left = [self.cell(target, f"{get_column_letter(c)}{r}")
                            for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
                else:
                    left = self.cell(target, address)
            elif kind == "name":
                name = token.upper()
                if name in ("TRUE", "FALSE"):
                    left = name == "TRUE"
                else:
                    if tokens[cursor][1] != "(": raise ValueError("Expected function arguments")
                    cursor += 1
                    args = []
                    while tokens[cursor][1] != ")":
                        args.append(expr())
                        if tokens[cursor][1] != ",": break
                        cursor += 1
                    if tokens[cursor][1] != ")": raise ValueError("Expected )")
                    cursor += 1
                    values = flatten(args)
                    numbers = [float(x) for x in values if isinstance(x, (int, float)) and not isinstance(x, bool)]
                    if name == "SUM": left = sum(numbers)
                    elif name == "MAX": left = max(numbers, default=0)
                    elif name == "MIN": left = min(numbers, default=0)
                    elif name == "COUNTA": left = sum(x is not None and x != "" for x in values)
                    elif name == "IF": left = args[1] if args[0] else args[2]
                    else: raise ValueError(f"Unsupported function {name}")
            else:
                raise ValueError(f"Unexpected token {token}")
            precedence = {"=": 5, "<>": 5, "<": 5, ">": 5, "<=": 5, ">=": 5,
                          "+": 10, "-": 10, "*": 20, "/": 20}
            while cursor < len(tokens):
                op = tokens[cursor][1]
                power = precedence.get(op, -1)
                if power < minimum: break
                cursor += 1
                right = expr(power + 1)
                # Excel treats referenced empty cells as zero in arithmetic.
                if op in ("+", "-", "*", "/"):
                    left = 0 if left is None else left
                    right = 0 if right is None else right
                if op == "+": left = left + right
                elif op == "-": left = left - right
                elif op == "*": left = left * right
                elif op == "/": left = left / right
                elif op == "=": left = left == right
                elif op == "<>": left = left != right
                elif op == "<": left = left < right
                elif op == ">": left = left > right
                elif op == "<=": left = left <= right
                elif op == ">=": left = left >= right
            return left

        answer = expr()
        if cursor != len(tokens): raise ValueError("Trailing formula tokens")
        return answer

    def evaluate_all(self):
        for ws in self.wb:
            for row in ws:
                for cell in row:
                    if cell.data_type == "f": self.cell(ws.title, cell.coordinate)
        return self.cache


def install_caches(path, cache, sheetnames):
    """Install independently evaluated <v> values without removing <f>."""
    staged = path.with_suffix(".cache.tmp.xlsx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(staged, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            content = zin.read(item.filename)
            match = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", item.filename)
            if match:
                title = sheetnames[int(match.group(1)) - 1]
                root = ET.fromstring(content)
                for cell in root.findall(f".//{{{NS}}}c"):
                    key = (title, cell.get("r"))
                    if key not in cache: continue
                    value = cache[key]
                    node = cell.find(f"{{{NS}}}v")
                    if node is None: node = ET.SubElement(cell, f"{{{NS}}}v")
                    if isinstance(value, bool):
                        cell.set("t", "b"); node.text = "1" if value else "0"
                    elif isinstance(value, (int, float)):
                        cell.set("t", "n"); node.text = repr(value)
                    else:
                        cell.set("t", "str"); node.text = "" if value is None else str(value)
                content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            zout.writestr(item, content)
    staged.replace(path)


def preview(wb, path, engine=None, names=None):
    """Compact contact sheet rendered from workbook values/styles using Pillow."""
    fonts = [Path("C:/Windows/Fonts/msyh.ttc"), Path("C:/Windows/Fonts/simhei.ttf")]
    font_path = next(p for p in fonts if p.exists())
    normal = ImageFont.truetype(str(font_path), 16)
    bold = ImageFont.truetype(str(font_path), 19)
    names = names or wb.sheetnames
    canvas = Image.new("RGB", (1660, 242 * len(names) + 20), "white")
    draw = ImageDraw.Draw(canvas)
    for panel, name in enumerate(names):
        ws = wb[name]; y0 = 15 + panel * 242
        draw.text((18, y0), name, font=bold, fill="#294B60")
        ncols = min(ws.max_column, 8)
        width = 1620 / max(1, ncols)
        for ri, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 5), max_col=ncols)):
            for ci, cell in enumerate(row):
                x, y = 18 + ci * width, y0 + 32 + ri * 37
                fill = "#294B60" if ri == 0 else ("#F2F6F8" if ri % 2 else "white")
                draw.rectangle((x, y, x + width - 2, y + 35), fill=fill)
                val = engine.cell(name, cell.coordinate) if engine else cell.value
                if isinstance(val, float): val = f"{val:,.3f}"
                txt = "" if val is None else str(val)
                limit = int(width // 9)
                if len(txt) > limit: txt = txt[:max(2, limit - 1)] + "…"
                draw.text((x + 5, y + 6), txt, font=normal, fill="white" if ri == 0 else "#172C39")
    canvas.save(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=HERE)
    parser.add_argument("--template", type=Path, default=HERE.parent / "问题二_结果提交.xlsx")
    parser.add_argument("--output", type=Path)
    parser.add_argument("--no-preview", action="store_true")
    args = parser.parse_args()
    folder = args.input_dir.resolve()
    output = (args.output or folder / "问题二_改进搜索结果.xlsx").resolve()
    template = args.template.resolve()
    if output == template or output == (HERE.parent / "问题二_结果提交.xlsx").resolve():
        raise ValueError("Refusing to overwrite the original submission workbook")
    output.parent.mkdir(parents=True, exist_ok=True)
    before_template = sha(template)
    final_path, archive_path = folder / "最终方案.json", folder / "联合非支配档案.json"
    final, archive = read_json(final_path), read_json(archive_path)
    if not isinstance(archive, list): raise ValueError("Expected an archive list")
    uavs, batteries, raw_inventory = inventory()
    expected = validate_solution(final, uavs, batteries)
    for solution in archive:
        validate_solution(solution, uavs, batteries)
    matches = [i for i, solution in enumerate(archive) if fingerprint(solution) == fingerprint(final)]
    if not matches: raise ValueError("Final solution is absent from archive")
    final_index = matches[0]
    source = openpyxl.load_workbook(template)
    if not args.no_preview:
        preview(source, folder / "导出_原模板预览.png", names=["Q2_运输架次", "核验_逐箱", "核验_资源汇总"])
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    header_style = copy.copy(source["Q2_运输架次"]["A1"]._style)
    body_style = copy.copy(source["Q2_运输架次"]["A2"]._style)

    def sheet(name, headers, template_name=None, widths=None):
        ws = wb.create_sheet(name)
        template_sheet = source[template_name] if template_name in source.sheetnames else None
        ws.append(headers)
        for col, cell in enumerate(ws[1], 1):
            cell._style = copy.copy(template_sheet.cell(1, min(col, template_sheet.max_column))._style) if template_sheet else copy.copy(header_style)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(col)].width = (widths or {}).get(col, 18)
        ws.row_dimensions[1].height = 38
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:1"
        return ws

    summary = sheet("指标汇总", ["指标", "Excel公式结果", "JSON记录/独立复算", "差值", "口径"], widths={1:27, 2:24, 3:25, 4:20, 5:78})
    rs = sheet("Q2_运输架次", [c.value for c in source["Q2_运输架次"][1]], "Q2_运输架次", {1:18, 2:16, 3:16, 4:18, 5:21, 6:54, 7:23, 8:22})
    bs = sheet("Q2_逐箱交付", [c.value for c in source["Q2_逐箱交付"][1]], "Q2_逐箱交付", {1:23, 2:18, 3:18, 4:25})
    rd = sheet("核验_架次详情", [c.value for c in source["核验_架次详情"][1]] + ["任务总时长s", "充电时长s", "访问次数", "箱数"], "核验_架次详情", {1:18, 2:24, 3:78, 4:18, 5:19, 6:21, 7:17, 8:24, 9:21, 10:21})
    bd = sheet("核验_逐箱", [c.value for c in source["核验_逐箱"][1]] + ["软箱优先系数", "软箱加权迟到s", "非医疗优先系数", "医疗标记"], "核验_逐箱", {1:23, 2:18, 3:18, 4:21})
    events = sheet("核验_资源事件", [c.value for c in source["核验_资源事件"][1]] + ["时长s"], "核验_资源事件", {1:18, 2:20, 3:18, 4:18, 5:23, 6:23, 7:23})
    resources = sheet("核验_资源汇总", [c.value for c in source["核验_资源汇总"][1]], "核验_资源汇总", {1:24, 2:21, 3:17, 4:16, 5:23, 6:23, 7:23, 8:23, 9:21})
    pf = sheet("联合非支配档案", ["方案编号", "最终提交标记", "架次数N", "能耗kWh", "最晚返场s", "加权迟到s", "软箱迟到箱数", "遗憾下界", "遗憾上界", "方案来源"], widths={1:18, 2:20, 3:17, 4:22, 5:23, 6:23, 7:20, 8:21, 9:21, 10:75})
    notes = sheet("求解说明", ["项目", "内容"], "求解说明", {1:34, 2:112})
    routes = sorted(final["routes"], key=lambda r: (r["start"], r["route_id"]))
    boxes = sorted(final["boxes"], key=lambda b: b["box"])
    row_map = {}
    for row, r in enumerate(routes, 2):
        row_map[r["route_id"]] = row
        rs.append([r["route_id"], r["uav"], r["vehicle"], r["battery_id"], r["start"],
                   "→".join(["O01"] + r["zones"] + ["O01"]), r["finish"], r["energy"]])
        rd.append([r["route_id"], r["candidate_id"], ";".join(r["box_ids"]), r["mass"], r["volume"],
                   r["takeoff"], r["soc"], r["recharge"],
                   f"='Q2_运输架次'!G{row}-'Q2_运输架次'!E{row}",
                   f"=H{row}-'Q2_运输架次'!G{row}", len(r["events"]), len(r["box_ids"])])
        rd.cell(row, 7).number_format = "0.00%"
    for row, b in enumerate(boxes, 2):
        bs.append([b["box"], b["route_id"], b["zone"], b["delivery"]])
        bd.append([b["box"], b["route_id"], b["zone"], b["kind"],
                   f"='Q2_逐箱交付'!D{row}", b["expected"], bool(b["hard"]), b["deadline"],
                   bool(b["first"]), b["priority"], f'=IF(G{row},H{row}-E{row},"")',
                   f"=MAX(0,E{row}-F{row})", f"=IF(G{row},0,J{row})", f"=M{row}*L{row}",
                   f"=IF(P{row},0,J{row})", b["kind"] == "医疗物资"])
    event_rows = []
    for r in routes:
        i = row_map[r["route_id"]]
        for typ, resource, phase, begin, end in [
            ("无人机", r["uav"], "任务", f"'Q2_运输架次'!E{i}", f"'Q2_运输架次'!G{i}"),
            ("电池", r["battery_id"], "任务", f"'Q2_运输架次'!E{i}", f"'Q2_运输架次'!G{i}"),
            ("电池", r["battery_id"], "充电", f"'Q2_运输架次'!G{i}", f"'核验_架次详情'!H{i}"),
        ]:
            event_rows.append((typ, resource, r["start"], phase, r["route_id"], begin, end))
    for row, (typ, resource, _, phase, rid, begin, end) in enumerate(sorted(event_rows), 2):
        events.append([typ, resource, rid, phase, "=" + begin, "=" + end, f"=F{row}-E{row}"])

    def refs(rows, sheetname, col):
        return ",".join(f"'{sheetname}'!{col}{i}" for i in rows)

    for typ, mapping, field in [("实体运输无人机", uavs, "uav"), ("共享电池", batteries, "battery_id")]:
        for ident, vehicle in sorted(mapping.items()):
            ids = [row_map[r["route_id"]] for r in routes if r[field] == ident]
            row = resources.max_row + 1
            resources.append([
                typ, ident, vehicle,
                f"=COUNTA({refs(ids, 'Q2_运输架次', 'A')})" if ids else 0,
                f"=SUM({refs(ids, '核验_架次详情', 'I')})" if ids else 0,
                (f"=SUM({refs(ids, '核验_架次详情', 'J')})" if ids else 0) if field == "battery_id" else None,
                f"=MAX({refs(ids, 'Q2_运输架次', 'G')})" if ids else 0,
                (f"=MAX({refs(ids, '核验_架次详情', 'H')})" if ids else 0) if field == "battery_id" else None,
                f"=E{row}/'指标汇总'!B4" if field == "uav" else None,
            ])
            resources.cell(row, 9).number_format = "0.00%"
    endr, endb = len(routes) + 1, len(boxes) + 1
    metric_rows = [
        ["架次数N", f"=COUNTA('Q2_运输架次'!A2:A{endr})", expected[0], "=B2-C2", "仅此最终方案的运输架次"],
        ["运输总能耗 kWh", f"=SUM('Q2_运输架次'!H2:H{endr})", expected[1], "=B3-C3", "统一题设与补充能耗假设；不解释为实飞测量"],
        ["最晚返场 s", f"=MAX('Q2_运输架次'!G2:G{endr})", expected[2], "=B4-C4", "全部运输机完成最后任务并返回O01；不含最终充电"],
        ["加权平均迟到 s", f"=SUM('核验_逐箱'!N2:N{endb})/SUM('核验_逐箱'!M2:M{endb})", expected[3], "=B5-C5", "仅无硬时限箱；分母为软箱应急优先系数和"],
        ["交付货箱数", f"=COUNTA('Q2_逐箱交付'!A2:A{endb})", len(boxes), "=B6-C6", "货箱编号唯一性另由导出检查确认"],
        ["硬时限箱数", f"=SUM('核验_逐箱'!G2:G{endb})", sum(b["hard"] for b in boxes), "=B7-C7", "医疗或首批的并集"],
        ["软时限箱数", "=B6-B7", sum(not b["hard"] for b in boxes), "=B8-C8", "无硬时限箱"],
        ["软箱优先系数和", f"=SUM('核验_逐箱'!M2:M{endb})", sum(b["priority"] for b in boxes if not b["hard"]), "=B9-C9", "当前加权迟到分母"],
        ["非医疗优先系数和", f"=SUM('核验_逐箱'!O2:O{endb})", sum(b["priority"] for b in boxes if b["kind"] != "医疗物资"), "=B10-C10", "用于辨认64非医疗箱与49软箱两种口径"],
        ["档案方案数", len(archive), len(archive), "=B11-C11", "独立档案页；不混入提交架次表"],
        ["最终档案编号", f"PF{final_index + 1:04d}", None, None, "按完整架次、资源与开始时刻匹配"],
        ["遗憾下界", final.get("regret_lower"), None, None, "原JSON记录；范围见求解说明"],
        ["遗憾上界", final.get("regret_upper"), None, None, "原JSON记录；不代表原题全局稳健最优认证"],
    ]
    for row in metric_rows: summary.append(row)
    # Excel SUM(range) ignores logical cells; expose numeric flags for this count.
    summary["B7"] = "+".join(f"IF('核验_逐箱'!G{i},1,0)" for i in range(2, endb + 1))
    summary["B7"] = "=" + summary["B7"].value
    for i, solution in enumerate(archive):
        vals = objective(solution)
        pf.append([f"PF{i+1:04d}", i == final_index, *vals,
                   sum(not b["hard"] and b["delivery"] > b["expected"] + TOL for b in solution["boxes"]),
                   solution.get("regret_lower"), solution.get("regret_upper"), solution.get("source", "")])
    extra = {}
    for filename in ["汇总.json", "candidate_scope.json", "search_config.json", "独立核验.json"]:
        if (folder / filename).exists(): extra[filename] = read_json(folder / filename)
    run_summary, scope = extra.get("汇总.json", {}), extra.get("candidate_scope.json", {})
    review = extra.get("独立核验.json", {})
    review_obj = review.get("recomputed", {}).get("objective")
    review_matches = bool(review_obj and all(close(a, b) for a, b in zip(review_obj, expected)))
    note_rows = [
        ("本次结果定位", "仅O01换电的可行基准；不代表允许服务区途中换电的完整专家规则下的全局结果。"),
        ("专家补充（用户原话记录）", "不能携带备用电池；可在有电池地方更换不限O01；初始共享电池都在O01。"),
        ("基准能源安排", "本方案每架次绑定一块电池，全程不携带备用电池，换电均在O01；服务区初始电池库存为0。"),
        ("尚待明确的专家规则", "服务区能否充电、途中换电是否必须满电尚待确认；本导出不设服务区充电能力或额外库存。"),
        ("最终方案来源", final.get("source", "未记录")),
        ("最终方案文件", str(final_path)), ("最终方案SHA256", sha(final_path)),
        ("非支配档案文件", str(archive_path)), ("档案SHA256", sha(archive_path)),
        ("格式参考", str(template)), ("参考工作簿SHA256", before_template),
        ("资源库存依据", str(raw_inventory)), ("资源库存SHA256", sha(raw_inventory)),
        ("提交范围", "Q2_运输架次与Q2_逐箱交付仅含最终方案；联合非支配档案独立列示备选指标。"),
        ("仅O01换电基准的搜索范围", run_summary.get("scope", "扩大有限候选库中的启发式搜索；未认证全部可能路线的全局最优。")),
        ("初始候选库数", scope.get("total", "未记录")),
        ("最终评价候选库数", run_summary.get("route_count", "未记录")),
        ("最终候选系数SHA256", run_summary.get("library_coefficient_sha256", "未记录")),
        ("候选库是否完整", bool(scope.get("complete", False))),
        ("物理模型", run_summary.get("physics", "继承q2_data的明确闭合假设，导出不改物理参数。")),
        ("软迟到定义", "软箱为既非医疗也非首批；逐箱L列为正迟到，M列为软箱权，N列为两者乘积。"),
        ("资源占用", "无人机任务区间[准备开始,返场]；电池任务同区间，返场后充至满电才可复用。资源事件包括最后一次任务后的充电。"),
        ("时间单位", "所有任务时刻、持续时间、充电时间与迟到值均为秒；开始时刻指开始准备。"),
        ("返航SOC", "核验_架次详情G列为0–1数值，以百分比格式显示。"),
        ("遗憾解释", "来自输入JSON中共同档案/候选范围的上下界；导出不重新证明稳健最优性，不与异库界无条件比较。"),
        ("公式缓存", "Excel公式保留，Python有限公式解释器从工作簿单元格重算缓存；Excel打开时自动重算。"),
        ("独立核验是否匹配本方案", review_matches),
        ("匹配核验记录all", review.get("checks", {}).get("all") if review_matches else "未匹配；不能借用旧核验结论"),
        ("导出核验边界", "检查输入一致性、覆盖去重、资源重叠、硬时限、公式与typed数值；不替代DEM/物理模型独立审查。"),
    ]
    for key, val in note_rows: notes.append([key, val])
    for ws in wb:
        ws.auto_filter.ref = ws.dimensions
        ws.print_area = ws.dimensions
        for row in ws.iter_rows(min_row=2):
            height = 25
            for cell in row:
                fmt = cell.number_format
                cell._style = copy.copy(body_style)
                cell.font = Font(name="微软雅黑", size=10, color="172C39")
                cell.alignment = Alignment(vertical="center", horizontal="left" if isinstance(cell.value, str) and cell.data_type != "f" else "right", wrap_text=True)
                if isinstance(cell.value, (float, int)) and not isinstance(cell.value, bool) or cell.data_type == "f":
                    cell.number_format = fmt if fmt != "General" else NUM
                if isinstance(cell.value, str) and cell.data_type != "f":
                    width = ws.column_dimensions[cell.column_letter].width
                    needed = math.ceil(sum(2 if ord(c)>127 else 1 for c in cell.value) / max(1, width - 2))
                    height = max(height, min(110, 16 * needed + 6))
            ws.row_dimensions[row[0].row].height = height
        if ws.title in ("指标汇总", "求解说明"): ws.auto_filter.ref = None
    for i in (2, 6, 7, 8, 9, 10, 11):
        for col in ("B", "C", "D"): summary[f"{col}{i}"].number_format = "0"
    for ws, cols in [(rd, [11, 12]), (resources, [4]), (pf, [3, 7])]:
        for row in ws.iter_rows(min_row=2):
            for col in cols: row[col-1].number_format = "0"
    for c in ["B2", "B3", "B4", "B5"]:
        summary[c].font = Font(name="微软雅黑", size=12, bold=True, color="294B60")
        summary[c].fill = PatternFill("solid", fgColor="E6EFF4")
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True, calcMode="auto")
    engine = FormulaEngine(wb); cache = engine.evaluate_all()
    for i, val in enumerate(expected, 2):
        if not close(cache[("指标汇总", f"B{i}")], val): raise ValueError("Formula result mismatch")
    wb.save(output)
    install_caches(output, cache, wb.sheetnames)
    formulas = openpyxl.load_workbook(output, data_only=False)
    values = openpyxl.load_workbook(output, data_only=True)
    for (sheetname, address), expected_value in cache.items():
        actual = values[sheetname][address].value
        if formulas[sheetname][address].data_type != "f": raise ValueError("Formula lost during export")
        equal = close(actual, expected_value) if isinstance(expected_value, (int, float)) else (actual or "") == (expected_value or "")
        if not equal: raise ValueError(f"Cached formula mismatch {sheetname}!{address}")
    for wsname, columns in [("Q2_运输架次", [5,7,8]), ("Q2_逐箱交付", [4])]:
        for row in values[wsname].iter_rows(min_row=2):
            for col in columns:
                if not isinstance(row[col-1].value, (int,float)) or row[col-1].data_type != "n":
                    raise ValueError("Numeric result exported as text")
    assert values["Q2_运输架次"].max_row == len(routes)+1
    assert values["Q2_逐箱交付"].max_row == len(boxes)+1
    assert values["核验_资源事件"].max_row == len(routes)*3+1
    assert sha(template) == before_template
    checks = {"source_objective": expected,
              "workbook_objective": [values["指标汇总"].cell(i,2).value for i in range(2,6)],
              "routes":len(routes), "boxes":len(boxes), "archive_solutions":len(archive),
              "uavs":len(uavs), "batteries":len(batteries), "formula_count":len(cache),
              "all_formula_caches_verified":True, "typed_numeric_columns_verified":True,
              "one_final_solution_only":True, "original_workbook_unchanged":True,
              "source_independent_audit_matches":review_matches,
              "output":str(output), "sha256":sha(output),
              "formula_method":"Python restricted formula parser; Excel auto-recalculation on opening"}
    (folder / "导出核验.json").write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding="utf-8")
    if not args.no_preview:
        preview(wb, folder / "导出_工作簿预览.png", engine)
    print(json.dumps(checks,ensure_ascii=False,indent=2))


if __name__ == "__main__":
    main()
