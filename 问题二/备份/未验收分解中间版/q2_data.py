"""Q2 raw data, event-based routes and physical coefficients (no Q1 dependency)."""
from __future__ import annotations
import os
os.environ['MKL_THREADING_LAYER'] = 'SEQUENTIAL'
os.environ['OMP_NUM_THREADS'] = '1'
os.environ['OPENBLAS_NUM_THREADS'] = '1'
os.environ['MKL_NUM_THREADS'] = '1'
import hashlib
import itertools
import json
import math
from collections import Counter
from pathlib import Path
import numpy as np
import openpyxl
from scipy.io import loadmat

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent
G, J_PER_KWH = 9.81, 3600000.


def clean(value):
    if isinstance(value, dict): return {str(k): clean(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, np.ndarray)): return [clean(x) for x in value]
    if isinstance(value, (float, np.floating)):
        return float(value) if math.isfinite(value) else None
    if isinstance(value, (np.integer,)): return int(value)
    if isinstance(value, (np.bool_,)): return bool(value)
    if isinstance(value, Path): return str(value)
    return value


def save_json(path, value):
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(clean(value), ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')


def locate(root, name):
    # Support both the local D题/ layout and a data/ repository layout.
    candidates = []
    for base in (Path(root)/'D题', Path(root)/'数据', Path(root)):
        if base.is_dir():
            candidates = [p for p in base.rglob(name) if '备份' not in p.parts]
            if candidates: break
    if len(candidates) != 1: raise ValueError(f'Expected exactly one {name}: {candidates}')
    return candidates[0]


def block(rows, header, n):
    begin = next(i for i, row in enumerate(rows) if row[0] == header)
    ans = []
    for row in rows[begin+1:]:
        if row[0] is None or all(x is None for x in row[1:n]): break
        ans.append(row[:n])
    return ans


def pixel_cover(x0, y0, x1, y1, shape):
    """All closed pixels crossed by a straight line, including corner touches."""
    ts = [0., 1.]
    for p, q in ((x0,x1),(y0,y1)):
        if abs(q-p) > 1e-14:
            ts.extend((k-p)/(q-p) for k in range(math.floor(min(p,q))+1, math.ceil(max(p,q))))
    ts = np.unique(ts)
    samples = np.r_[ts, (ts[:-1]+ts[1:])/2]
    cells = set()
    for t in samples:
        x, y = x0+t*(x1-x0), y0+t*(y1-y0)
        for epsx, epsy in itertools.product((-1e-9,0.,1e-9), repeat=2):
            c, r = math.floor(x+epsx), math.floor(y+epsy)
            if 0 <= r < shape[0] and 0 <= c < shape[1]: cells.add((r,c))
    return sorted(cells)


def load_data(root=ROOT):
    paths = {k:locate(root,k) for k in ['调度中心与服务区.xlsx','运输无人机数据.xlsx',
                                      '物资需求与配送时限.xlsx','镇龙乡及周边30米DEM.mat']}
    raw = list(openpyxl.load_workbook(paths['运输无人机数据.xlsx'],data_only=True).active.values)
    inventory = {r[0]:r[1:] for r in block(raw,'共享电池库存',3)} if False else None
    sep = next(i for i,r in enumerate(raw) if r[0]=='共享电池库存')
    inventory = {r[0]:(int(r[1]),float(r[2])) for r in block(raw[sep+1:],'机型编号',3)}
    vehicles = {}
    for r in block(raw,'机型编号',18):
        g = r[0]
        vehicles[g] = dict(zip(['empty','capacity','volume','speed','r0','rf','battery','rho',
                               'prep','load','handoff0','handoff1','up','down','eta','descent'], map(float,r[2:18])))
        vehicles[g]['rho'] /= 100
        vehicles[g]['tfull'] = inventory[g][1]
        if vehicles[g]['descent'] != 0: raise ValueError('This declared energy closure requires zero descent supplement')
    uavs = {r[0]:r[1] for r in block(raw,'无人机编号',3)}
    if any(r[2]!='O01' for r in block(raw,'无人机编号',3)): raise ValueError('Nondepot initial UAV')
    bats = {f'{g}-BAT-{i:02d}':g for g,(n,_) in inventory.items() for i in range(1,n+1)}
    wb = openpyxl.load_workbook(paths['物资需求与配送时限.xlsx'],data_only=True)
    boxes = {}
    for r in list(wb['逐箱货箱清单'].values)[1:]:
        if r[0] is None: continue
        hard = []
        if r[5]=='是': hard.append(float(r[6]))
        if r[2]=='医疗物资': hard.append(float(r[7]))
        boxes[r[0]]=dict(box=r[0],zone=r[1],kind=r[2],mass=float(r[3]),volume=float(r[4]),
                         first=r[5]=='是',first_deadline=r[6],expected=float(r[7]),priority=float(r[8]),
                         hard=bool(hard),deadline=min(hard) if hard else None)
    raw = list(openpyxl.load_workbook(paths['调度中心与服务区.xlsx'],data_only=True).active.values)
    depot = block(raw,'调度中心编号',5)[0]
    zs = block(raw,'服务区编号',6)
    nodes = {r[0]:dict(lon=float(r[2]),lat=float(r[3]),ground=float(r[4])) for r in [depot]+zs}
    lon0,lat0=nodes['O01']['lon'], nodes['O01']['lat']; phi=math.radians(lat0)
    e2,a=6.6943799901413165e-3,6378137.
    nv=a/math.sqrt(1-e2*math.sin(phi)**2); mv=a*(1-e2)/(1-e2*math.sin(phi)**2)**1.5
    for z,v in nodes.items():
        v['x']=math.radians(v['lon']-lon0)*nv*math.cos(phi)
        v['y']=math.radians(v['lat']-lat0)*mv
        v['operation']=v['ground']+(0 if z=='O01' else 30)
    raster=loadmat(paths['镇龙乡及周边30米DEM.mat'])
    dem=raster['dem']; dx,_,left,_,dy,top=raster['transform'].ravel()
    geos={}
    for i,j in itertools.permutations(nodes,2):
        ni,nj=nodes[i],nodes[j]
        endpoints=[((n['lon']-left)/dx,(n['lat']-top)/dy) for n in (ni,nj)]
        if any(not(0<=x<dem.shape[1] and 0<=y<dem.shape[0]) for x,y in endpoints): raise ValueError('Out of DEM')
        cells=pixel_cover(*endpoints[0],*endpoints[1],dem.shape)
        vals=np.asarray([dem[r,c] for r,c in cells],float)
        if not np.isfinite(vals).all() or np.any(vals==raster['nodata'].item()): raise ValueError('DEM nodata')
        cruise=float(vals.max())+50; up=cruise-ni['operation']; down=cruise-nj['operation']
        if min(up,down)<0: raise ValueError('Cruise below operation; clarify geometry instead of silently clipping')
        geos[(i,j)]=dict(distance=math.hypot(nj['x']-ni['x'],nj['y']-ni['y']),up=up,down=down,
                         cruise=cruise,terrain=float(vals.max()),cells=len(cells))
    manifest=[dict(path=str(p),sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in paths.values()]
    return dict(boxes=boxes,vehicles=vehicles,uavs=uavs,batteries=bats,nodes=nodes,geos=geos,manifest=manifest)


def charge(v,soc):
    return v['tfull']*(.65*(.9-soc)/.9+.35) if soc<.9 else v['tfull']*.35*(1-soc)/.1


def route_key(r):
    return (r['vehicle'],tuple((e['zone'],tuple(sorted(e['boxes']))) for e in r['events']))


def make_route(data, vehicle, events, stats=None):
    def reject(reason):
        if stats is not None: stats[reason]+=1
        return None
    if stats is not None: stats['attempts']+=1
    v=data['vehicles'][vehicle]; bs=data['boxes']
    events=[dict(zone=e['zone'],boxes=sorted(e['boxes'])) for e in events]
    ids=[b for e in events for b in e['boxes']]
    if not ids or len(ids)!=len(set(ids)): return reject('empty_or_duplicate_box')
    if any(not e['boxes'] or any(b not in bs or bs[b]['zone']!=e['zone'] for b in e['boxes']) for e in events): return reject('destination')
    if any(a['zone']==b['zone'] for a,b in zip(events,events[1:])): return reject('adjacent_same_station')
    mass=sum(bs[b]['mass'] for b in ids); vol=sum(bs[b]['volume'] for b in ids)
    if mass>v['capacity']+1e-9: return reject('mass')
    if vol>v['volume']+1e-10: return reject('volume')
    q=mass; t=v['prep']+len(ids)*v['load']; prep=t; energy=0.; offsets={}; segs=[]; here='O01'
    for event in events+[dict(zone='O01',boxes=[])]:
        there=event['zone']; geo=data['geos'][(here,there)]
        length=v['r0']-(v['r0']-v['rf'])*(max(0.,q)/v['capacity'])**1.5
        eh=v['battery']*geo['distance']/length
        eu=(v['empty']+q)*G*geo['up']/(v['eta']*J_PER_KWH)
        dt=geo['up']/v['up']+geo['distance']/v['speed']+geo['down']/v['down']
        segs.append(dict(**geo,**{'from':here,'to':there},load=q,energy=eh+eu,horizontal=eh,climb=eu,time=dt,
                         departure=t,arrival=t+dt))
        energy+=eh+eu; t+=dt
        if event['boxes']:
            event['arrival_offset']=t
            t+=v['handoff0']+len(event['boxes'])*v['handoff1']
            event['delivery_offset']=t
            for b in event['boxes']: offsets[b]=t
            q-=sum(bs[b]['mass'] for b in event['boxes'])
        here=there
    if energy>v['battery']*(1-v['rho'])+1e-9: return reject('energy')
    latest=min((bs[b]['deadline']-offsets[b] for b in ids if bs[b]['hard']),default=math.inf)
    if latest < -1e-7: return reject('hard_deadline_at_zero')
    soc=1-energy/v['battery']
    r=dict(vehicle=vehicle,events=events,box_ids=sorted(ids),zones=[e['zone'] for e in events],nbox=len(ids),
           mass=mass,volume=vol,energy=energy,duration=t,prep_s=prep,offsets=offsets,
           latest_start=latest,soc=soc,charge_s=charge(v,soc),segments=segs)
    r['candidate_id']='K'+hashlib.sha256(repr(route_key(r)).encode()).hexdigest()[:16]
    if stats is not None: stats['accepted']+=1
    return r


def decode(data, records):
    routes=sorted(records,key=lambda r:(r['start'],r['candidate_id']))
    assigned=[]
    for i,r in enumerate(routes,1):
        r['route_id']=r.get('route_id',f'Q2-{i:04d}')
        r['takeoff']=r['start']+r['prep_s'];r['finish']=r['start']+r['duration'];r['recharge']=r['finish']+r['charge_s']
        for h,e in enumerate(r['events'],1):
            for b in e['boxes']:
                assigned.append(dict(**data['boxes'][b],route_id=r['route_id'],candidate_id=r['candidate_id'],event=h,
                                     delivery=r['start']+r['offsets'][b],uav=r['uav'],battery_id=r['battery_id']))
    soft=[b for b in assigned if not b['hard']]
    den=sum(b['priority'] for b in soft)
    obj=[len(routes),sum(r['energy'] for r in routes),max((r['finish'] for r in routes),default=0.),
         sum(b['priority']*max(0,b['delivery']-b['expected']) for b in soft)/(den or 1)]
    return dict(routes=routes,boxes=assigned,objective=obj)


def load_legacy(data,path):
    old=json.loads(Path(path).read_text(encoding='utf-8'));records=[]
    for r in old['routes']:
        boxes=[b['box'] for b in old['boxes'] if b['route_id']==r['route_id']]
        events=[dict(zone=z,boxes=[b for b in boxes if data['boxes'][b]['zone']==z]) for z in r['zones']]
        cand=make_route(data,r['vehicle'],events)
        if cand is None: raise ValueError('Invalid legacy route')
        records.append(dict(cand,start=r['start'],uav=r['uav'],battery_id=r['battery_id'],route_id=r['route_id']))
    return decode(data,records)
