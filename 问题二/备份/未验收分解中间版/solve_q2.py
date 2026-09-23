# -*- coding: utf-8 -*-
"""Question 2 solver: exact-box route library plus exact resource scheduling."""
from __future__ import annotations
import itertools, json, math, random, sys, time
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from copy import deepcopy
from pathlib import Path
import numpy as np
import pandas as pd
import psutil

OUT=Path(__file__).resolve().parent; ROOT=OUT.parent
sys.path.insert(0,str(OUT));sys.path.insert(0,str(ROOT/'问题一'))
from q2_data import load_data, make_route, save_json, route_key
from q2_model import RouteMaster, FixedSchedule, archive_add, preference_vertices
from q2_verify import verify_plan

SEEDS=list(range(202601,202611)); RHO=.20; MASTER_TIME=4.; SCHEDULE_TIME=4.

def _json(v):
    if isinstance(v,np.ndarray): return v.tolist()
    if isinstance(v,(np.integer,np.floating)): return v.item()
    if isinstance(v,Path): return str(v)
    if isinstance(v,dict): return {str(k):_json(x) for k,x in v.items()}
    if isinstance(v,(list,tuple)): return [_json(x) for x in v]
    return v

def save_csv(path,rows):
    Path(path).parent.mkdir(parents=True,exist_ok=True)
    pd.DataFrame(rows).to_csv(path,index=False,encoding='utf-8-sig')

def pattern_variants(p,identities,max_variants=5):
    groups=[]
    for key,count in p['counts'].items(): groups.append((list(identities[key]),int(count)))
    out=[]
    for v in range(max_variants):
        chosen=[]
        for j,(ids,n) in enumerate(groups):
            if not n: continue
            if n>=len(ids): take=ids[:]
            else:
                st=(v*(j+1)+j)%len(ids);take=[ids[(st+k)%len(ids)] for k in range(n)]
            chosen.extend(take)
        if len(chosen)==len(set(chosen)) and chosen not in out: out.append(chosen)
    return out

def generate_candidates(data):
    """Finite route library; repeated count patterns get distinct exact box IDs."""
    from q1_core import load_inputs,generate_patterns,active_patterns
    q1boxes,q1vehicles,q1geo,_,_=load_inputs(); universe,_,identities=generate_patterns(q1boxes,q1vehicles,q1geo)
    active,_=active_patterns(universe,rho=RHO); attempts=Counter(); allroutes={}; byz=defaultdict(list)
    # Explicit singleton routes guarantee every individual box can be represented.
    for b,vb in data['boxes'].items():
        for g in data['vehicles']:
            r=make_route(data,g,[{'zone':vb['zone'],'boxes':[b]}],attempts)
            if r: allroutes[route_key(r)]=r;byz[vb['zone']].append(r)
    bypat=defaultdict(list)
    for p in active: bypat[p['zone']].append(p)
    for z,ps in bypat.items():
        ps=sorted(ps,key=lambda p:(p['energy'],p['operation_s'],-p['nbox']))
        keep=[p for p in ps if p['nbox']==1]+ps[:40]+sorted(ps,key=lambda p:(-p['nbox'],p['energy']))[:30]
        seen=set()
        for p in keep:
            if p['pattern'] in seen: continue
            seen.add(p['pattern'])
            for ids in pattern_variants(p,identities,6):
                r=make_route(data,p['vehicle'],[{'zone':z,'boxes':ids}],attempts)
                if r: allroutes[route_key(r)]=r;byz[z].append(r)
    bestz={}
    for z,rs in byz.items():
        rs=list({route_key(r):r for r in rs}.values());rs.sort(key=lambda r:(r['energy']/max(1,r['nbox']),r['duration'],-r['nbox']))
        bestz[z]=rs[:32]
    zones=sorted(bestz);rng=random.Random(20260924)
    for z1,z2 in itertools.combinations(zones,2):
        local=[]
        for a in bestz[z1][:12]:
            for b in bestz[z2][:12]:
                if a['vehicle']!=b['vehicle']: continue
                for order in ((z1,z2),(z2,z1)):
                    src={z1:a,z2:b}
                    r=make_route(data,a['vehicle'],[{'zone':z,'boxes':src[z]['box_ids']} for z in order],attempts)
                    if r: local.append(r)
        for r in sorted({route_key(r):r for r in local}.values(),key=lambda r:(r['energy'],r['duration']))[:16]:allroutes[route_key(r)]=r
    triples=list(itertools.combinations(zones,3));rng.shuffle(triples)
    # Include deterministic singleton triples first; they are useful for
    # testing the repeated multi-point representation even when larger bundles
    # are energetically infeasible.
    singlez={z:[r for r in bestz[z] if r['nbox']==1] for z in zones}
    triple_iter=triples[:100]
    for tri in triple_iter:
        aa=singlez[tri[0]][:2] or bestz[tri[0]][:3]
        bb=singlez[tri[1]][:2] or bestz[tri[1]][:3]
        cc=singlez[tri[2]][:2] or bestz[tri[2]][:3]
        for a in aa:
            for b in bb:
                for c in cc:
                    if not (a['vehicle']==b['vehicle']==c['vehicle']): continue
                    src={tri[0]:a,tri[1]:b,tri[2]:c};order=list(tri);rng.shuffle(order)
                    r=make_route(data,a['vehicle'],[{'zone':z,'boxes':src[z]['box_ids']} for z in order],attempts)
                    if r:allroutes[route_key(r)]=r
    routes=list(allroutes.values())
    for i,r in enumerate(sorted(routes,key=lambda x:(len(x['zones']),x['energy'],x['candidate_id'])),1):r['route_id']=f'K{i:05d}'
    attempts['final']=len(routes);attempts['single']=sum(len(r['zones'])==1 for r in routes);attempts['double']=sum(len(r['zones'])==2 for r in routes);attempts['triple']=sum(len(r['zones'])==3 for r in routes)
    return routes,attempts,{'universe_patterns':len(universe),'active_patterns':len(active),'identities':len(identities)}

def gpu_check(routes):
    """Independent GPU batch sum of horizontal and climb components."""
    try:
        import torch
        if torch.cuda.device_count()<2:return {'available':torch.cuda.device_count(),'error':'fewer than two GPUs'},0.
        chunks=np.array_split(np.arange(len(routes)),2);out=[];diff=[]
        def job(dev,idx):
            vals=np.array([[sum(s['horizontal'] for s in r['segments']),sum(s['climb'] for s in r['segments'])] for r in [routes[i] for i in idx]],float)
            with torch.cuda.device(dev):
                x=torch.as_tensor(vals,device=dev,dtype=torch.float64);y=x.sum(1);torch.cuda.synchronize(dev);free,total=torch.cuda.mem_get_info(dev)
            return {'device':dev,'name':torch.cuda.get_device_name(dev),'routes':len(idx),'free_bytes':int(free),'total_bytes':int(total)},y.cpu().numpy(),idx
        with ThreadPoolExecutor(max_workers=2) as ex:
            fs=[ex.submit(job,i,chunks[i]) for i in range(2)]
            for f in as_completed(fs):
                rec,val,idx=f.result();out.append(rec);diff.extend(np.abs(val-np.array([routes[i]['energy'] for i in idx])))
        return {'devices':sorted(out,key=lambda x:x['device'])},float(max(diff) if diff else 0.)
    except Exception as e:return {'error':repr(e)},None

def evaluate(data,master,job,seed,logs):
    label=job['label'];selected,rm=master.solve(job['w'],capN=job.get('capN'),capE=job.get('capE'),seconds=MASTER_TIME,seed=seed,noise=job.get('noise',0.))
    rec={'job':label,'seed':seed,'route_master':rm}
    if selected is None:logs.append(rec|{'status':'no_route_incumbent'});return None
    try:p,sm=FixedSchedule(data,selected).solve_lex(seconds=SCHEDULE_TIME,seed=seed,capL=job.get('capL'))
    except Exception as e:logs.append(rec|{'status':'schedule_build_error','error':repr(e)});return None
    rec['schedule']=sm
    if p is None:logs.append(rec|{'status':'no_schedule_incumbent'});return None
    checks,errors,_=verify_plan(data,p);rec['verification']=checks;rec['errors']=errors;rec['status']='feasible' if checks['all'] else 'verification_failed';logs.append(rec)
    if checks['all']:
        p['source']={'job':label,'seed':seed,'route_master':rm,'schedule':sm};return p
    return None

def destroy_repair(data,master,current,rng,fraction):
    routes=current['routes'];nr=max(1,min(len(routes)-1,max(1,round(len(routes)*fraction)))) if len(routes)>1 else 1
    destroy=set(rng.sample(range(len(routes)),nr));fixed=[routes[i] for i in range(len(routes)) if i not in destroy];released=set(b for i,r in enumerate(routes) if i in destroy for b in r['box_ids'])
    pool=[r for r in master.routes if set(r['box_ids'])<=released]
    if not pool:return None
    local_data=deepcopy(data);local_data['boxes']={b:data['boxes'][b] for b in released}
    try:local,_=RouteMaster(local_data,pool).solve(w=(.4,.3,.2,.1),seconds=MASTER_TIME,seed=rng.randrange(10**9),noise=.001)
    except Exception:return None
    if local is None:return None
    try:
        p,_=FixedSchedule(data,fixed+local).solve_lex(seconds=SCHEDULE_TIME,seed=rng.randrange(10**9));checks,_,_=verify_plan(data,p) if p else ({'all':False},[],None)
        return p if p is not None and checks['all'] else None
    except Exception:return None

def plan_rows(plan):
    return [{'架次编号':r['route_id'],'无人机编号':r['uav'],'机型编号':r['vehicle'],'电池编号':r['battery_id'],'开始时刻（s）':round(r['start'],6),'实际起飞时刻（s）':round(r['takeoff'],6),'访问服务区顺序':'→'.join(['O01']+r['zones']+['O01']),'返回O01时刻（s）':round(r['finish'],6),'电池充满时刻（s）':round(r['recharge'],6),'架次能耗（kWh）':round(r['energy'],9),'返场SOC':round(r['soc'],9),'货箱数':r['nbox'],'货箱编号':'、'.join(r['box_ids'])} for r in plan['routes']]

def box_rows(plan):
    return [{'货箱编号':b['box'],'架次编号':b['route_id'],'服务区编号':b['zone'],'访问事件':b['event'],'交付完成时刻（s）':round(b['delivery'],6),'期望送达时间（s）':b['expected'],'适用硬截止时间（s）':b['deadline'] if b['hard'] else '','时限裕度或迟到量（s）':round((b['deadline']-b['delivery']) if b['hard'] else max(0,b['delivery']-b['expected']),6)} for b in sorted(plan['boxes'],key=lambda z:(z['delivery'],z['box']))]

def resource_rows(plan,data):
    cmax=plan['objective'][2];rows=[]
    for u,g in data['uavs'].items():
        rr=[r for r in plan['routes'] if r['uav']==u];rows.append({'资源类型':'实体运输无人机','资源编号':u,'机型编号':g,'执行次数':len(rr),'任务占用总时长s':sum(r['duration'] for r in rr),'最后返场s':max((r['finish'] for r in rr),default=0),'占用率':sum(r['duration'] for r in rr)/max(1,cmax)})
    for p,g in data['batteries'].items():
        rr=[r for r in plan['routes'] if r['battery_id']==p];rows.append({'资源类型':'共享电池','资源编号':p,'机型编号':g,'执行次数':len(rr),'任务占用总时长s':sum(r['duration'] for r in rr),'充电总时长s':sum(r['charge_s'] for r in rr),'最后充满s':max((r['recharge'] for r in rr),default=0)})
    return rows

def write_xlsx(path,rows,boxes,summary,resources):
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        from q2_data import locate
        wb=load_workbook(locate(ROOT,'结果提交模板.xlsx'))
        for name,data in [('Q2_运输架次',rows),('Q2_逐箱交付',boxes)]:
            if name not in wb.sheetnames:wb.create_sheet(name)
            ws=wb[name];ws.delete_rows(1,ws.max_row);keys=list(data[0]) if data else [];ws.append(keys)
            for rec in data:ws.append([rec.get(k,'') for k in keys])
            for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1F4E78');c.alignment=Alignment(horizontal='center')
        for name,obj in [('Q2_结果摘要',summary),('Q2_资源核验',resources)]:
            if name not in wb.sheetnames:wb.create_sheet(name)
            ws=wb[name];ws.delete_rows(1,ws.max_row)
            if isinstance(obj,dict):
                for k,v in obj.items():ws.append([k,json.dumps(_json(v),ensure_ascii=False) if isinstance(v,(dict,list)) else v])
            else:
                keys=list(obj[0]) if obj else [];ws.append(keys)
                for rec in obj:ws.append([rec.get(k,'') for k in keys])
        wb.save(path)
    except Exception as e:(OUT/'xlsx_warning.txt').write_text(repr(e),encoding='utf-8')

def main():
    t0=time.perf_counter();data=load_data();routes,gen_stats,gen_meta=generate_candidates(data);master=RouteMaster(data,routes)
    (OUT/'子问题一').mkdir(exist_ok=True);(OUT/'子问题二').mkdir(exist_ok=True)
    save_csv(OUT/'子问题一'/'候选架次库.csv',[{k:v for k,v in r.items() if k not in ('segments','events','offsets')} for r in routes])
    save_json(OUT/'子问题一'/'候选生成统计.json',{'stats':gen_stats,'meta':gen_meta,'candidate_count':len(routes)})
    gpu_info,gpu_diff=gpu_check(routes);logs=[];plans=[]
    anchors=[('架次数基准',(.80,.05,.10,.05)),('能耗基准',(.05,.80,.10,.05)),('完成时间基准',(.05,.10,.80,.05)),('及时性基准',(.05,.10,.10,.75)),('等权基准',(.25,.25,.25,.25))]
    jobs=[{'label':x,'w':tuple(w)} for x,w in anchors];rng=random.Random(20260924)
    for i in range(20):
        w=np.array([rng.random() for _ in range(4)]);jobs.append({'label':f'随机偏好{i+1:02d}','w':tuple(w/w.sum()),'noise':.002})
    def run(j,i):return evaluate(data,master,j,SEEDS[i%len(SEEDS)],logs)
    with ThreadPoolExecutor(max_workers=4) as ex:
        fs=[ex.submit(run,j,i) for i,j in enumerate(jobs)]
        for f in as_completed(fs):
            p=f.result()
            if p:plans.append(p)
    current=min(plans,key=lambda p:tuple(p['objective'])) if plans else None
    if current is not None:
        for it in range(15):
            trial=destroy_repair(data,master,current,rng,(.10,.20,.35)[it%3])
            if trial is None:continue
            a=np.asarray(trial['objective']);b=np.asarray(current['objective'])
            score=lambda x:float(.25*(x[0]/80+x[1]/100+x[2]/10000+x[3]/1000))
            if np.all(a<=b+np.array([0,1e-7,1e-5,1e-7])) or rng.random()<math.exp(min(0.,-(score(a)-score(b))/(.03*(1-it/55)))):
                current=trial;plans.append(trial)
    if plans:
        aa=np.array([p['objective'] for p in plans]);n0,n1=int(np.floor(aa[:,0].min())),int(np.ceil(aa[:,0].max()));e0,e1=float(aa[:,1].min()),float(aa[:,1].max());l0,l1=float(aa[:,3].min()),float(aa[:,3].max())
        nc=sorted(set([n0,n1,n1+2]));ec=sorted(set([max(0.,e0*.98),e0,e1*1.05]));lc=sorted(set([max(0.,l0),l1+max(1.,.2*l1)]));eps=[]
        for n,e,l in list(itertools.product(nc,ec,lc))[:12]:eps.append({'label':f'自适应ε-N{n}-E{e:.3f}-L{l:.3f}','w':(.20,.20,.40,.20),'capN':n,'capE':e,'capL':l})
        with ThreadPoolExecutor(max_workers=4) as ex:
            fs=[ex.submit(run,j,i) for i,j in enumerate(eps)]
            for f in as_completed(fs):
                p=f.result()
                if p:plans.append(p)
    archive=[]
    for p in plans:archive_add(archive,p)
    vertices=preference_vertices(delta=.30)
    with ThreadPoolExecutor(max_workers=4) as ex:
        fs=[ex.submit(run,{'label':f'偏好极点{h+1:02d}','w':tuple(w)},h) for h,w in enumerate(vertices)]
        for f in as_completed(fs):
            p=f.result()
            if p:plans.append(p);archive_add(archive,p)
    if not archive:raise RuntimeError('No independently verified feasible Question 2 plan')
    obj=np.array([p['objective'] for p in archive]);ideal=obj.min(0);scale=np.where(obj.max(0)-ideal>1e-8,obj.max(0)-ideal,1.)
    vals=[min(float(np.dot(w,(np.asarray(p['objective'])-ideal)/scale)) for p in archive) for w in vertices]
    for p in archive:p['archive_regret']=float(max(np.dot(w,(np.asarray(p['objective'])-ideal)/scale)-v for w,v in zip(vertices,vals)))
    final=min(archive,key=lambda p:(p['archive_regret'],p['objective'][2],p['objective'][3],p['objective'][0]));checks,errors,_=verify_plan(data,final)
    rows=plan_rows(final);br=box_rows(final);resources=resource_rows(final,data)
    save_csv(OUT/'子问题二'/'最终方案_运输架次.csv',rows);save_csv(OUT/'子问题二'/'最终方案_逐箱交付.csv',br);save_csv(OUT/'子问题二'/'资源使用汇总.csv',resources);save_csv(OUT/'子问题二'/'求解运行日志.csv',logs)
    save_csv(OUT/'子问题二'/'非支配方案.csv',[{'方案编号':f'PF{i+1:03d}','架次数':int(p['objective'][0]),'能耗kWh':p['objective'][1],'最晚返场s':p['objective'][2],'加权迟到s':p['objective'][3],'档案内最坏遗憾':p['archive_regret']} for i,p in enumerate(sorted(archive,key=lambda q:tuple(q['objective'])))])
    save_json(OUT/'子问题二'/'最终方案.json',final);save_json(OUT/'子问题二'/'独立核验.json',{'checks':checks,'errors':errors})
    events=[]
    for r in final['routes']:
        events += [{'资源类型':'运输无人机','资源编号':r['uav'],'架次编号':r['route_id'],'阶段':'任务占用','开始时刻s':r['start'],'结束时刻s':r['finish']},{'资源类型':'共享电池','资源编号':r['battery_id'],'架次编号':r['route_id'],'阶段':'任务占用','开始时刻s':r['start'],'结束时刻s':r['finish']},{'资源类型':'共享电池','资源编号':r['battery_id'],'架次编号':r['route_id'],'阶段':'充电','开始时刻s':r['finish'],'结束时刻s':r['recharge']}]
    save_csv(OUT/'子问题二'/'最终方案_资源事件.csv',events)
    summary={'状态':'可行且经原始数据独立核验','候选架次总数':len(routes),'单点候选数':gen_stats['single'],'两点候选数':gen_stats['double'],'三点候选数':gen_stats['triple'],'货箱数':len(data['boxes']),'服务区数':len({v['zone'] for v in data['boxes'].values()}),'硬时限货箱数':sum(v['hard'] for v in data['boxes'].values()),'非支配方案数':len(archive),'最终架次数':int(final['objective'][0]),'总能耗kWh':float(final['objective'][1]),'最晚返场s':float(final['objective'][2]),'普通物资加权迟到s':float(final['objective'][3]),'普通物资迟到箱数':sum((not b['hard']) and b['delivery']>b['expected']+1e-6 for b in final['boxes']),'最小硬时限裕度s':min((b['deadline']-b['delivery'] for b in final['boxes'] if b['hard']),default=None),'档案内最坏遗憾':float(final['archive_regret']),'遗憾认证范围':'统一有限候选库、已运行偏好极点和精确固定路线排程内；不是完整路线空间全局遗憾','候选生成统计':gen_stats,'GPU核验':gpu_info,'GPU最大能耗差kWh':gpu_diff,'核验':checks,'运行秒数':time.perf_counter()-t0}
    save_json(OUT/'求解摘要.json',summary);save_json(OUT/'计算资源记录.json',{'CPU逻辑线程':psutil.cpu_count(),'CPU物理核':psutil.cpu_count(logical=False),'内存总字节':psutil.virtual_memory().total,'MILP并行任务':4,'MILP单任务线程':1,'GPU':gpu_info,'GPU最大能耗差kWh':gpu_diff});write_xlsx(OUT/'问题二_结果提交.xlsx',rows,br,summary,resources)
    md=['# 问题二求解结果','',f'- 候选架次：{len(routes)}（单点{gen_stats["single"]}、两点{gen_stats["double"]}、三点{gen_stats["triple"]}）',f'- 最终代表：{int(final["objective"][0])}架次，{final["objective"][1]:.6f} kWh，最晚返场 {final["objective"][2]:.3f} s，加权迟到 {final["objective"][3]:.6f} s',f'- 普通物资迟到箱数：{summary["普通物资迟到箱数"]}，最小硬时限裕度：{summary["最小硬时限裕度s"]:.6f} s',f'- 档案内最坏遗憾：{final["archive_regret"]:.6f}（有限候选库范围）','', '## 独立核验',json.dumps(_json({'checks':checks,'errors':errors}),ensure_ascii=False,indent=2),'','## 最终架次表','', '|架次编号|无人机|机型|电池|开始时刻(s)|访问顺序|返场(s)|能耗(kWh)|','|---|---|---|---|---:|---|---:|---:|']
    md += ['|'+'|'.join(str(r[k]) for k in ('架次编号','无人机编号','机型编号','电池编号','开始时刻（s）','访问服务区顺序','返回O01时刻（s）','架次能耗（kWh）'))+'|' for r in rows];(OUT/'问题二求解结果.md').write_text('\n'.join(md),encoding='utf-8');print(json.dumps(_json(summary),ensure_ascii=False,indent=2))

if __name__=='__main__':main()
