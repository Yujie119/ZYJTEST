# -*- coding: utf-8 -*-
"""问题二：异构无人机多点多架次运输调度求解。

实现：候选架次生成、稀疏集合划分、实体无人机/共享电池排程、
多目标档案、独立核验和双 GPU 批量复核。所有输入物理规则继承问题一。
"""
from __future__ import annotations
import os
os.environ.setdefault("MKL_THREADING_LAYER", "SEQUENTIAL")
os.environ.setdefault("OMP_NUM_THREADS", "1")
os.environ.setdefault("OPENBLAS_NUM_THREADS", "1")
os.environ.setdefault("MKL_NUM_THREADS", "1")
import itertools, json, math, random, sys, time, warnings
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
import psutil
from scipy.io import loadmat
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import csc_matrix, lil_matrix

ROOT = Path(__file__).resolve().parents[1]
Q1 = ROOT / "问题一"
sys.path.insert(0, str(Q1))
from q1_core import (load_inputs, generate_patterns, active_patterns, supercover,
                     G, J_PER_KWH)

OUT = Path(__file__).resolve().parent
SUB1, SUB2 = OUT / "子问题一", OUT / "子问题二"
for d in (OUT, SUB1, SUB2): d.mkdir(parents=True, exist_ok=True)
RHO = .20
MIP_TIME = 20.0
SEEDS = list(range(202601, 202611))
ROUTES = []

def jr(x):
    if isinstance(x, np.ndarray): return x.tolist()
    if isinstance(x, (np.integer, np.floating)): return x.item()
    if isinstance(x, Path): return str(x)
    if isinstance(x, dict): return {str(k): jr(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)): return [jr(v) for v in x]
    return x
def save_json(path, obj): Path(path).write_text(json.dumps(jr(obj), ensure_ascii=False, indent=2), encoding="utf-8")
def save_csv(path, rows): pd.DataFrame(rows).to_csv(path, index=False, encoding="utf-8-sig")

def hi(c, integer, bounds, cons, tl=MIP_TIME):
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Unrecognized options detected")
        return milp(np.asarray(c, float), integrality=integer, bounds=bounds,
                    constraints=cons, options={"time_limit": tl, "mip_rel_gap": 5e-4,
                    "threads": 1, "mip_feasibility_tolerance": 1e-8})

@dataclass
class VInfo:
    id: str; capacity: float; volume: float; empty: float; speed: float
    r0: float; rf: float; battery: float; rho_floor: float
    prep: float; load: float; handoff0: float; handoff1: float
    up: float; down: float; eta: float; tfull: float

def read_inputs():
    boxes0, vehicles, _, nodes, manifest = load_inputs()
    fpath = next((ROOT / "D题/数据/无人机应急物资运输基础数据").glob("运输无人机数据.xlsx"))
    raw = pd.read_excel(fpath, header=None)
    end = raw.index[raw.iloc[:, 0].eq("逐架无人机清单")][0]
    vs = {}
    for row in raw.iloc[2:end].dropna(subset=[raw.columns[0]]).itertuples(index=False, name=None):
        if row[0] in vehicles:
            vs[row[0]] = VInfo(row[0], float(row[3]), float(row[4]), float(row[2]), float(row[5]),
                float(row[6]), float(row[7]), float(row[8]), float(row[9])/100, float(row[10]),
                float(row[11]), float(row[12]), float(row[13]), float(row[14]), float(row[15]),
                float(row[16]), {"A":1800., "B":2400., "C":3000.}[row[0]])
    bpath = next((ROOT / "D题/数据/无人机应急物资运输基础数据").glob("物资需求与配送时限.xlsx"))
    b = pd.read_excel(bpath, sheet_name="逐箱货箱清单").rename(columns={
        "货箱编号":"box", "服务区编号":"zone", "物资类型":"kind",
        "单箱质量（kg）":"mass", "单箱体积（m³）":"volume", "是否首批保障":"first",
        "首批截止时间（s）":"first_deadline", "期望送达时间（s）":"expected", "应急优先系数":"priority"})
    b["hard"] = b["first"].astype(str).eq("是") | b["kind"].eq("医疗物资")
    b["deadline"] = np.minimum(b["expected"].astype(float), b["first_deadline"].fillna(np.inf).astype(float))
    b["category"] = b["zone"].astype(str) + "|" + b["kind"].astype(str)
    return boxes0, b, vs, nodes, manifest

def all_geometry(nodes):
    path = next((ROOT / "D题/数据/镇龙乡地理空间数据").rglob("镇龙乡及周边30米DEM.mat"))
    z = loadmat(path); dem, tr = z["dem"], np.asarray(z["transform"]).ravel()
    dx, _, left, _, dy, top = tr
    cs = (nodes.lon.to_numpy()-left)/dx; rs = (nodes.lat.to_numpy()-top)/dy
    out = {}
    for i in range(len(nodes)):
        for j in range(len(nodes)):
            if i == j: continue
            cells = supercover(cs[i], rs[i], cs[j], rs[j], dem.shape)
            val = np.array([dem[r,c] for r,c in cells], float); cruise = float(val.max())+50
            out[(i,j)] = {"from":i, "to":j, "distance":float(np.hypot(nodes.x_m.iloc[i]-nodes.x_m.iloc[j], nodes.y_m.iloc[i]-nodes.y_m.iloc[j])),
                "cruise":cruise, "up":max(0.,cruise-float(nodes.operation_m.iloc[i])),
                "down":max(0.,cruise-float(nodes.operation_m.iloc[j])), "terrain":float(val.max()), "cells":len(cells)}
    return out

def seg(v, g, q):
    length = v.r0-(v.r0-v.rf)*max(0., q/v.capacity)**1.5
    horizontal = v.battery*g["distance"]/length
    climb = (v.empty+q)*G*g["up"]/(v.eta*J_PER_KWH)
    return {"energy":horizontal+climb, "horizontal":horizontal, "climb":climb,
            "time":g["up"]/v.up+g["distance"]/v.speed+g["down"]/v.down,
            "distance":g["distance"],"up":g["up"],"down":g["down"]}
def charge(v, soc):
    if soc < .9: return v.tfull*(.65*(.9-soc)/.9+.35)
    return v.tfull*.35*(1-soc)/.1

def cat_counts(p, zone, boxes):
    z = boxes[boxes.zone.eq(zone)]; groups = list(z.groupby(["mass","volume"], sort=True)); out = defaultdict(int)
    for key, n in p["counts"].items():
        gid = int(key.split(":")[-1][1:]); (mass, vol), _ = groups[gid]
        kinds = z[np.isclose(z.mass, mass) & np.isclose(z.volume, vol)].kind.unique()
        if len(kinds): out[f"{zone}|{kinds[0]}"] += int(n)
    return dict(out)

def make_route(parts, order, boxes, nodes, geos, vs):
    vname = next(iter(parts.values()))["vehicle"]
    if any(p["vehicle"] != vname for p in parts.values()): return None
    v = vs[vname]; counts = defaultdict(int); deliveries=[]; mass=vol=0.; nbox=0
    for z in order:
        cc=cat_counts(parts[z], z, boxes)
        if not cc: return None
        deliveries.append({"zone":z,"counts":cc})
        for c,n in cc.items():
            counts[c]+=n; zz=boxes[boxes.category.eq(c)].iloc[0]; mass += float(zz.mass)*n; vol += float(zz.volume)*n
        nbox += sum(cc.values())
    if mass>v.capacity+1e-8 or vol>v.volume+1e-8: return None
    ids=[0]+[int(nodes.index[nodes.zone.eq(z)][0]) for z in order]+[0]
    q=mass; energy=0.; t=v.prep+nbox*v.load; offsets={}; segments=[]
    for h,z in enumerate(order):
        s=seg(v,geos[(ids[h],ids[h+1])],q); energy+=s["energy"]; t+=s["time"]
        t += v.handoff0 + sum(deliveries[h]["counts"].values())*v.handoff1
        for c in deliveries[h]["counts"]: offsets[c]=t
        segments.append({"from":nodes.zone.iloc[ids[h]],"to":z,"load":q,**s})
        q -= sum(float(boxes[boxes.category.eq(c)].iloc[0].mass)*n for c,n in deliveries[h]["counts"].items())
    s=seg(v,geos[(ids[-2],0)],0.); energy+=s["energy"]; t+=s["time"]
    segments.append({"from":order[-1],"to":"O01","load":0.,**s})
    if energy>v.battery*(1-RHO)+1e-8: return None
    soc=1-energy/v.battery; latest=float("inf"); urgent=float("inf")
    for c in counts:
        cb=boxes[boxes.category.eq(c)]; x=cb.iloc[0]
        # Medical boxes are all hard by the statement.  A first-batch box is
        # hard individually, but the route may carry later boxes of the same
        # category; its first-batch slot is enforced after resource scheduling.
        if str(x.kind)=="医疗物资": latest=min(latest,float(x.expected)-offsets[c])
        first=cb[cb["first"].astype(str).eq("是")]
        if len(first): urgent=min(urgent,float(first.iloc[0].first_deadline)-offsets[c])
    return {"vehicle":vname,"zones":list(order),"counts":dict(counts),"deliveries":deliveries,
        "nbox":nbox,"mass":mass,"volume":vol,"energy":energy,
        "horizontal_energy":sum(s["horizontal"] for s in segments),"climb_energy":sum(s["climb"] for s in segments),
        "duration":t,"prep_s":v.prep+nbox*v.load,"offsets":offsets,"latest_start":latest,"urgent_start":urgent,
        "soc":soc,"charge_s":charge(v,soc),"segments":segments}

def generate_routes(boxes, nodes, geos, vs):
    q1boxes, q1vehicles, q1geo, _, _ = load_inputs(); universe, _, _ = generate_patterns(q1boxes, q1vehicles, q1geo)
    active, raw = active_patterns(universe, rho=RHO)
    byz=defaultdict(list)
    for p in active: byz[p["zone"]].append(p)
    rng=random.Random(20260923); selected={}
    for z,ps in byz.items():
        ps=sorted(ps,key=lambda p:(p["energy"],p["operation_s"],-p["nbox"]))
        selected[z]=ps[:22]+(rng.sample(ps[22:], min(8,max(0,len(ps)-22))) if len(ps)>22 else [])
    local=[]
    for z,ps in selected.items():
        for p in ps:
            r=make_route({z:p},[z],boxes,nodes,geos,vs)
            if r: local.append(r)
    zones=sorted(selected)
    for z1,z2 in itertools.combinations(zones,2):
        pairs=list(itertools.product(selected[z1],selected[z2]))
        if len(pairs)>120: pairs=rng.sample(pairs,120)
        for p1,p2 in pairs:
            for order in ((z1,z2),(z2,z1)):
                r=make_route({z1:p1,z2:p2},list(order),boxes,nodes,geos,vs)
                if r: local.append(r)
    for tri in itertools.combinations(zones,3):
        for _ in range(10):
            parts={z:rng.choice(selected[z]) for z in tri}; order=list(tri); rng.shuffle(order)
            r=make_route(parts,order,boxes,nodes,geos,vs)
            if r: local.append(r)
    best={}
    for r in local:
        key=(r["vehicle"],tuple(r["zones"]),tuple(sorted(r["counts"].items())))
        if key not in best or (r["energy"],r["duration"])<(best[key]["energy"],best[key]["duration"]): best[key]=r
    out=list(best.values())
    for i,r in enumerate(out,1): r["route_id"]=f"R{i:06d}"
    return out, active, raw

def gpu_check(routes, vs):
    try:
        import torch
        chunks=np.array_split(np.arange(len(routes)),2); records=[]; diffs=[]
        def job(device, idx):
            arr=np.array([[r["horizontal_energy"],r["climb_energy"]] for r in [routes[i] for i in idx]],float)
            with torch.cuda.device(device):
                x=torch.as_tensor(arr,device=device,dtype=torch.float64); y=x.sum(1); torch.cuda.synchronize(device)
                val=y.cpu().numpy(); free,total=torch.cuda.mem_get_info(device)
                return val,{"device":device,"name":torch.cuda.get_device_name(device),"routes":len(idx),"free_bytes":int(free),"total_bytes":int(total)}
        with ThreadPoolExecutor(max_workers=2) as pool:
            fs=[pool.submit(job,d,chunks[d]) for d in range(2)]
            for f in as_completed(fs):
                val,rec=f.result(); records.append(rec); diffs.extend(np.abs(val-np.array([routes[i]["energy"] for i in chunks[rec["device"]]])))
        return records,float(max(diffs) if diffs else 0.)
    except Exception as e: return [{"error":repr(e)}],None

def build_matrix(routes, boxes):
    cats=sorted(boxes.category.unique()); totals=np.array([(boxes.category==c).sum() for c in cats],float)
    routes=[r for r in routes if r["latest_start"]>=-1e-7]
    A=csc_matrix(np.array([[r["counts"].get(c,0) for r in routes] for c in cats],float))
    return routes,A,cats,totals
def choose(routes,A,cats,totals,w,capN=None,capE=None,tl=MIP_TIME):
    c=w[0]+w[1]*np.array([r["energy"] for r in routes])/60+w[2]*np.array([r["duration"] for r in routes])/10000
    latest=np.array([r["latest_start"] for r in routes], float)
    urgency=np.where(np.isfinite(latest), np.maximum(latest, 0.0), 0.0)
    type_uav={"A":4., "B":2., "C":2.}; type_bat={"A":6., "B":4., "C":4.}
    c += w[3]*urgency/10000
    c += .04*np.array([r["duration"]/type_uav[r["vehicle"]] + r["charge_s"]/type_bat[r["vehicle"]] for r in routes])/1000
    cons=[LinearConstraint(A,totals,totals)]; rows=[]; ub=[]
    if capN is not None: rows.append(np.ones(len(routes))); ub.append(capN)
    if capE is not None: rows.append(np.array([r["energy"] for r in routes])); ub.append(capE)
    if rows: cons.append(LinearConstraint(csc_matrix(np.array(rows)),-np.inf,np.array(ub)))
    res=hi(c,np.ones(len(routes)),Bounds(np.zeros(len(routes)),np.ones(len(routes))),cons,tl)
    if res.x is None or not res.success: return None,dict(status=int(res.status),message=str(res.message))
    x=np.rint(res.x).astype(int)
    return [routes[i] for i,v in enumerate(x) if v],dict(status=int(res.status),mip_gap=float(res.mip_gap or 0),nodes=int(res.mip_node_count or 0))

def schedule(routes, boxes, vs, seed=0, mode=0):
    uavs={"A":["U01","U02","U03","U04"],"B":["U05","U06"],"C":["U07","U08"]}
    bats={g:[f"{g}-BAT-{i:02d}" for i in range(1,n+1)] for g,n in [("A",6),("B",4),("C",4)]}
    ua={x:0. for xs in uavs.values() for x in xs}; ba={x:0. for xs in bats.values() for x in xs}
    ordered=list(routes); rng=random.Random(seed)
    keys=[lambda r:(min(r["latest_start"],r.get("urgent_start",float("inf"))),-r["nbox"],r["energy"]),lambda r:(min(r["latest_start"],r.get("urgent_start",float("inf"))),r["duration"],-r["energy"]),lambda r:(-r["energy"],min(r["latest_start"],r.get("urgent_start",float("inf"))),r["duration"]),lambda r:(-r["nbox"],min(r["latest_start"],r.get("urgent_start",float("inf"))),r["duration"])]
    if mode<4: ordered.sort(key=keys[mode])
    else: rng.shuffle(ordered); ordered.sort(key=lambda r:r["latest_start"])
    out=[]
    for r in ordered:
        g=r["vehicle"]; u=min(uavs[g],key=ua.get); p=min(bats[g],key=ba.get); start=max(ua[u],ba[p])
        if start>r["latest_start"]+1e-6: return {"_fail":"medical_deadline_start","route":r["route_id"],"start":start,"latest":r["latest_start"]}
        f=start+r["duration"]; out.append({**r,"uav":u,"battery_id":p,"start":start,"takeoff":start+r["prep_s"],"finish":f,"recharge":f+r["charge_s"]}); ua[u]=f; ba[p]=f+r["charge_s"]
    slots=defaultdict(list)
    for r in out:
        for d in r["deliveries"]:
            for c,n in d["counts"].items(): slots[c].extend([(r,r["start"]+r["offsets"][c])] * n)
    assigned=[]
    for cat,g in boxes.groupby("category",sort=False):
        bxs=g.sort_values(["hard","deadline","box"],ascending=[False,True,True]).to_dict("records"); ss=sorted(slots[cat],key=lambda x:x[1])
        if len(bxs)!=len(ss): return {"_fail":"category_slot_mismatch","category":cat,"boxes":len(bxs),"slots":len(ss)}
        assigned.extend([{**b,"route_id":r["route_id"],"delivery":t,"uav":r["uav"],"battery_id":r["battery_id"]} for b,(r,t) in zip(bxs,ss)])
    late=[b for b in assigned if b["hard"] and b["delivery"]>b["deadline"]+1e-6]
    if late: return {"_fail":"first_or_medical_deadline","count":len(late),"worst":max(float(b["delivery"]-b["deadline"]) for b in late)}
    den=sum(float(b["priority"]) for b in assigned if not b["hard"])
    L=sum(float(b["priority"])*max(0,float(b["delivery"]-b["expected"])) for b in assigned if not b["hard"])/(den or 1)
    obj=np.array([len(out),sum(r["energy"] for r in out),max(r["finish"] for r in out),L])
    out.sort(key=lambda r:r["start"]); return {"routes":out,"boxes":assigned,"objective":obj}

def schedule_milp(routes, boxes, vs, time_limit=35.0):
    """Exact resource assignment/scheduling for a fixed route set."""
    n=len(routes)
    if n==0: return None
    uavs={"A":["U01","U02","U03","U04"],"B":["U05","U06"],"C":["U07","U08"]}
    bats={"A":[f"A-BAT-{i:02d}" for i in range(1,7)],"B":[f"B-BAT-{i:02d}" for i in range(1,5)],"C":[f"C-BAT-{i:02d}" for i in range(1,5)]}
    H=max(20000., max((r["latest_start"] if np.isfinite(r["latest_start"]) else 0.)+r["duration"]+r["charge_s"] for r in routes)+5000.)
    lb=[]; ub=[]; integ=[]; obj=[]; names=[]; idx_s=list(range(n))
    for i,r in enumerate(routes): lb.append(0.); ub.append(float(r["latest_start"] if np.isfinite(r["latest_start"]) else H)); integ.append(0); obj.append(0.); names.append(("s",i))
    cidx=n; lb.append(0.); ub.append(H); integ.append(0); obj.append(1.); names.append(("cmax",0))
    uidx={}; bidx={};
    for i,r in enumerate(routes):
        for d in uavs[r["vehicle"]]: uidx[(i,d)]=len(lb); lb.append(0); ub.append(1); integ.append(1); obj.append(0); names.append(("u",i,d))
        for p in bats[r["vehicle"]]: bidx[(i,p)]=len(lb); lb.append(0); ub.append(1); integ.append(1); obj.append(0); names.append(("b",i,p))
    yidx={}
    for i,j in itertools.combinations(range(n),2):
        if routes[i]["vehicle"]!=routes[j]["vehicle"]: continue
        g=routes[i]["vehicle"]
        for d in uavs[g]:
            for a,b in ((i,j),(j,i)):
                yidx[("u",a,b,d)]=len(lb); lb.append(0); ub.append(1); integ.append(1); obj.append(0); names.append(("yu",a,b,d))
        for p in bats[g]:
            for a,b in ((i,j),(j,i)):
                yidx[("b",a,b,p)]=len(lb); lb.append(0); ub.append(1); integ.append(1); obj.append(0); names.append(("yb",a,b,p))
    # Select which route delivers the single first-batch box of each category.
    # This avoids treating every later box in that category as a hard task.
    first_cats={}
    for cat,g in boxes.groupby("category",sort=False):
        f=g[g["first"].astype(str).eq("是")]
        if len(f): first_cats[cat]=float(f.iloc[0].first_deadline)
    yfidx={}
    for cat,deadline in first_cats.items():
        for i,r in enumerate(routes):
            if r["counts"].get(cat,0)>0:
                yfidx[(cat,i)]=len(lb); lb.append(0); ub.append(1); integ.append(1); obj.append(0); names.append(("yf",cat,i))
    rows=[]; lo=[]; hi_b=[]
    def add(co, low=-np.inf, high=np.inf): rows.append(co); lo.append(low); hi_b.append(high)
    for i,r in enumerate(routes):
        add({uidx[(i,d)]:1 for d in uavs[r["vehicle"]]},1,1)
        add({bidx[(i,p)]:1 for p in bats[r["vehicle"]]},1,1)
        add({cidx:1, idx_s[i]:-1},r["duration"],np.inf)
    M=H+max(r["duration"]+r["charge_s"] for r in routes)+1000.
    for cat,deadline in first_cats.items():
        cand=[i for (c,i) in yfidx if c==cat]
        if not cand: return None
        add({yfidx[(cat,i)]:1 for i in cand},1,1)
        for i in cand:
            off=routes[i]["offsets"][cat]
            add({idx_s[i]:1,yfidx[(cat,i)]:M},-np.inf,deadline+M-off)
    for i,j in itertools.combinations(range(n),2):
        if routes[i]["vehicle"]!=routes[j]["vehicle"]: continue
        g=routes[i]["vehicle"]
        for d in uavs[g]:
            y1=yidx[("u",i,j,d)]; y2=yidx[("u",j,i,d)]; ui=uidx[(i,d)]; uj=uidx[(j,d)]
            add({y1:1,y2:1,ui:-1,uj:-1},-1, np.inf); add({y1:1,y2:1,ui:-1},-np.inf,0); add({y1:1,y2:1,uj:-1},-np.inf,0)
            add({idx_s[j]:1,idx_s[i]:-1,y1:-M},routes[i]["duration"]-M,np.inf)
            add({idx_s[i]:1,idx_s[j]:-1,y2:-M},routes[j]["duration"]-M,np.inf)
        for p in bats[g]:
            y1=yidx[("b",i,j,p)]; y2=yidx[("b",j,i,p)]; bi=bidx[(i,p)]; bj=bidx[(j,p)]
            add({y1:1,y2:1,bi:-1,bj:-1},-1,np.inf); add({y1:1,y2:1,bi:-1},-np.inf,0); add({y1:1,y2:1,bj:-1},-np.inf,0)
            add({idx_s[j]:1,idx_s[i]:-1,y1:-M},routes[i]["duration"]+routes[i]["charge_s"]-M,np.inf)
            add({idx_s[i]:1,idx_s[j]:-1,y2:-M},routes[j]["duration"]+routes[j]["charge_s"]-M,np.inf)
    mat=lil_matrix((len(rows),len(lb)),dtype=float)
    for ri,co in enumerate(rows):
        for ci,val in co.items(): mat[ri,ci]=val
    res=hi(np.array(obj),np.array(integ),Bounds(np.array(lb),np.array(ub)),LinearConstraint(mat.tocsc(),np.array(lo),np.array(hi_b)),time_limit)
    if res.x is None or not res.success: return None
    out=[]
    for i,r in enumerate(routes):
        u=max(uavs[r["vehicle"]],key=lambda d:res.x[uidx[(i,d)]])
        p=max(bats[r["vehicle"]],key=lambda z:res.x[bidx[(i,z)]])
        start=float(res.x[idx_s[i]]); f=start+r["duration"]
        out.append({**r,"uav":u,"battery_id":p,"start":start,"takeoff":start+r["prep_s"],"finish":f,"recharge":f+r["charge_s"]})
    slots=defaultdict(list)
    for r in out:
        for d in r["deliveries"]:
            for c,nc in d["counts"].items(): slots[c].extend([(r,r["start"]+r["offsets"][c])]*nc)
    assigned=[]
    for cat,g in boxes.groupby("category",sort=False):
        bxs=g.sort_values(["hard","deadline","box"],ascending=[False,True,True]).to_dict("records")
        ss=sorted(slots[cat],key=lambda x:x[1])
        if len(bxs)!=len(ss): return None
        assigned.extend([{**b,"route_id":r["route_id"],"delivery":t,"uav":r["uav"],"battery_id":r["battery_id"]} for b,(r,t) in zip(bxs,ss)])
    late=[b for b in assigned if b["hard"] and b["delivery"]>b["deadline"]+1e-6]
    if late: return None
    den=sum(float(b["priority"]) for b in assigned if not b["hard"])
    L=sum(float(b["priority"])*max(0,float(b["delivery"]-b["expected"])) for b in assigned if not b["hard"])/(den or 1.)
    out.sort(key=lambda r:r["start"])
    return {"routes":out,"boxes":assigned,"objective":np.array([n,sum(r["energy"] for r in out),max(r["finish"] for r in out),L])}

def verify(plan, boxes, vs):
    checks={}; ids=[x["box"] for x in plan["boxes"]]; checks["coverage"]=len(ids)==len(set(ids))==len(boxes) and set(ids)==set(boxes.box)
    checks["hard_deadlines"]=all((not b["hard"]) or b["delivery"]<=b["deadline"]+1e-6 for b in plan["boxes"])
    checks["route_load"]=all(r["mass"]<=vs[r["vehicle"]].capacity+1e-7 and r["volume"]<=vs[r["vehicle"]].volume+1e-8 for r in plan["routes"])
    recalc=[]
    for r in plan["routes"]:
        v=vs[r["vehicle"]]; e=0.
        for s in r["segments"]:
            length=v.r0-(v.r0-v.rf)*max(0.,float(s["load"])/v.capacity)**1.5
            e += v.battery*float(s["distance"])/length + (v.empty+float(s["load"]))*G*float(s["up"])/(v.eta*J_PER_KWH)
        recalc.append(e)
    checks["route_energy"]=all(e<=vs[r["vehicle"]].battery*(1-RHO)+1e-7 and abs(e-r["energy"])<=1e-6 for e,r in zip(recalc,plan["routes"]))
    bu=defaultdict(list); bp=defaultdict(list)
    for r in plan["routes"]: bu[r["uav"]].append((r["start"],r["finish"])); bp[r["battery_id"]].append((r["start"],r["recharge"]))
    checks["uav_nonoverlap"]=all(all(a[1]<=b[0]+1e-6 for a,b in zip(sorted(v),sorted(v)[1:])) for v in bu.values())
    checks["battery_nonoverlap"]=all(all(a[1]<=b[0]+1e-6 for a,b in zip(sorted(v),sorted(v)[1:])) for v in bp.values())
    checks["soc"]=all(r["soc"]>=vs[r["vehicle"]].rho_floor-1e-8 for r in plan["routes"]); checks["all"]=all(checks.values()); return checks

def dom(a,b): return np.all(a<=b+np.array([0,1e-7,1e-5,1e-7])) and np.any(a<b-np.array([0,1e-7,1e-5,1e-7]))
def add_archive(arc,p):
    f=p["objective"]
    if any(np.allclose(f,q["objective"],atol=[0,1e-6,1e-4,1e-7]) or dom(q["objective"],f) for q in arc): return
    arc[:]=[q for q in arc if not dom(f,q["objective"])]; arc.append(p)

def output_workbook(path, rows, dro, summary, resources):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb=load_workbook(ROOT/"D题/结果提交模板.xlsx")
    for name,data in [("Q2_运输架次",rows),("Q2_逐箱交付",dro)]:
        ws=wb[name]; ws.delete_rows(2,max(0,ws.max_row-1))
        for i,rec in enumerate(data,2):
            for j,c in enumerate(ws[1],1): c2=ws.cell(i,j); c2.value=rec.get(c.value)
        for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E78"); c.alignment=Alignment(horizontal="center")
    for name,data in [("Q2_结果摘要",summary),("Q2_资源核验",resources)]:
        ws=wb[name] if name in wb.sheetnames else wb.create_sheet(name); ws.delete_rows(1,ws.max_row)
        for i,(k,v) in enumerate(data.items(),1): ws.cell(i,1).value=k; ws.cell(i,2).value=str(v)
    wb.save(path)

def main():
    t0=time.perf_counter(); boxes0,boxes,vs,nodes,manifest=read_inputs(); geos=all_geometry(nodes)
    save_csv(SUB1/"全节点DEM航段几何.csv",[{"from":nodes.zone.iloc[i],"to":nodes.zone.iloc[j],**g} for (i,j),g in geos.items()])
    routes,active,raw=generate_routes(boxes,nodes,geos,vs)
    save_csv(SUB1/"候选架次库.csv",[{k:(json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v) for k,v in r.items() if k not in ("segments","deliveries")} for r in routes])
    routes,A,cats,totals=build_matrix(routes,boxes); gpu_info,gpu_diff=gpu_check(routes,vs)
    jobs=[("架次数",[.75,.1,.1,.05]),("能耗",[.1,.75,.1,.05]),("完成时间",[.1,.1,.75,.05]),("及时性",[.1,.1,.1,.7]),("单点可行性",[.2,.1,.6,.1])]
    rng=random.Random(20260923)
    for i in range(20):
        w=np.array([rng.random() for _ in range(4)]); jobs.append((f"随机{i+1:02d}",(w/w.sum()).tolist()))
    results=[]; logs=[]
    def run(job,idx):
        label,w=job; sel,meta=choose(routes,A,cats,totals,w,tl=MIP_TIME); best=None
        if label.startswith("单点"):
            sr=[r for r in routes if len(r["zones"])==1]
            sr,sa,sc,st=build_matrix(sr,boxes); sel,meta=choose(sr,sa,sc,st,w,tl=MIP_TIME)
        if sel:
            p=schedule_milp(sel,boxes,vs,time_limit=35.0)
            if p is not None and (best is None or tuple(p["objective"])<tuple(best["objective"])): best=p
            elif p is None: meta["failure_mode"]="resource_or_first_deadline_milp"
        meta=dict(meta, selected_routes=len(sel) if sel else 0,
                  selected_hard_routes=sum(np.isfinite(r["latest_start"]) for r in sel) if sel else 0)
        return label,meta,best
    def run_batch(batch, offset):
        with ThreadPoolExecutor(max_workers=min(8,max(1,(os.cpu_count() or 2)//2))) as pool:
            fs=[pool.submit(run,j,offset+i) for i,j in enumerate(batch)]
            for f in as_completed(fs):
                label,meta,p=f.result(); logs.append({"job":label,"meta":meta,"status":"feasible" if p else "no_feasible_schedule"})
                if p: results.append(p)
    # Anchor solves first; their values define the adaptive epsilon region.
    run_batch(jobs[:5],0)
    if results:
        aa=np.array([p["objective"] for p in results]); nlo,nhi=int(aa[:,0].min()),int(aa[:,0].max()); elo,ehi=float(aa[:,1].min()),float(aa[:,1].max())
        eps_jobs=[]
        for nc in range(nlo,nhi+1):
            for frac in (0.,.25,.5,.75,1.):
                eps_jobs.append((f"自适应ε-N{nc}-E{frac:.2f}",[.65,.15,.15,.05]))
                # Store the cap in a side table consumed by this batch closure.
        # attach caps by label; the objective remains the same but the feasible region changes.
        cap_map={f"自适应ε-N{nc}-E{frac:.2f}":(nc,elo+frac*(ehi-elo)) for nc in range(nlo,nhi+1) for frac in (0.,.25,.5,.75,1.)}
        old_run=run
        def run(job,idx):
            label,w=job
            if label.startswith("自适应ε-"):
                capN,capE=cap_map[label]; sel,meta=choose(routes,A,cats,totals,w,capN=capN,capE=capE,tl=MIP_TIME); best=None
                if sel:
                    p=schedule_milp(sel,boxes,vs,time_limit=35.0); best=p if p is not None else None
                meta=dict(meta, selected_routes=len(sel) if sel else 0, epsilon_N=capN, epsilon_E=capE)
                return label,meta,best
            return old_run(job,idx)
        run_batch(eps_jobs,5)
    run_batch(jobs[5:],5+len(locals().get('eps_jobs',[])))
    if not results:
        save_csv(SUB2/"求解运行日志_失败.csv",logs)
        raise RuntimeError("没有得到经独立核验的完整可行排程")
    archive=[]
    for p in results:
        if verify(p,boxes,vs)["all"]: add_archive(archive,p)
    if not archive: raise RuntimeError("可行排程未通过独立核验")
    obj=np.array([p["objective"] for p in archive]); ideal=obj.min(0); scale=np.where(obj.max(0)>ideal+1e-8,obj.max(0)-ideal,1.)
    weights=[np.array([.25]*4)]
    for i in range(4):
        for s in (-.15,.15):
            w=np.full(4,.25); w[i]+=s; w-=s/3; weights.append(w)
    bestv=np.array([min(w@((p["objective"]-ideal)/scale) for p in archive) for w in weights])
    for p in archive: p["regret"]=float(max(w@((p["objective"]-ideal)/scale)-v for w,v in zip(weights,bestv)))
    final=min(archive,key=lambda p:(p["regret"],p["objective"][2],p["objective"][3])); checks=verify(final,boxes,vs)
    rows=[{"架次编号":r["route_id"],"无人机编号":r["uav"],"机型编号":r["vehicle"],"电池编号":r["battery_id"],"开始时刻（s）":round(max(0.,r["start"]),3),"访问服务区顺序":"→".join(["O01"]+r["zones"]+["O01"]),"返回O01时刻（s）":round(r["finish"],3),"架次能耗（kWh）":round(r["energy"],6)} for r in final["routes"]]
    dro=[{"货箱编号":b["box"],"架次编号":b["route_id"],"服务区编号":b["zone"],"交付完成时刻（s）":round(float(b["delivery"]),3)} for b in sorted(final["boxes"],key=lambda x:(x["delivery"],x["box"]))]
    save_csv(SUB2/"最终方案_运输架次.csv",rows); save_csv(SUB2/"最终方案_逐箱交付.csv",dro); save_csv(SUB2/"求解运行日志.csv",logs)
    events=[]
    for r in final["routes"]:
        events.append({"资源类型":"运输无人机","资源编号":r["uav"],"架次编号":r["route_id"],"阶段":"任务占用","开始时刻s":r["start"],"结束时刻s":r["finish"]})
        events.append({"资源类型":"共享电池","资源编号":r["battery_id"],"架次编号":r["route_id"],"阶段":"任务占用","开始时刻s":r["start"],"结束时刻s":r["finish"]})
        events.append({"资源类型":"共享电池","资源编号":r["battery_id"],"架次编号":r["route_id"],"阶段":"充电","开始时刻s":r["finish"],"结束时刻s":r["recharge"]})
    save_csv(SUB2/"最终方案_资源事件.csv",events)
    save_csv(SUB2/"非支配方案.csv",[{"方案编号":f"PF{i+1:03d}","架次数":int(p["objective"][0]),"能耗kWh":p["objective"][1],"最晚返场s":p["objective"][2],"加权迟到s":p["objective"][3],"最坏遗憾":p["regret"]} for i,p in enumerate(sorted(archive,key=lambda x:tuple(x["objective"])))])
    save_json(SUB2/"最终方案.json",final); save_json(SUB2/"独立核验.json",checks)
    resources={"CPU逻辑线程":psutil.cpu_count(),"CPU物理核":psutil.cpu_count(logical=False),"内存总字节":psutil.virtual_memory().total,"MILP并行任务":min(8,max(1,(os.cpu_count() or 2)//2)),"MILP单任务线程":1,"GPU核验":gpu_info,"GPU最大能耗差kWh":gpu_diff}
    save_json(OUT/"计算资源记录.json",resources)
    summary={"状态":"可行且经独立核验","候选架次总数":len(routes),"单点候选数":sum(len(r["zones"])==1 for r in routes),"两点候选数":sum(len(r["zones"])==2 for r in routes),"三点候选数":sum(len(r["zones"])==3 for r in routes),"货箱数":len(boxes),"服务区数":boxes.zone.nunique(),"硬时限货箱数":int(boxes.hard.sum()),"非支配方案数":len(archive),"最终架次数":int(final["objective"][0]),"总能耗kWh":float(final["objective"][1]),"最晚返场s":float(final["objective"][2]),"普通物资加权迟到s":float(final["objective"][3]),"档案内最坏遗憾":float(final["regret"]),"核验":checks,"认证范围":"给定候选架次库和计算预算下的近似非支配方案；遗憾为档案内评价","运行秒数":time.perf_counter()-t0}
    save_json(OUT/"求解摘要.json",summary); output_workbook(OUT/"问题二_结果提交.xlsx",rows,dro,summary,resources)
    md=["# 问题二求解结果","",f"- 候选架次：{len(routes)}（单点{sum(len(r['zones'])==1 for r in routes)}、两点{sum(len(r['zones'])==2 for r in routes)}、三点{sum(len(r['zones'])==3 for r in routes)}）",f"- 最终代表：{int(final['objective'][0])} 架次，{final['objective'][1]:.6f} kWh，最晚返场 {final['objective'][2]:.3f} s，加权迟到 {final['objective'][3]:.6f} s",f"- 档案内最坏遗憾：{final['regret']:.6f}","", "## 硬约束核验",json.dumps(checks,ensure_ascii=False,indent=2),"","## 最终架次表","","|架次编号|无人机|机型|电池|开始时刻(s)|访问顺序|返场(s)|能耗(kWh)|","|---|---|---|---|---:|---|---:|---:|"]
    for r in rows: md.append("|"+"|".join(str(r[k]) for k in ["架次编号","无人机编号","机型编号","电池编号","开始时刻（s）","访问服务区顺序","返回O01时刻（s）","架次能耗（kWh）"])+"|")
    (OUT/"问题二求解结果.md").write_text("\n".join(md),encoding="utf-8")
    print(json.dumps(summary,ensure_ascii=False,indent=2))
if __name__=="__main__": main()
