"""Fresh audit: raw XLSX + GeoTIFF, no imports of q2 modules or cached geometry."""
from pathlib import Path
from collections import Counter, defaultdict
import json, math, hashlib
import numpy as np
import openpyxl, tifffile
from scipy.io import loadmat

ROOT=Path(__file__).resolve().parents[3]
OUT=Path(__file__).resolve().parent
RAW=ROOT/'D题/数据/无人机应急物资运输基础数据'

def dump(name,obj):
    (OUT/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')

def raw_data():
    ws=openpyxl.load_workbook(RAW/'运输无人机数据.xlsx',data_only=True).active
    vehicles={ws.cell(i,1).value:list(next(ws.iter_rows(min_row=i,max_row=i,values_only=True)))[2:18] for i in (3,4,5)}
    uavs={ws.cell(i,1).value:ws.cell(i,2).value for i in range(9,17)}
    inventory={ws.cell(i,1).value:(ws.cell(i,2).value,ws.cell(i,3).value) for i in range(20,23)}
    bs=openpyxl.load_workbook(RAW/'物资需求与配送时限.xlsx',data_only=True)['逐箱货箱清单']
    boxes={r[0]:r for r in list(bs.values)[1:] if r[0]}
    ns=openpyxl.load_workbook(RAW/'调度中心与服务区.xlsx',data_only=True).active
    nodes={ns.cell(i,1).value:[ns.cell(i,j).value for j in (3,4,5)] for i in [3,*range(7,22)]}
    dempath=next((ROOT/'D题').rglob('镇龙乡及周边30米DEM.tif'))
    with tifffile.TiffFile(dempath) as tf:
        dem=tf.asarray(); meta=tf.geotiff_metadata
    assert int(meta['GeographicTypeGeoKey'])==4326
    dx,sy,_=meta['ModelPixelScale'];_,_,_,tie_x,tie_y,_=meta['ModelTiepoint']
    # GeoTIFF RasterPixelIsPoint stores centers. Convert to closed-pixel edges.
    assert int(meta['GTRasterTypeGeoKey'])==2
    left=tie_x-dx/2;top=tie_y+sy/2;dy=-sy
    mat=loadmat(next((ROOT/'D题').rglob('镇龙乡及周边30米DEM.mat')))
    provenance={'geotiff_sha256':hashlib.sha256(dempath.read_bytes()).hexdigest(),
        'mat_tif_pixels_equal':bool(np.array_equal(dem,mat['dem'])),
        'max_transform_difference':float(np.max(np.abs(np.array([dx,0,left,0,dy,top])-mat['transform'].ravel()))),
        'geotiff_raster_type':'PixelIsPoint, half-pixel correction applied'}
    lon0,lat0,_=nodes['O01']; phi=math.radians(lat0); e2=6.6943799901413165e-3;a=6378137.
    nv=a/math.sqrt(1-e2*math.sin(phi)**2); mv=a*(1-e2)/(1-e2*math.sin(phi)**2)**1.5
    xy={k:(math.radians(v[0]-lon0)*nv*math.cos(phi),math.radians(v[1]-lat0)*mv) for k,v in nodes.items()}
    geos={}
    for i in nodes:
        for j in nodes:
            if i==j:continue
            xi,yi=(nodes[i][0]-left)/dx,(nodes[i][1]-top)/dy
            xj,yj=(nodes[j][0]-left)/dx,(nodes[j][1]-top)/dy
            # Independent rectangle/slab intersection, not production sampling.
            cc,rr=np.meshgrid(np.arange(max(0,math.floor(min(xi,xj))-1),min(dem.shape[1],math.ceil(max(xi,xj))+1)),
                np.arange(max(0,math.floor(min(yi,yj))-1),min(dem.shape[0],math.ceil(max(yi,yj))+1)))
            lo=np.zeros(cc.shape);hi=np.ones(cc.shape)
            for start,end,c in [(xi,xj,cc),(yi,yj,rr)]:
                if abs(end-start)<1e-14:hi=np.where((c<=start)&(start<=c+1),hi,-1)
                else:
                    t0=(c-start)/(end-start); t1=(c+1-start)/(end-start)
                    lo=np.maximum(lo,np.minimum(t0,t1));hi=np.minimum(hi,np.maximum(t0,t1))
            hit=lo<=hi+1e-12;terrain=float(dem[rr[hit],cc[hit]].max());cruise=terrain+50
            zi=nodes[i][2]+(0 if i=='O01' else 30);zj=nodes[j][2]+(0 if j=='O01' else 30)
            geos[i,j]=(math.dist(xy[i],xy[j]),cruise-zi,cruise-zj,terrain,int(hit.sum()))
    return vehicles,uavs,inventory,boxes,nodes,geos,provenance

def verify(plan,raw):
    vehicles,uavs,inventory,boxes,nodes,geos,_=raw
    assertions=[];records=[];delivered={};uv=defaultdict(list);bat=defaultdict(list)
    def check(label,ok,actual=None,expected=None):
        assertions.append({'check':label,'passed':bool(ok),'actual':actual,'expected':expected})
    for r in plan['routes']:
        g=r['vehicle']; empty,cap,volume,speed,r0,rf,B,rho,prep,load,hand0,hand1,up,down,eta,desc=vehicles[g]
        ids=[b for e in r['events'] for b in e['boxes']]; totalmass=sum(boxes[b][3] for b in ids);totalvol=sum(boxes[b][4] for b in ids)
        check(r['route_id']+' capacity',totalmass<=cap+1e-9 and totalvol<=volume+1e-10)
        check(r['route_id']+' identifiers',uavs.get(r['uav'])==g and r['battery_id'].startswith(g+'-BAT-') and 1<=int(r['battery_id'].split('-')[-1])<=inventory[g][0])
        time=prep+load*len(ids);energy=0.;mass=totalmass;here='O01'
        for e in [*r['events'],{'zone':'O01','boxes':[]}]:
            there=e['zone'];distance,climb,descent,terrain,cells=geos[here,there]
            assert min(climb,descent)>=0
            time+=climb/up+distance/speed+descent/down
            energy+=B*distance/(r0-(r0-rf)*(mass/cap)**1.5)+(empty+mass)*9.81*climb/eta/3600000
            if e['boxes']:
                time+=hand0+hand1*len(e['boxes'])
                for b in e['boxes']:
                    row=boxes[b];check(b+' unique',b not in delivered);delivered[b]=r['start']+time
                    check(b+' destination',row[1]==there)
                    limits=([row[6]] if row[5]=='是' else [])+([row[7]] if row[2]=='医疗物资' else [])
                    if limits:check(b+' hard time',delivered[b]<=min(limits)+2e-5,delivered[b],min(limits))
                mass-=sum(boxes[b][3] for b in e['boxes'])
            here=there
        soc=1-energy/B; full=inventory[g][1]
        charge=(max(0,.9-soc)/.9*.65+min(.1,1-soc)/.1*.35)*full
        finish=r['start']+time;ready=finish+charge
        check(r['route_id']+' energy',abs(energy-r['energy'])<1e-7,energy,r['energy'])
        check(r['route_id']+' finish',abs(finish-r['finish'])<2e-5,finish,r['finish'])
        check(r['route_id']+' recharge',abs(ready-r['recharge'])<2e-5,ready,r['recharge'])
        check(r['route_id']+' soc',soc>=rho/100-1e-9,soc,rho/100)
        check(r['route_id']+' start',r['start']>=-2e-5)
        uv[r['uav']].append((r['start'],finish,r['route_id']));bat[r['battery_id']].append((r['start'],ready,r['route_id']))
        records.append({'route_id':r['route_id'],'candidate_id':r['candidate_id'],'energy':energy,'finish':finish,'soc':soc,'start':r['start']})
    check('exact coverage',set(delivered)==set(boxes))
    for kind,iv in [('uav',uv),('battery',bat)]:
        for name,rr in iv.items():
            rr=sorted(rr)
            for a,b in zip(rr,rr[1:]):check(kind+' '+name+' '+a[2]+' before '+b[2],a[1]<=b[0]+2e-5,a[1],b[0])
    soft=[b for b,v in boxes.items() if v[5]!='是' and v[2]!='医疗物资']
    obj=[len(records),sum(r['energy'] for r in records),max(r['finish'] for r in records),sum(boxes[b][8]*max(0,delivered[b]-boxes[b][7]) for b in soft)/sum(boxes[b][8] for b in soft)]
    for j,(a,b) in enumerate(zip(obj,plan['objective'])):check('objective '+str(j),abs(a-b)<2e-5,a,b)
    return {'all':all(x['passed'] for x in assertions),'checks_count':len(assertions),'objective':obj,'assertions':assertions,'delivery':delivered,'routes':records}

if __name__=='__main__':
    raw=raw_data();base=json.loads((ROOT/'问题二/子问题二/最终方案.json').read_text(encoding='utf-8'))
    witness=json.loads((OUT/'fixed_resources_timing_witness.json').read_text(encoding='utf-8'))['plan']
    a,b=verify(base,raw),verify(witness,raw)
    dump('independent_raw_verification.json',{'provenance':raw[-1],'original':a,'witness':b})
    changes=[];orig={r['candidate_id']:r for r in base['routes']}
    for r in witness['routes']:
        old=orig[r['candidate_id']]
        assert r['uav']==old['uav'] and r['battery_id']==old['battery_id']
        if abs(r['start']-old['start'])>2e-5:changes.append({'original_route_id':old['route_id'],'candidate_id':r['candidate_id'],'uav':r['uav'],'battery_id':r['battery_id'],'old_start':old['start'],'new_start':r['start'],'shift':r['start']-old['start']})
    dump('witness_timing_changes.json',changes)
    print(json.dumps({'provenance':raw[-1],'original':[a['all'],a['checks_count'],a['objective']],'witness':[b['all'],b['checks_count'],b['objective']],'changed_routes':len(changes)},ensure_ascii=False))
