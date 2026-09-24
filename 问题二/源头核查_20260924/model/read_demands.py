from openpyxl import load_workbook
p=r'D:\研究生\比赛\数学建模\2026\D题\数据\无人机应急物资运输基础数据\物资需求与配送时限.xlsx'
wb=load_workbook(p,data_only=True,read_only=True)
for ws in wb.worksheets:
 print('\n===',ws.title)
 for i,row in enumerate(ws.iter_rows(values_only=True),1):
  print(i,'\t'.join('' if x is None else str(x) for x in row))
