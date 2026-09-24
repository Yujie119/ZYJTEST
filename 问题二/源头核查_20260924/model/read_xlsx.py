from openpyxl import load_workbook
from pathlib import Path
for p in Path(r'D:\研究生\比赛\数学建模\2026\D题\数据\无人机应急物资运输基础数据').glob('*.xlsx'):
 print('\n===',p.name)
 wb=load_workbook(p,data_only=False,read_only=True)
 print('sheets',wb.sheetnames)
 for ws in wb.worksheets:
  print('--',ws.title,ws.max_row,ws.max_column)
  for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,25),values_only=True): print('\t'.join('' if x is None else str(x) for x in row))
