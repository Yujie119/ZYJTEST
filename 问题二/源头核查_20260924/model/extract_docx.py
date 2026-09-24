from docx import Document
p=r"D:\研究生\比赛\数学建模\2026\D题\山区洪涝灾害下无人机运输与通信协同优化.docx"
d=Document(p)
for i,para in enumerate(d.paragraphs,1):
 t=para.text.strip()
 if any(k in t for k in ['问题二','医疗','首批','期望','时限','附录 2','能源','充电','准备','交接','返回']): print(f'{i}: {t}')
for ti,t in enumerate(d.tables,1):
 for ri,row in enumerate(t.rows,1):
  s=' | '.join(c.text.replace('\n',';') for c in row.cells)
  if any(k in s for k in ['医疗','首批','时限','能源','充电','无人机','交付']): print(f'T{ti}R{ri}: {s}')
