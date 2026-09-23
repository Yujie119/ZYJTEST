"""Rebuild all exports from one verified Q2 plan; read back the official sheets."""
from __future__ import annotations
import os
os.environ['MKL_THREADING_LAYER']='SEQUENTIAL'
os.environ['OMP_NUM_THREADS']='1'
os.environ['OPENBLAS_NUM_THREADS']='1'
os.environ['MKL_NUM_THREADS']='1'
import hashlib,json,platform,sys
from collections import Counter
from datetime import datetime,timezone
from pathlib import Path
import numpy as np
import pandas as pd
import scipy
from scipy.optimize._highspy import _core as hc
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Font,PatternFill
from q2_data import OUT,ROOT,load_data,load_baseline,locate,save_json
from q2_verify import verify_plan,TIME_TOL

SUB=OUT/'子问题二'
def read(name):return json.loads((OUT/name).read_text(encoding='utf-8'))
def csv(name,rows):
    table=pd.DataFrame(rows);table.to_csv(SUB/name,index=False,encoding='utf-8-sig',float_format='%.12g');return table
def digest(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def mdtable(rows):
    if isinstance(rows,pd.DataFrame):rows=rows.to_dict('records')
    if not rows:return ''
    keys=list(rows[0]);lines=['|'+'|'.join(keys)+'|','|'+'|'.join('---' for _ in keys)+'|']
    for r in rows:
        def show(v):
            if v is None:return '—'
            if isinstance(v,(float,np.floating)):return f'{v:.6f}'.rstrip('0').rstrip('.')
            return str(v).replace('|','／')
        lines.append('|'+'|'.join(show(r.get(k)) for k in keys)+'|')
    return '\n'.join(lines)

def main():
    data=load_data();plan=read('子问题二/最终方案.json');checks,errors,recalc=verify_plan(data,plan)
    if not checks['all']:raise AssertionError(errors)
    run=read('联合运行摘要.json');config=read('联合求解配置.json');frozen=read('冻结评价配置.json')
    archive=read('子问题二/联合非支配档案.json');preferences=read('子问题二/偏好极点基准.json')
    logs=read('子问题二/完整联合运行日志.json');pool=read('子问题一/联合候选库.json')
    library_hash=hashlib.sha256('\n'.join(sorted(r['candidate_id'] for r in pool)).encode()).hexdigest()
    assert library_hash==run['library_sha256']==frozen['library_sha256']
    assert set(r['candidate_id'] for r in plan['routes'])<=set(r['candidate_id'] for r in pool)
    for p in archive:
        c,e,_=verify_plan(data,p)
        if not c['all']:raise AssertionError(e)
        vals=np.array([np.dot(v['coefficients'],p['objective']) for v in preferences])
        interval=[max(0.,max(vals-np.array([v['upper'] for v in preferences]))),max(0.,max(vals-np.array([v['lower'] for v in preferences])))]
        assert np.max(abs(np.array(interval)-[p['regret_lower'],p['regret_upper']]))<1e-7
    assert abs(plan['regret_upper']-min(p['regret_upper'] for p in archive))<1e-7
    routes=plan['routes'];boxes=sorted(plan['boxes'],key=lambda b:b['box']);N,E,C,L=plan['objective']
    wb=load_workbook(locate(ROOT,'结果提交模板.xlsx'))
    fh=[x.value for x in wb['Q2_运输架次'][1]];bh=[x.value for x in wb['Q2_逐箱交付'][1]]
    fv=[[r['route_id'],r['uav'],r['vehicle'],r['battery_id'],r['start'],'→'.join(['O01']+r['zones']+['O01']),r['finish'],r['energy']] for r in routes]
    bv=[[b['box'],b['route_id'],b['zone'],b['delivery']] for b in boxes]
    flight_rows=[dict(zip(fh,r)) for r in fv];box_rows=[dict(zip(bh,r)) for r in bv]
    csv('最终方案_运输架次.csv',flight_rows);csv('最终方案_逐箱交付.csv',box_rows)
    detail=[];events=[];visits=[];legs=[];waits=[];ua={u:0. for u in data['uavs']};ba={p:0. for p in data['batteries']}
    for r in sorted(routes,key=lambda r:(r['start'],r['route_id'])):
        detail.append(dict(架次编号=r['route_id'],候选编号=r['candidate_id'],货箱编号列表=';'.join(r['box_ids']),总质量kg=r['mass'],总体积m3=r['volume'],实际起飞s=r['takeoff'],返航SOC=r['soc'],充满时刻s=r['recharge']))
        for kind,rid,phase,start,end in [('无人机',r['uav'],'任务',r['start'],r['finish']),('电池',r['battery_id'],'任务',r['start'],r['finish']),('电池',r['battery_id'],'充电',r['finish'],r['recharge'])]:
            events.append(dict(资源类型=kind,资源编号=rid,架次编号=r['route_id'],阶段=phase,开始s=start,结束s=end))
        for h,e in enumerate(r['events'],1):visits.append(dict(架次编号=r['route_id'],事件序号=h,服务区编号=e['zone'],货箱编号列表=';'.join(e['boxes']),到达s=r['start']+e['arrival_offset'],交接完成s=r['start']+e['delivery_offset']))
        for h,s in enumerate(r['segments'],1):legs.append(dict(架次编号=r['route_id'],航段=h,起点=s['from'],终点=s['to'],剩余载荷kg=s['load'],水平距离m=s['distance'],爬升m=s['up'],下降m=s['down'],巡航海拔m=s['cruise'],离开s=r['start']+s['departure'],到达s=r['start']+s['arrival'],水平能耗kWh=s['horizontal'],爬升能耗kWh=s['climb'],航段能耗kWh=s['energy']))
        ad,ap=ua[r['uav']],ba[r['battery_id']];earliest=max(ad,ap)
        reason='初始可用' if earliest<TIME_TOL else '电池就绪' if ap>ad+TIME_TOL else '无人机返场' if ad>ap+TIME_TOL else '两资源同时就绪'
        waits.append(dict(架次编号=r['route_id'],无人机先前返场s=ad,电池先前充满s=ap,给定资源顺序下最早开始s=earliest,实际开始s=r['start'],额外等待s=max(0.,r['start']-earliest),电池相对无人机延后s=max(0.,ap-ad),限制条件=reason))
        ua[r['uav']]=r['finish'];ba[r['battery_id']]=r['recharge']
    csv('最终方案_架次详情.csv',detail);event_table=csv('最终方案_资源事件.csv',events)
    csv('最终方案_逐站事件.csv',visits);csv('最终方案_航段明细.csv',legs);csv('资源等待条件.csv',waits)
    allboxes=[dict(货箱编号=b['box'],架次编号=b['route_id'],服务区编号=b['zone'],物资类型=b['kind'],交付完成时刻s=b['delivery'],期望时刻s=b['expected'],硬时限=b['hard'],适用截止时刻s=b['deadline'],首批保障=b['first'],优先系数=b['priority'],时限裕度s=(b['deadline']-b['delivery']) if b['hard'] else None,迟到量s=max(0.,b['delivery']-b['expected']) if not b['hard'] else None) for b in boxes]
    csv('最终方案_逐箱核验.csv',allboxes)
    hard=[b for b in allboxes if b['硬时限']];soft=[b for b in allboxes if not b['硬时限']]
    csv('硬时限核验.csv',hard);csv('普通物资及时性.csv',soft)
    resource=[]
    for kind,raw,key in [('实体运输无人机',data['uavs'],'uav'),('共享电池',data['batteries'],'battery_id')]:
        for rid,g in raw.items():
            rr=[r for r in routes if r[key]==rid];work=sum(r['duration'] for r in rr)
            resource.append(dict(资源类型=kind,资源编号=rid,机型编号=g,使用次数=len(rr),任务总时长s=work,充电总时长s=sum(r['charge_s'] for r in rr) if key=='battery_id' else None,最后返场s=max((r['finish'] for r in rr),default=0.),最后充满s=max((r['recharge'] for r in rr),default=0.) if key=='battery_id' else None,无人机占用率=work/C if key=='uav' else None))
    resource_table=csv('资源使用汇总.csv',resource)
    archive.sort(key=lambda p:tuple(p['objective']));ar=[]
    for i,p in enumerate(archive,1):ar.append(dict(方案编号=f'PF{i:03d}',来源=p.get('source','未记录'),架次数=p['objective'][0],能耗kWh=p['objective'][1],最晚返场s=p['objective'][2],加权迟到s=p['objective'][3],遗憾下界=p['regret_lower'],遗憾上界=p['regret_upper'],选为提交方案=all(abs(x-y)<1e-7 for x,y in zip(p['objective'],plan['objective']))))
    csv('非支配方案.csv',ar)
    named=[('原发布方案',load_baseline(data)),('固定原资源修正',read('子问题二/固定原资源反例复现.json')['plan'])]
    for j,name in enumerate(['架次侧重基准','能耗侧重基准','完成时间侧重基准','及时性侧重基准']):named.append((name,min(archive,key=lambda p:(p['objective'][j],tuple(p['objective'])))))
    named.append(('最终代表方案',plan));comparison=[]
    for name,p in named:comparison.append(dict(方案=name,架次数N=p['objective'][0],能耗kWh=p['objective'][1],最晚返场s=p['objective'][2],加权迟到s=p['objective'][3],普通物资迟到箱数=sum(not b['hard'] and b['delivery']>b['expected']+TIME_TOL for b in p['boxes']),遗憾下界=p.get('regret_lower'),遗憾上界=p.get('regret_upper')))
    compare_table=csv('四目标基准比较.csv',comparison)
    sr=[]
    for log in logs:
        for index,m in enumerate(log['stages'],1):sr.append(dict(任务=log['job']['label'],阶段=index,阶段目标=m.get('stage',str(m['weights'])),状态=m['status'],已有可行解=m['incumbent'],最优性认证=m['optimal'],下界=m['lower_bound'],上界=m['upper_bound'],相对界差=m['gap'],实际求解秒=m['seconds'],变量数=m['variables'],约束数=m['constraints'],模型时域s=m['horizon'],阈值=json.dumps(m['epsilon'],ensure_ascii=False),先前阶段全部认证=m.get('prior_stages_certified'),种子=m['seed']))
    csv('求解运行日志.csv',sr)
    csv('偏好极点与有效界.csv',[dict(极点=i+1,架次权重=v['weight'][0],能耗权重=v['weight'][1],完成时间权重=v['weight'][2],迟到权重=v['weight'][3],加权目标下界=v['lower'],加权目标上界=v['upper'],MILP状态=v['status'],LP状态=v['lp_status']) for i,v in enumerate(preferences)])
    pd.DataFrame([dict(候选编号=r['candidate_id'],机型编号=r['vehicle'],访问顺序='→'.join(r['zones']),货箱编号列表=';'.join(r['box_ids']),货箱数=r['nbox'],质量kg=r['mass'],体积m3=r['volume'],能耗kWh=r['energy'],持续时间s=r['duration'],返航SOC=r['soc'],最晚开始s=r['latest_start']) for r in pool]).to_csv(OUT/'子问题一/候选架次库.csv',index=False,encoding='utf-8-sig')
    pd.DataFrame([dict(起点=i,终点=j,**g) for (i,j),g in data['geos'].items()]).to_csv(OUT/'子问题一/全节点DEM航段几何.csv',index=False,encoding='utf-8-sig')
    for sheet,headers,values in [('Q2_运输架次',fh,fv),('Q2_逐箱交付',bh,bv)]:
        ws=wb[sheet]
        if ws.max_row>1:ws.delete_rows(2,ws.max_row-1)
        for row in values:ws.append(row)
        assert [c.value for c in ws[1]]==headers
    extra={'核验_架次详情':pd.DataFrame(detail),'核验_逐箱':pd.DataFrame(allboxes),'核验_资源事件':event_table,'核验_资源汇总':resource_table,'目标权衡':compare_table,'求解说明':pd.DataFrame([dict(项目='计算范围',内容='冻结K0候选库；未证明完整原题全局最优'),dict(项目='最终方案选择',内容='当前核验档案内最坏遗憾上界最小'),dict(项目='遗憾区间',内容=str([plan['regret_lower'],plan['regret_upper']])),dict(项目='候选库SHA256',内容=library_hash)])}
    for name,table in extra.items():
        ws=wb.create_sheet(name);ws.append(list(table.columns))
        for row in table.itertuples(index=False,name=None):ws.append([None if isinstance(v,float) and np.isnan(v) else v for v in row])
    for name in ['Q2_运输架次','Q2_逐箱交付']+list(extra):
        ws=wb[name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(name='微软雅黑',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='294B60');c.alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[1].height=32
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(52,max(16,len(str(col[0].value or ''))*1.7))
            for cell in col[1:]:
                cell.font=Font(name='微软雅黑',size=10)
                if isinstance(cell.value,float):cell.number_format='0.000000'
    target=OUT/'问题二_结果提交.xlsx';temporary=OUT/'问题二_结果提交.tmp.xlsx';wb.save(temporary);os.replace(temporary,target)
    reopened=load_workbook(target,data_only=True)
    for name,headers,expected in [('Q2_运输架次',fh,fv),('Q2_逐箱交付',bh,bv)]:
        actual=list(reopened[name].values);assert list(actual[0])==headers and len(actual)==len(expected)+1
        for a,b in zip(actual[1:],expected):
            for x,y in zip(a,b):assert abs(x-y)<1e-8 if isinstance(y,(int,float)) else x==y
    initial=read('子问题二/构造初始化档案.json');direct=read('子问题二/直接联合档案.json')
    nnew=sum(not any(np.all(abs(np.array(p['objective'])-q['objective'])<=np.array([0,1e-7,1e-5,1e-7])) for q in initial) for p in direct)
    summary=dict(objective=plan['objective'],regret_interval=[plan['regret_lower'],plan['regret_upper']],global_optimality_certified=False,chosen_source=plan.get('source'),archive_count=len(archive),box_count=len(boxes),hard_box_count=len(hard),soft_box_count=len(soft),late_soft_boxes=sum(b['迟到量s']>TIME_TOL for b in soft),minimum_hard_margin_s=min(b['时限裕度s'] for b in hard),minimum_return_soc=min(r['soc'] for r in routes),route_type_counts=Counter(r['vehicle'] for r in routes),route_visit_counts=Counter(len(r['events']) for r in routes),last_return_route=max(routes,key=lambda r:r['finish'])['route_id'],last_return_uav=max(routes,key=lambda r:r['finish'])['uav'],inventory_uavs=Counter(data['uavs'].values()),inventory_batteries=Counter(data['batteries'].values()),direct_new_objective_vectors_vs_initializers=nnew,independent_checks=checks,template_readback_passed=True,library_count=len(pool),library_sha256=library_hash,normalization=frozen,run_times=run,final_plan_sha256=digest(SUB/'最终方案.json'),scope='Physical rules audited under stated assumptions. Bounds apply only to K0; 10-seed comparative experiments remain planned.')
    save_json(OUT/'求解摘要.json',summary);save_json(SUB/'独立核验.json',dict(checks=checks,errors=errors,recomputed=recalc,final_plan_sha256=summary['final_plan_sha256']))
    gpu=read('子问题二/双GPU物理复算.json')
    save_json(OUT/'计算资源记录.json',dict(python=sys.version,executable=sys.executable,numpy=np.__version__,scipy=scipy.__version__,highs=hc._Highs().version(),platform=platform.platform(),logical_cpu_count=os.cpu_count(),parallel_jobs=config['parallel_jobs'],threads_per_solver=1,gpu=gpu,interpretation='CPU parallel MILPs; both GPUs independently recompute physics; no GPU MIP or speedup claim.'))
    text=['# 问题二求解结果（审计修正版）',f'最终代表方案为 {N} 架次，总能耗 {E:.6f} kWh，最晚返场 {C:.6f} s，普通物资加权平均迟到 {L:.6f} s。80箱全部交付，31箱硬时限全部满足。',f'方案从当前 {len(archive)} 个非支配方案中按最坏遗憾上界最小准则选择；K0范围内遗憾区间为 [{plan["regret_lower"]:.6f}, {plan["regret_upper"]:.6f}]。未认证完整路线问题的全局最优，也未认证库内精确最小最大遗憾。','## 四目标取舍',mdtable(compare_table),'每一行是一套完整可行方案；四个侧重基准指当前档案中该指标最小的方案，重复行表示基准重合。遗憾采用同一尺度和同一K0基准上下界。','## 最终运输架次',mdtable(flight_rows),'完整80箱交付、逐站事件、航段、资源与充电记录见子问题二目录；两个Q2提交页严格使用原模板表头，仅提交这套代表方案。','## 可行性与资源',f'硬时限最小裕度 {summary["minimum_hard_margin_s"]:.6f} s，最小返航SOC {100*summary["minimum_return_soc"]:.6f}%；普通物资迟到 {summary["late_soft_boxes"]} 箱。',mdtable(resource_table),'资源表列出全部8架无人机和14块电池，未使用设备保留为0次。电池最后充满可晚于运输完成。等待表描述给定顺序下的就绪限制，不将全部空闲归因于充电。','## 实际计算与结论范围',f'完整单点候选枚举 {run["single_point_catalog"]} 条；实际联合MILP使用K0的 {len(pool)} 条，包含三种机型与四站路线。直接阶段相对初始化新增 {nnew} 个目标向量；来源保存在JSON及非支配方案表。',f'直接阶段 {run["direct_seconds"]:.3f} s，ALNS对照 {run["alns_seconds"]:.3f} s，整次求解 {run["total_seconds"]:.3f} s。时限参数不等于墙钟耗时，预处理和退出开销计入日志；两方法预算不同，不作等预算优劣比较。',f'两块GPU从箱重、机型参数及已审计DEM几何重算 {gpu["route_count"]} 条候选的能耗、时长、SOC及充电；MILP由CPU执行。','10种子对照、候选扩展及资源/偏好敏感性仍为计划。物理能耗继承论文公开的航程校准与势能爬升假设。限时、候选范围和偏好界差不自动等于代码错误。',f'最终方案SHA256：{summary["final_plan_sha256"]}。',f'候选库SHA256：{library_hash}。']
    (OUT/'问题二求解结果.md').write_text('\n\n'.join(text)+'\n',encoding='utf-8')
    save_json(OUT/'正文结果数据.json',dict(summary=summary,comparison=comparison,flight_rows=flight_rows,resources=resource))
    inputs=[p for p in OUT.glob('*.py')]+[OUT/'联合求解配置.json',OUT/'冻结评价配置.json',OUT/'初始化/审计基线方案.json',ROOT/'正文/正文.md']
    outputs=[target,OUT/'求解摘要.json',OUT/'问题二求解结果.md']+list(SUB.glob('*.csv'))+list(SUB.glob('*.json'))
    outputs += [p for p in [OUT/'审计修正与复核报告.md',OUT/'复核_独立审阅_修正版.md'] if p.exists()]
    save_json(OUT/'复现指纹.json',dict(generated_at=datetime.now(timezone.utc).isoformat(),inputs={str(p.relative_to(ROOT)):digest(p) for p in inputs},outputs={str(p.relative_to(ROOT)):digest(p) for p in outputs},source_data=data['manifest']))
    print(json.dumps({k:summary[k] for k in ['objective','regret_interval','minimum_hard_margin_s','minimum_return_soc','template_readback_passed','chosen_source','direct_new_objective_vectors_vs_initializers']},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
