"""Targeted event-geometry tests and cross-file submission verification."""
from q3_common import *
from q3_communication import blocked_events, merged, sphere_events
from verify_q3 import Verifier
from openpyxl import load_workbook
import hashlib

def main():
    checks=[]
    dem=np.zeros((5,5));dem[2,2]=8.
    q=np.array([.5,2.5,10.]);a=np.array([4.5,2.5,2.]);b=np.array([4.5,2.5,18.])
    ints=merged(blocked_events(a,b,q,dem,-9999))
    # The ray touches the right edge x=3 of an 8m obstacle at endpoint z=6.8m.
    # z(u)=2+16u gives the manually derived transition u=0.3.
    assert len(ints)==1 and abs(ints[0][0])<1e-6 and abs(ints[0][1]-.3)<2e-6,ints
    checks.append('vertical_motion_manual_obstacle_transition_u_0.3')
    assert not merged(blocked_events(b,b,q,dem,-9999))
    checks.append('stationary_clear_link')
    touch=np.array([4.5,2.5,6.8]);assert merged(blocked_events(touch,touch,q,dem,-9999))
    checks.append('closed_pixel_terrain_touch_is_blocked')
    roots=sphere_events(np.array([-2.,0,0]),np.array([2.,0,0]),np.zeros(3),1.)
    assert np.allclose(sorted(roots),[.25,.75]);checks.append('sphere_threshold_roots')

    p=load(OUT/'子问题二/最终方案_待独立核验.json');f=load(OUT/'子问题二/最终方案.json')
    v=Verifier(p)
    mapping=pd.read_csv(OUT/'子问题二/架次编号映射.csv')
    tids=dict(zip(mapping.loc[mapping.type=='transport','candidate_id'],mapping.loc[mapping.type=='transport','submission_id']))
    rids=dict(zip(mapping.loc[mapping.type=='relay','candidate_id'],mapping.loc[mapping.type=='relay','submission_id']))
    assert np.allclose(p['objective'],f['objective'],atol=1e-8)
    assert set(b['box'] for b in f['boxes'])==set(v.boxes.box)
    assert len(f['boxes'])==80 and len(set(b['box'] for b in f['boxes']))==80
    wb=load_workbook(OUT/'问题三_结果提交.xlsx',data_only=True)
    template=load_workbook(ROOT/'D题/结果提交模板.xlsx')
    for sn,base in [('Q3_中继架次','Q3_中继架次'),('Q3_通信保障','Q3_通信保障'),('Q3_运输架次','Q2_运输架次'),('Q3_逐箱交付','Q2_逐箱交付')]:
        expected=[c.value for c in template[base][1] if c.value]
        assert [wb[sn].cell(1,i+1).value for i in range(len(expected))]==expected
    checks.append('submission_headers_match_original_template')
    tf=pd.read_csv(OUT/'子问题二/最终运输架次.csv').set_index('架次编号')
    rf=pd.read_csv(OUT/'子问题二/最终中继架次.csv').set_index('中继架次编号')
    bf=pd.read_csv(OUT/'子问题二/最终逐箱交付.csv').set_index('货箱编号')
    for r in p['routes']:
        row=tf.loc[tids[r['route_id']]]
        assert abs(row['开始时刻（s）']-r['start'])<1e-8 and abs(row['架次能耗（kWh）']-r['energy'])<1e-8
    for r in p['relays']:
        row=rf.loc[rids[r['relay_id']]]
        assert abs(row['悬停海拔（m）']-r['position']['z'])<1e-8
        assert abs(row['返回O01时刻（s）']-r['finish'])<1e-8
    for b in p['boxes']:
        assert bf.loc[b['box'],'架次编号']==tids[b['route_id']]
        assert abs(bf.loc[b['box'],'交付完成时刻（s）']-b['delivery'])<1e-8
    checks.append('final_json_csv_identifiers_and_floating_values')
    for sn,fp in [('Q3_运输架次','最终运输架次'),('Q3_中继架次','最终中继架次'),('Q3_逐箱交付','最终逐箱交付'),('Q3_通信保障','最终通信保障')]:
        frame=pd.read_csv(OUT/f'子问题二/{fp}.csv').fillna('')
        ws=wb[sn]
        for i,row in enumerate(frame.itertuples(index=False,name=None),2):
            for j,value in enumerate(row,1):
                actual=ws.cell(i,j).value
                if isinstance(value,(int,float)):
                    assert isinstance(actual,(int,float)) and abs(actual-value)<1e-8,(sn,i,j)
                else:assert (actual or '')==value,(sn,i,j)
    checks.append('excel_numeric_cells_equal_csv')
    comm=load(OUT/'子问题二/通信状态分段.json')
    # Dense points are a cross-check of the analytic event partition, not the
    # source of the continuous certificate (which is verified separately).
    sampled=0
    for r in p['routes']:
        rr=sorted([c for c in comm if c['route_id']==r['route_id']],key=lambda c:c['t0'])
        assert abs(rr[0]['t0']-r['takeoff'])<2e-4 and abs(rr[-1]['t1']-r['finish'])<2e-4
        for a,b in zip(rr,rr[1:]):assert abs(a['t1']-b['t0'])<2e-4
        stages,_,_,_=v.route_stages(r)
        for st in stages:
            for t in np.linspace(st['t0'],st['t1'],max(3,int((st['t1']-st['t0'])/3)+1))[1:-1]:
                c=next(c for c in rr if c['t0']-1e-7<=t<=c['t1']+1e-7)
                point=v.position(st,t);mu,_=v.link_margin_q(point,point,v.gateway,122)
                assert c['mode']==('直连' if mu>=-1e-7 else '中继'),(r['route_id'],t,mu,c)
                if c['mode']=='中继':assert c['relay_id'] in rids
                sampled+=1
    checks.append(f'event_partition_vs_independent_point_checks_{sampled}')
    av=load(OUT/'子问题一/非支配档案独立核验.json')
    assert all(a['passed'] for a in av)
    checks.append(f'nondominated_candidates_verified_{len(av)}')
    hashes={f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in OUT.glob('*.py')}
    save(OUT/'程序文件校验.json',hashes)
    save(OUT/'子问题二/交付交叉核验.json',dict(status='PASS',checks=checks,point_crosschecks=sampled,
          caveat='展示/交叉检查采样不替代独立全区间证书及事件端点核验。'))
    print('Q3_DELIVERY_PASS',len(checks),'checks',sampled,'point cross-checks',flush=True)
if __name__=='__main__':main()
