"""Create the Q3-only submission workbook, tables, figures and result report."""
from q3_common import *
from verify_q3 import Verifier
from q3_communication import export_communication
import copy, hashlib, platform
from collections import Counter,defaultdict
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from matplotlib.lines import Line2D

plt.rcParams.update({'font.family':['Microsoft YaHei','DejaVu Sans'],'axes.unicode_minus':False,
 'font.size':10,'axes.spines.top':False,'axes.spines.right':False,'svg.fonttype':'none',
 'pdf.fonttype':42,'figure.dpi':120,'savefig.dpi':240})
COL={'A':'#7B4AB1','B':'#0072B2','C':'#009E73','R01':'#D55E00','R02':'#CC79A7'}

def safe(obj):
    if isinstance(obj,(float,np.floating)) and not math.isfinite(obj):return None
    if isinstance(obj,dict):return {k:safe(v) for k,v in obj.items()}
    if isinstance(obj,(list,tuple)):return [safe(v) for v in obj]
    return q2.jr(obj)

def strict_save(path,obj):
    Path(path).write_text(json.dumps(safe(obj),ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')

def table(headers,rows):
    return '\n'.join(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']+
                     ['| '+' | '.join(str(x) for x in row)+' |' for row in rows])

def savefig(fig,name):
    fig.savefig(OUT/f'图表/{name}.png',bbox_inches='tight')
    fig.savefig(OUT/f'图表/{name}.svg',bbox_inches='tight')
    plt.close(fig)

def figures(ctx,p,comm,archive,tid,rid):
    # True-coordinate map, without invented circular communication coverage.
    fig,ax=plt.subplots(figsize=(10.5,7.7),layout='constrained')
    dx,_,left,_,dy,top=ctx.tr
    xs=(left+np.arange(ctx.dem.shape[1]+1)*dx-ctx.lon0)*ctx.mx/1000
    ys=(top+np.arange(ctx.dem.shape[0]+1)*dy-ctx.lat0)*ctx.my/1000
    im=ax.imshow(ctx.dem,extent=[xs[0],xs[-1],ys[-1],ys[0]],cmap='gist_earth',alpha=.5,vmin=100,vmax=1000)
    drawn=set()
    for r in p['routes']:
        seq=[0]+[int(z[1:]) for z in r['zones']]+[0]
        for a,b in zip(seq,seq[1:]):
            key=(min(a,b),max(a,b),r['vehicle'])
            if key in drawn:continue
            drawn.add(key);xy=ctx.xyz[[a,b],:2]/1000
            style={'A':(':',2.6,5),'B':('-',2.,3),'C':('--',1.5,4)}[r['vehicle']]
            ax.plot(xy[:,0],xy[:,1],c=COL[r['vehicle']],lw=style[1],ls=style[0],alpha=.8,zorder=style[2])
    ax.scatter(ctx.xyz[1:,0]/1000,ctx.xyz[1:,1]/1000,s=28,c='white',edgecolor='#222',zorder=4)
    for i in range(1,16):ax.annotate(ctx.nodes.zone.iloc[i],ctx.xyz[i,:2]/1000,xytext=(4,5),textcoords='offset points',fontsize=8)
    ax.scatter([0],[0],marker='*',s=210,c='#202020',zorder=5);ax.text(.1,-.35,'O01 / G01',weight='bold')
    for r in p['relays']:
        q=r['position'];x,y=q['x']/1000,q['y']/1000
        ax.plot([0,x],[0,y],ls='--',lw=1.15,c=COL[r['uav']],alpha=.7)
        ax.scatter(x,y,marker='D',s=68,c=COL[r['uav']],edgecolor='white',zorder=6)
        label=f"{rid[r['relay_id']]} · {r['uav']}\n离地{q['height']:.0f} m · {r['alpha']/60:.1f}–{r['beta']/60:.1f} min"
        offset=(-185,18) if rid[r['relay_id']]=='Q3-R004' else (7,-22)
        ax.annotate(label,(x,y),xytext=offset,textcoords='offset points',fontsize=7.8,
                    bbox=dict(boxstyle='round,pad=.25',facecolor='white',edgecolor='none',alpha=.9))
    ax.set(xlim=(-7.4,6.4),ylim=(-1.1,8.9),xlabel='相对 O01 东向距离 / km',ylabel='相对 O01 北向距离 / km',
           title='最终联合方案：运输路线与定点中继任务')
    ax.set_aspect('equal');ax.grid(alpha=.13);fig.colorbar(im,ax=ax,shrink=.67,label='DEM 地面高程 / m')
    ax.legend(handles=[Line2D([0],[0],color=COL[g],ls={'A':':','B':'-','C':'--'}[g],lw=2,label=f'{g} 型运输') for g in ['A','B','C']]+
                      [Line2D([0],[0],color=COL['R01'],ls='--',label='中继赴返路线')],loc='lower left',fontsize=8)
    savefig(fig,'图6-4_路线与中继部署')

    fig,axes=plt.subplots(2,1,figsize=(12,12),layout='constrained',gridspec_kw={'height_ratios':[1,1.7]},sharex=True)
    names=[f'U{i:02d}' for i in range(1,9)]+['R01','R02']
    for r in p['routes']:
        y=names.index(r['uav']);axes[0].barh(y,(r['finish']-r['start'])/60,left=r['start']/60,height=.62,color=COL[r['vehicle']],alpha=.9)
        axes[0].text((r['start']+r['finish'])/120,y,tid[r['route_id']].replace('Q3-',''),ha='center',va='center',fontsize=7,color='white')
    for r in p['relays']:
        y=names.index(r['uav']);c=COL[r['uav']]
        for start,end,alpha in [(r['start'],r['alpha'],.35),(r['alpha'],r['beta'],.9),(r['beta'],r['finish'],.35),(r['finish'],r['body_ready'],.12)]:
            axes[0].barh(y,(end-start)/60,left=start/60,height=.62,color=c,alpha=alpha)
        axes[0].text((r['alpha']+r['beta'])/120,y,rid[r['relay_id']].replace('Q3-',''),ha='center',va='center',fontsize=7)
    axes[0].set(yticks=range(len(names)),yticklabels=names,title='a  运输任务与中继阶段（中继浅色：准备/赴返；中继深色：服务）')
    axes[0].invert_yaxis()
    bn=[f'{g}-BAT-{i:02d}' for g,n in [('A',6),('B',4),('C',4)] for i in range(1,n+1)]+[f'R-EC-{i:02d}' for i in range(1,7)]
    for r in p['routes']+p['relays']:
        b=r.get('battery_id',r.get('component'));y=bn.index(b);c=COL.get(r.get('vehicle',r['uav']),'#888')
        axes[1].barh(y,(r['finish']-r['start'])/60,left=r['start']/60,height=.62,color=c,alpha=.9)
        axes[1].barh(y,r['charge_s']/60,left=r['finish']/60,height=.62,color=c,alpha=.25,hatch='///',edgecolor=c,lw=.3)
    axes[1].set(yticks=range(len(bn)),yticklabels=bn,xlabel='任务启动后的时刻 / min',title='b  电池及能源组件（实色：任务占用；斜线：充电）');axes[1].invert_yaxis()
    for ax in axes:
        ax.axvline(p['objective'][2]/60,c='#333',ls=':',lw=1);ax.grid(axis='x',alpha=.15);ax.set_axisbelow(True)
    savefig(fig,'图6-5_联合资源时间轴')

    fig,axes=plt.subplots(1,2,figsize=(12,4.5),layout='constrained')
    F=np.array([a['objective'] for a in archive]);lo=F.min(0);hi=F.max(0);span=np.maximum(hi-lo,1e-5)
    for a in archive:
        f=np.array(a['objective']);chosen=np.max(abs(f-np.array(p['objective'])))<1e-6
        axes[0].plot(range(4),(f-lo)/span,marker='o',c='#D55E00' if chosen else '#8FA7B7',lw=2.6 if chosen else 1,alpha=1 if chosen else .5)
    axes[0].set(xticks=range(4),xticklabels=['总架次','总能耗','联合时间','加权迟到'],ylabel='档案内显示尺度（0 最低，1 最高）',title='a  已认证非支配方案；橙色为代表方案')
    axes[0].grid(alpha=.15)
    for n,marker in [(24,'o'),(25,'s'),(26,'^')]:
        mask=F[:,0]==n
        axes[1].scatter(F[mask,1],F[mask,3]/60,s=65,marker=marker,c=F[mask,2]/60,cmap='viridis',vmin=F[:,2].min()/60,vmax=F[:,2].max()/60,label=f'{n} 架次',edgecolor='#333',lw=.4)
    axes[1].scatter([p['objective'][1]],[p['objective'][3]/60],marker='*',s=240,c='#D55E00',edgecolor='white',zorder=5)
    axes[1].set(xlabel='运输与中继总能耗 / kWh',ylabel='普通物资加权平均迟到 / min',title='b  能耗与配送及时性的取舍')
    axes[1].legend(fontsize=8);axes[1].grid(alpha=.15)
    sm=plt.cm.ScalarMappable(cmap='viridis',norm=plt.Normalize(F[:,2].min()/60,F[:,2].max()/60));fig.colorbar(sm,ax=axes[1],label='联合完成时间 / min',shrink=.8)
    savefig(fig,'图6-3_四目标权衡')

    fig,axes=plt.subplots(2,1,figsize=(12,9),layout='constrained',gridspec_kw={'height_ratios':[2.1,1]})
    ordered=sorted(p['routes'],key=lambda r:(r['start'],r['route_id']));ix={r['route_id']:i for i,r in enumerate(ordered)}
    rbody={r['relay_id']:r['uav'] for r in p['relays']}
    for c in comm:
        color='#7E9AAB' if c['mode']=='直连' else COL[rbody[c['relay_id']]]
        axes[0].barh(ix[c['route_id']],(c['t1']-c['t0'])/60,left=c['t0']/60,height=.72,color=color)
    axes[0].set(yticks=range(len(ordered)),yticklabels=[tid[r['route_id']] for r in ordered],title='a  直连优先的连续通信状态；准备和空闲留白');axes[0].invert_yaxis()
    axes[0].legend(handles=[Patch(color='#7E9AAB',label='G01 直连'),Patch(color=COL['R01'],label='R01 中继'),Patch(color=COL['R02'],label='R02 中继')],loc='upper right',ncol=3,fontsize=8)
    ss=pd.read_csv(OUT/'子问题二/直连裕量展示采样.csv');selected=max(ordered,key=lambda r:len(r['links']))
    data=ss[ss.route_id==selected['route_id']]
    axes[1].scatter(data.time_s/60,data.direct_margin_dB,s=9,c='#0072B2',label='直连裕量展示采样（约10 s）')
    cert=pd.read_csv(OUT/'子问题二/全区间通信证书.csv');cert=cert[cert.route_id==selected['route_id']]
    for _,r in cert.iterrows():axes[1].plot([r.t0/60,r.t1/60],[r.lower_margin_dB]*2,c='#D55E00',lw=1.6)
    axes[1].plot([],[],c='#D55E00',label='实际保障策略的区间裕量下界')
    axes[1].axhline(0,c='#222',ls='--',lw=1)
    axes[1].set(xlabel='任务启动后的时刻 / min',ylabel='双向链路裕量 / dB',title=f"b  {tid[selected['route_id']]}：直连失效阶段由有效中继保障")
    axes[1].legend(fontsize=8,ncol=2)
    for ax in axes:ax.grid(axis='x',alpha=.15);ax.set_axisbelow(True)
    savefig(fig,'图6-6_连续通信与裕量')

def main():
    p=load(OUT/'子问题二/最终方案_待独立核验.json')
    v=Verifier(p);v.verify_routes();v.verify_boxes();v.verify_relays();vr=v.finish()
    if not vr['passed']:raise RuntimeError(vr['failures'])
    strict_save(OUT/'子问题二/独立核验结果.json',vr)
    csv(OUT/'子问题二/全区间通信证书.csv',v.comm_certificates)
    csv(OUT/'子问题二/独立核验明细.csv',v.rows)
    comm=export_communication(p);ctx=Context()
    archive=load(OUT/'子问题一/非支配方案明细.json')
    tid={r['route_id']:f'Q3-T{i+1:03d}' for i,r in enumerate(sorted(p['routes'],key=lambda r:(r['start'],r['route_id'])))}
    rid={r['relay_id']:f'Q3-R{i+1:03d}' for i,r in enumerate(sorted(p['relays'],key=lambda r:(r['start'],r['relay_id'])))}
    trans=[{'架次编号':tid[r['route_id']],'无人机编号':r['uav'],'机型编号':r['vehicle'],'电池编号':r['battery_id'],
            '开始时刻（s）':r['start'],'访问服务区顺序':'→'.join(['O01']+r['zones']+['O01']),
            '返回O01时刻（s）':r['finish'],'架次能耗（kWh）':r['energy']} for r in sorted(p['routes'],key=lambda r:tid[r['route_id']])]
    boxes=[{'货箱编号':b['box'],'架次编号':tid[b['route_id']],'服务区编号':b['zone'],'交付完成时刻（s）':b['delivery']} for b in sorted(p['boxes'],key=lambda b:b['box'])]
    relays=[{'中继架次编号':rid[r['relay_id']],'中继无人机编号':r['uav'],'能源组件编号':r['component'],
             '开始时刻（s）':r['start'],'悬停经度（°）':r['position']['lon'],'悬停纬度（°）':r['position']['lat'],
             '悬停海拔（m）':r['position']['z'],'建链完成时刻（s）':r['alpha'],'服务结束时刻（s）':r['beta'],
             '返回O01时刻（s）':r['finish'],'架次能耗（kWh）':r['energy']} for r in sorted(p['relays'],key=lambda r:rid[r['relay_id']])]
    commrows=[{'运输架次编号':tid[c['route_id']],'通信阶段':c['phase'],'开始时刻（s）':c['t0'],'结束时刻（s）':c['t1'],
               '保障方式':c['mode'],'中继架次编号':rid[c['relay_id']] if c['relay_id'] else ''} for c in comm]
    for filename,data in [('最终运输架次',trans),('最终逐箱交付',boxes),('最终中继架次',relays),('最终通信保障',commrows)]:csv(OUT/f'子问题二/{filename}.csv',data)
    csv(OUT/'子问题二/架次编号映射.csv',[dict(type='transport',candidate_id=k,submission_id=val) for k,val in tid.items()]+[dict(type='relay',candidate_id=k,submission_id=val) for k,val in rid.items()])
    events=pd.read_csv(OUT/'子问题二/通信事件端点核验.csv')
    events['运输架次编号']=events.route_id.map(tid);events['保障主体']=events.provider.map(lambda x:rid.get(x,x))
    events.to_csv(OUT/'子问题二/通信事件端点核验.csv',index=False,encoding='utf-8-sig')
    resources=[]
    for r in p['routes']:
        resources += [dict(resource_type='运输机',resource=r['uav'],sortie=tid[r['route_id']],start=r['start'],task_end=r['finish'],ready=r['finish']),
                      dict(resource_type='运输电池',resource=r['battery_id'],sortie=tid[r['route_id']],start=r['start'],task_end=r['finish'],ready=r['recharge'])]
    for r in p['relays']:
        resources += [dict(resource_type='中继机',resource=r['uav'],sortie=rid[r['relay_id']],start=r['start'],task_end=r['finish'],ready=r['body_ready']),
                      dict(resource_type='能源组件',resource=r['component'],sortie=rid[r['relay_id']],start=r['start'],task_end=r['finish'],ready=r['recharge'])]
    csv(OUT/'子问题二/最终资源占用.csv',resources)
    named=copy.deepcopy(p)
    for r in named['routes']:r['candidate_id']=r['route_id'];r['route_id']=tid[r['route_id']]
    for r in named['relays']:r['solver_relay_id']=r['relay_id'];r['relay_id']=rid[r['relay_id']]
    for b in named['boxes']:b['route_id']=tid[b['route_id']]
    named['communication']=commrows;named['communication_endpoint_file']='通信事件端点核验.csv'
    named['source_snapshot']='缓存/问题二初始化快照.json';named['verification_status']='PASS'
    strict_save(OUT/'子问题二/最终方案.json',named)
    wb=load_workbook(ROOT/'D题/结果提交模板.xlsx')
    wb.copy_worksheet(wb['Q2_运输架次']).title='Q3_运输架次'
    wb.copy_worksheet(wb['Q2_逐箱交付']).title='Q3_逐箱交付'
    datasets={'Q3_中继架次':relays,'Q3_通信保障':commrows,'Q3_运输架次':trans,'Q3_逐箱交付':boxes}
    for name,data in datasets.items():
        ws=wb[name];headers=list(data[0]);nc=len(headers)
        for merged_range in list(ws.merged_cells.ranges):ws.unmerge_cells(str(merged_range))
        for row in ws.iter_rows(min_row=2):
            for c in row:c.value=None
        for i,header in enumerate(headers,1):ws.cell(1,i,header)
        for rn,row in enumerate(data,2):
            for cn,header in enumerate(headers,1):
                cell=ws.cell(rn,cn,row[header]);cell.font=Font(name='微软雅黑',size=10)
                cell.alignment=Alignment(vertical='center',horizontal='left' if isinstance(row[header],str) else 'right')
                cell.fill=PatternFill('solid',fgColor='F1F6FA' if rn%2==0 else 'FFFFFF')
                if isinstance(row[header],(float,int)):cell.number_format='0.00000000' if '经度' in header or '纬度' in header else '0.000000' if '能耗' in header else '0.000'
            ws.row_dimensions[rn].height=22
        for i,header in enumerate(headers,1):
            c=ws.cell(1,i);c.font=Font(name='微软雅黑',size=10,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24465A')
            c.alignment=Alignment(wrap_text=True,vertical='center')
            ws.column_dimensions[get_column_letter(i)].width=30 if '顺序' in header else 25 if '时间' in header or '时刻' in header else 22
        ws.row_dimensions[1].height=34;ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:{get_column_letter(nc)}{len(data)+1}'
        ws.print_area=ws.auto_filter.ref;ws.sheet_view.showGridLines=False
        ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.print_title_rows='1:1'
    desc=wb.create_sheet('Q3_说明')
    notes=[['说明','内容'],['方案范围','仅填问题三；Q1/Q2/Q4原模板页保持空白。Q3运输与逐箱补充页沿用原Q2字段。'],
           ['开始时刻','均指开始准备；实际起飞及充电事件见JSON和资源CSV。'],['通信边界','通信表记录区间内部状态；地形/距离临界端点以通信事件端点核验.csv的保障主体为准。'],
           ['数值精度','Excel保留原始双精度，只改变显示格式。不要用显示舍入值重新计算冲突。'],['最优性范围','已认证有限候选中的最小遗憾上界代表；未证明完整连续空间全局最优。']]
    for row in notes:desc.append(row)
    desc.column_dimensions['A'].width=22;desc.column_dimensions['B'].width=110
    for row in desc:
        for c in row:c.font=Font(name='微软雅黑',size=11);c.alignment=Alignment(wrap_text=True,vertical='center')
        desc.row_dimensions[row[0].row].height=40
    wb.active=wb.sheetnames.index('Q3_中继架次')
    excel=OUT/'问题三_结果提交.xlsx';wb.save(excel)
    check=load_workbook(excel,data_only=True)
    for name,data in datasets.items():
        ws=check[name]
        for rn,row in enumerate(data,2):
            for cn,k in enumerate(row,1):
                value=ws.cell(rn,cn).value;expected=row[k]
                if isinstance(expected,(float,int)):
                    assert isinstance(value,(float,int)) and abs(value-expected)<1e-8,(name,rn,cn)
                else:assert (value or '')==expected,(name,rn,cn)
    strict_save(OUT/'子问题二/Excel回读核验.json',dict(status='PASS',sheets={k:len(v) for k,v in datasets.items()},numeric_values_preserved=True))
    figures(ctx,p,comm,archive,tid,rid)
    comparison=[]
    frozen=load(OUT/'子问题一/对照_frozen_Q2.json')
    retimed=min([a for a in archive if a['source'] in ['final_polish','anchor_N','anchor_L','anchor_C','anchor_E']],key=lambda a:a.get('regret_upper',1e9))
    for label,a in [('A 冻结Q2快照后配置中继',frozen),('B 固定箱组/机型/资源顺序重排',retimed),('C 联合搜索最终代表',p)]:
        comparison.append(dict(method=label,NT=a['NT'],NR=a['NR'],ET=a['ET'],ER=a['ER'],C=a['objective'][2],L=a['objective'][3],
                               service_duration_s=sum(r['duration'] for r in a['relays'])))
    csv(OUT/'子问题一/方案对照.csv',comparison)
    par=[]
    for i,a in enumerate(archive):par.append(dict(archive_id=f'P{i+1:02d}',source=a['source'],NT=a['NT'],NR=a['NR'],N=a['objective'][0],E=a['objective'][1],C=a['objective'][2],L=a['objective'][3],
                      regret_lower=a['regret_lower'],regret_upper=a['regret_upper'],selected=np.max(abs(np.array(a['objective'])-p['objective']))<1e-6))
    csv(OUT/'子问题一/非支配方案与遗憾.csv',par)
    hard=[b for b in p['boxes'] if b['hard']];soft=[b for b in p['boxes'] if not b['hard']]
    resource_counts={g:len(set(r['uav'] for r in p['routes'] if r['vehicle']==g)) for g in ['A','B','C']}
    battery_counts={g:len(set(r['battery_id'] for r in p['routes'] if r['vehicle']==g)) for g in ['A','B','C']}
    summary=dict(objective=p['objective'],NT=p['NT'],NR=p['NR'],ET=p['ET'],ER=p['ER'],boxes=len(boxes),hard_boxes=len(hard),
       soft_boxes=len(soft),late_soft_boxes=sum(b['delivery']>b['expected']+2e-4 for b in soft),
       hard_deadline_min_margin_s=min(b['deadline']-b['delivery'] for b in hard),transport_uavs=resource_counts,
       transport_batteries=battery_counts,relay_uavs=len(set(r['uav'] for r in p['relays'])),relay_components=len(set(r['component'] for r in p['relays'])),
       min_transport_soc=min(r['soc'] for r in p['routes']),min_relay_soc=min(r['soc'] for r in p['relays']),
       archive_count=len(archive),regret_archive=p['regret_archive'],regret_upper=p['regret_upper'],
       completion_time_lower_bound=p['completion_time_lower_bound'],fixed_pool_bound_gap=(p['objective'][2]-p['completion_time_lower_bound'])/p['objective'][2],
       verification_checks=len(vr['checks']),verification='PASS',communication=load(OUT/'子问题二/通信事件核验摘要.json'),
       last_return=max([(r['finish'],tid[r['route_id']]) for r in p['routes']]+[(r['finish'],rid[r['relay_id']]) for r in p['relays']])[1],
       candidate_routes=len(load(OUT/'子问题一/累计候选路线.json')),source=p['source'])
    strict_save(OUT/'求解摘要.json',summary)
    manifest=[]
    paths=list((ROOT/'D题/数据/无人机应急物资运输基础数据').glob('*.xlsx'))+[next((ROOT/'D题/数据').rglob('镇龙乡及周边30米DEM.mat')),OUT/'缓存/问题二初始化快照.json']
    for f in paths:manifest.append(dict(path=str(f),bytes=f.stat().st_size,sha256=hashlib.sha256(f.read_bytes()).hexdigest()))
    strict_save(OUT/'输入文件校验.json',manifest)
    hardware=dict(gpu=load(OUT/'缓存/GPU位置筛选记录.json'),preprocessing=load(OUT/'缓存/预处理资源.json'),
       formal_search=load(OUT/'求解运行配置.json'),supplement=load(OUT/'子问题一/补充修复与冻结对照.json')['elapsed_s'],
       cpu_parallel_workers=4,geometry_threads=10,solver_threads_per_job=1,python=platform.python_version(),
       scipy=__import__('scipy').__version__,numpy=np.__version__,
       note='双GPU仅参与批量几何候选筛选；全区间认证和MILP在CPU执行。内存缓存复用DEM和240航段；未人为填满内存或显存。各阶段耗时不含此前失败调试，未将其当作等预算算法胜负证据。')
    strict_save(OUT/'计算资源记录.json',hardware)
    mr=lambda a:[f"{a['NT']}+{a['NR']}",f"{a['ET']+a['ER']:.6f}",f"{a['C']:.3f}",f"{a['L']:.3f}"]
    textmd=f'''# 问题三求解结果

最终提交的是一套联合调度方案：**{p['NT']}个运输架次、{p['NR']}个中继架次，80箱全部交付**。31箱医疗或首批保障物资满足硬时限。普通物资的期望时间为软目标，49箱中{summary['late_soft_boxes']}箱迟到，不能表述为全部80箱按期望时间送达。

## 一、子问题一：联合优化与目标权衡

{table(['指标','结果'],[['运输+中继总架次',p['objective'][0]],['运输能耗 / kWh',f"{p['ET']:.6f}"],['中继能耗 / kWh',f"{p['ER']:.6f}"],['总能耗 / kWh',f"{p['objective'][1]:.6f}"],['联合完成时间 / s',f"{p['objective'][2]:.6f}"],['普通物资加权平均迟到 / s',f"{p['objective'][3]:.6f}"],['档案内遗憾',f"{p['regret_archive']:.8f}"],['保守遗憾上界',f"{p['regret_upper']:.8f}"]])}

采用自适应ε约束、运输—中继关联移除/重构、局部MILP修复和偏好不确定性决策。偏好中心为四目标等权，L1半径0.3，计算12个极点。最终按有效外松弛得到的遗憾上界最小选择；这不是把四项单独最好的数值拼接成一行。

正式搜索为2轮、每轮4条种子链、每链8步；链之间共享初始/轮间档案，不能称为8次独立完整重复实验。补充计算对适配任务尝试A型机改派，并复用相同需求、地形和通信认证联合调整中继与时刻。资源敏感性得到的方案若在原库存下仍可行，也合并到档案后重新比较。共保存{len(archive)}个已独立认证的非支配方案；不将尚未完成的消融或10次独立重复写成已完成。

{table(['方案','运输+中继架次','总能耗 / kWh','联合完成 / s','加权迟到 / s'],[[a['method']]+mr(a) for a in comparison])}

A使用问题二初始化快照，并不代表问题二目录的最新结果。B固定箱组、机型及运输资源身份/顺序，只开放时刻与中继；C开放结构邻域及改派。此处是不同求解范围下的可行方案对照，各方法累计预算不同，不据此宣称算法统计优越。

![四目标权衡](图表/图6-3_四目标权衡.png)

固定累计路线库外松弛的完成时间下界为{p['completion_time_lower_bound']:.3f} s，相对可行上界差为{100*summary['fixed_pool_bound_gap']:.2f}%。该下界先删除全部通信采样约束，再放松资源先后、整数和中继约束，仅保留分类覆盖、机型总工作量以及适用的架次/能耗上限；有效但较弱，**不能证明全路线、连续空间的全局最优**。归一化尺度由初始方案冻结，后续方案优于初始参考时允许出现负归一化值。

## 二、子问题二：最终联合调度与可行性

### 运输架次

{table(list(trans[0]),[[f'{x:.3f}' if isinstance(x,float) else x for x in row.values()] for row in trans])}

### 中继架次

{table(list(relays[0]),[[f'{x:.8f}' if '经度' in k or '纬度' in k else f'{x:.3f}' if isinstance(x,float) else x for k,x in row.items()] for row in relays])}

时刻均从任务启动算起。“开始”是开始准备；Excel保留原始精度，以上只作显示舍入。最终使用运输机A/B/C分别为{resource_counts['A']}/{resource_counts['B']}/{resource_counts['C']}架，匹配电池分别{battery_counts['A']}/{battery_counts['B']}/{battery_counts['C']}组；中继使用{summary['relay_uavs']}架机体和{summary['relay_components']}个不同能源组件。未使用的库存仍存在，不把实际使用量当成独立完成任务的最小资源需求。

硬时限最小裕度为{summary['hard_deadline_min_margin_s']:.6f} s；存在恰好到达截止时刻的任务。最低运输返航SOC为{100*summary['min_transport_soc']:.4f}%，最低中继返航SOC为{100*summary['min_relay_soc']:.4f}%。实际最后返場任务是{summary['last_return']}；末次充电与中继周转不延长题目定义的联合完成时间。

![路线与部署](图表/图6-4_路线与中继部署.png)

![资源占用](图表/图6-5_联合资源时间轴.png)

### 连续通信与独立核验

从原始Excel、原DEM和实际箱号重新计算；不复用求解器缓存的真假覆盖标记。共{len(vr['checks'])}项核验通过，覆盖箱组、时限、三维飞行与能源、机体/电池/组件占用及通信。每个仿射运动阶段利用三维扫掠三角形与闭DEM像元裁剪得到整段保障证书，直连失败时核验同一中继的接入和回传。当前最小证书裕量下界为{vr['minimum_certificate_lower_margin_dB']:.6f} dB。

通信结果表含{len(commrows)}段。直连优先状态由地形遮挡事件和距离门限球交点划分，并对{summary['communication']['event_points']}个阶段/切换端点单独检查。边界时刻以`通信事件端点核验.csv`的实际保障主体为准；共享边界不表示同时接受两架中继服务。展示用约10 s裕量采样只用于图形，连续结论来自区间及端点认证。通信规则仅相对于题给30 m DEM和简化传播模型。

![通信保障](图表/图6-6_连续通信与裕量.png)

### 输出与复现

- `问题三_结果提交.xlsx`：仅填Q3原有两表及Q3运输、逐箱补充页；所有数值来自同一最终方案。
- `子问题二/最终方案.json`：正式编号方案，含通信关系，可供问题四冻结继承。
- `子问题二/最终运输架次.csv`、`最终中继架次.csv`、`最终逐箱交付.csv`、`最终通信保障.csv`、`最终资源占用.csv`：完整可机读结果。
- `子问题一/非支配方案与遗憾.csv`、`资源敏感性.json`、`偏好半径敏感性.json`：权衡与有限对照。
- `子问题二/独立核验结果.json`、`全区间通信证书.csv`、`通信事件端点核验.csv`、`Excel回读核验.json`：核验依据。

数值范围：24个中继位置/高度候选，100/200/300 m高度层；水平初始区间200 m、垂直100 m，未认证区间最多细分4层；每实体中继最多3个架次槽，搜索时域30000 s；累计{summary['candidate_routes']}个运输候选及其显式副本。候选路线允许1至3个不同服务区，本次代表最终选中的运输架次均为单服务区；这不是证明多点访问无益。

两张RTX 3080 Ti实际参与候选传播/遮挡筛选，每卡完成52,952,280次射线采样计算；其后由CPU作完整几何认证与MILP。CPU采用4个并行优化任务，每任务1个求解线程，几何预计算10线程。DEM、航段和证书驻留内存重复使用，具体测量与环境见`计算资源记录.json`和`运行环境.md`。没有为了占满资源制造无效计算。运行命令见`README.md`。
'''
    (OUT/'问题三求解结果.md').write_text(textmd,encoding='utf-8')
    # A readable table preview for visual QA; the workbook values are checked
    # independently above, and the preview uses the same exported dataset.
    fig,ax=plt.subplots(figsize=(14,2.8));ax.axis('off')
    compact=[[r['中继架次编号'],r['中继无人机编号'],r['能源组件编号'],f"{r['开始时刻（s）']:.2f}",f"{r['建链完成时刻（s）']:.2f}",f"{r['服务结束时刻（s）']:.2f}",f"{r['返回O01时刻（s）']:.2f}",f"{r['架次能耗（kWh）']:.4f}"] for r in relays]
    tb=ax.table(cellText=compact,colLabels=['中继架次','实体机','能源组件','开始/s','建链/s','服务结束/s','返场/s','能耗/kWh'],loc='center',cellLoc='center')
    tb.auto_set_font_size(False);tb.set_fontsize(10);tb.scale(1,1.9)
    for (row,col),cell in tb.get_celld().items():
        cell.set_edgecolor('white');cell.set_facecolor('#24465A' if row==0 else '#F1F6FA' if row%2 else 'white')
        if row==0:cell.get_text().set_color('white')
    fig.savefig(OUT/'图表/结果表核对预览.png',bbox_inches='tight');plt.close(fig)
    print('REPORT_COMPLETE',summary['objective'],'checks',len(vr['checks']), 'workbook',str(excel),flush=True)
if __name__=='__main__':main()
