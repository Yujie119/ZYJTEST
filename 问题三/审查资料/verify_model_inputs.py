"""问题三建模审查：只复算输入、链路门限和充电变换，不求解调度。"""

from __future__ import annotations

import json
import math
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "D题" / "数据" / "无人机应急物资运输基础数据"
OUT = Path(__file__).resolve().parent


def rows_of(path: Path, sheet: str = "数据") -> list[tuple]:
    wb = load_workbook(path, data_only=True, read_only=True)
    rows = list(wb[sheet].values)
    wb.close()
    return rows


def main() -> None:
    boxes = [r for r in rows_of(DATA / "物资需求与配送时限.xlsx", "逐箱货箱清单")[1:] if r[0]]
    medical = {r[0] for r in boxes if r[2] == "医疗物资"}
    first = {r[0] for r in boxes if r[5] == "是"}
    hard = medical | first

    communication = rows_of(DATA / "通信链路参数.xlsx")
    parameters = {(r[0], r[3]): r[4] for r in communication[2:] if r[0] is not None}
    frequency = float(parameters["传播参数", "f"])
    loss_system = float(parameters["传播参数", "Lsys"])
    loss_blockage = float(parameters["传播参数", "Lobs"])
    threshold = float(parameters["接收参数", "Psens"]) + float(parameters["接收参数", "M"])
    links = []
    for label, a, b in [
        ("运输—网关", "运输无人机", "固定网关 G01"),
        ("运输—中继接入", "运输无人机", "中继接入端"),
        ("中继回传—网关", "中继回传端", "固定网关 G01"),
    ]:
        pa, ga = float(parameters[a, "Pt"]), float(parameters[a, "G"])
        pb, gb = float(parameters[b, "Pt"]), float(parameters[b, "G"])
        forward = pa + ga + gb - loss_system - threshold
        reverse = pb + gb + ga - loss_system - threshold
        bidirectional = min(forward, reverse)
        clear_km = 10 ** ((bidirectional - 32.45 - 20 * math.log10(frequency)) / 20)
        links.append({
            "link": label,
            "forward_allowed_loss_db": forward,
            "reverse_allowed_loss_db": reverse,
            "bidirectional_allowed_loss_db": bidirectional,
            "unblocked_budget_distance_km": clear_km,
            "blocked_budget_distance_km": clear_km / 10 ** (loss_blockage / 20),
        })

    relay_rows = rows_of(DATA / "中继无人机数据.xlsx")
    relay = dict(zip(relay_rows[1], relay_rows[2]))
    full_charge_s = float(relay_rows[11][2])
    capacity = float(relay["能源组件可用能量（kWh）"])
    reserve = float(relay["返航电量下限（%）"]) / 100
    service_kw = float(relay["悬停功率（kW）"]) + float(relay["通信附加功率（kW）"])

    def appendix_charge(soc: float) -> float:
        if soc < 0.9:
            return full_charge_s * (0.65 * (0.9 - soc) / 0.9 + 0.35)
        return full_charge_s * 0.35 * (1 - soc) / 0.1

    def energy_form_charge(soc: float) -> float:
        e = 1 - soc
        return full_charge_s * (3.5 * e if e <= 0.1 else 0.35 + 0.65 / 0.9 * (e - 0.1))

    charge_rows = [{"soc": s, "appendix_s": appendix_charge(s), "eq_6_17_s": energy_form_charge(s)}
                   for s in (1, 0.95, 0.9, 0.5, 0.2)]
    max_difference = max(abs(appendix_charge(i / 10000) - energy_form_charge(i / 10000))
                         for i in range(10001))

    template = load_workbook(ROOT / "D题" / "结果提交模板.xlsx", data_only=True, read_only=True)
    schema = {ws.title: [str(c.value) for c in ws[1] if c.value is not None] for ws in template}
    template.close()
    result = {
        "scope": "原始参数与公式复算；没有执行问题三调度、DEM连续通信认证或全局最优性证明",
        "boxes": {"total": len(boxes), "medical": len(medical), "first_batch": len(first),
                  "overlap": len(medical & first), "hard_deadline": len(hard),
                  "soft_tardiness": len(boxes) - len(hard)},
        "communication": {"frequency_mhz": frequency, "system_loss_db": loss_system,
                          "blockage_loss_db": loss_blockage, "reception_threshold_dbm": threshold,
                          "links": links},
        "relay": {"parameters": relay, "vehicle_ids": [relay_rows[i][0] for i in (6, 7)],
                  "component_count": relay_rows[11][1], "full_charge_s": full_charge_s,
                  "sortie_energy_budget_kwh": (1 - reserve) * capacity,
                  "service_power_kw": service_kw,
                  "assumed_setup_energy_kwh": service_kw * float(relay["建链时间（s）"]) / 3600},
        "charge_comparison": {"examples": charge_rows, "grid_count": 10001,
                              "max_absolute_difference_s": max_difference,
                              "without_sos2_example": {"energy_kwh": 0.1 * capacity,
                                  "correct_charge_s": 0.35 * full_charge_s,
                                  "nonadjacent_convex_combination_s": 0.1 * full_charge_s}},
        "template": {"sheet_count": len(schema), "headers": schema},
        "checks": {"unique_box_ids": len({r[0] for r in boxes}) == len(boxes),
                   "paper_hard_deadline_count_matches": len(hard) == 31,
                   "paper_link_thresholds_match": [r["bidirectional_allowed_loss_db"] for r in links] == [122, 116, 126],
                   "charge_formula_matches_appendix": max_difference < 1e-9},
    }
    destination = OUT / "关键参数核验.json"
    destination.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"saved": str(destination), "checks": result["checks"],
                      "boxes": result["boxes"], "template_sheet_count": len(schema)}, ensure_ascii=False, indent=2))
    if not all(result["checks"].values()):
        raise SystemExit("存在参数或公式核验未通过项。")


if __name__ == "__main__":
    main()
