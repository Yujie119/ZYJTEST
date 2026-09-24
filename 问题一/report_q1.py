"""Export official Q1 results, research figures and a Chinese results report."""
from pathlib import Path
from copy import copy
import json
import math
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.comments import Comment
from PIL import Image, ImageDraw, ImageFont
from q1_core import WORK, ROOT

S1,S2,S3=[WORK/f"子问题{s}" for s in ("一","二","三")]
NAVY,TEAL,ORANGE,RED="#244563","#168C8C","#DC8C28","#C84A4A"
plt.rcParams.update({"font.family":"Microsoft YaHei","axes.unicode_minus":False,
                     "font.size":10,"axes.spines.top":False,"axes.spines.right":False,
                     "pdf.fonttype":42,"savefig.facecolor":"white"})


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def table(headers,rows):
    def show(v):
        return "—" if v is None else str(v).replace("|","/").replace("\n"," ")
    return "\n".join(["| "+" | ".join(headers)+" |","| "+" | ".join(["---"]*len(headers))+" |"]+
                    ["| "+" | ".join(show(x) for x in row)+" |" for row in rows])


def save_figure(fig,path):
    fig.savefig(path.with_suffix(".png"),dpi=200,bbox_inches="tight")
    fig.savefig(path.with_suffix(".pdf"),bbox_inches="tight")
    plt.close(fig)


def wrap_pixels(text,font,width,draw):
    lines=[]
    for original in str(text).split("\n"):
        line=""
        for char in original:
            if line and draw.textlength(line+char,font=font)>width:
                lines.append(line)
                line=char
            else:
                line+=char
        lines.append(line)
    return lines


def render_sheet(sheet,path,last_row=None):
    """Deterministic cell preview using saved dimensions and a Chinese font."""
    last_col=max(c.column for c in sheet[1] if c.value is not None)
    last_row=last_row or max(c.row for row in sheet for c in row if c.value is not None)
    widths=[]
    for c in range(1,last_col+1):
        letter=openpyxl.utils.get_column_letter(c)
        width=sheet.column_dimensions[letter].width if letter in sheet.column_dimensions else 13
        widths.append(int(width*8+7))
    heights=[int((sheet.row_dimensions[r].height or 25)*1.5) for r in range(1,last_row+1)]
    font=ImageFont.truetype("C:/Windows/Fonts/msyh.ttc",16)
    image=Image.new("RGB",(sum(widths)+2,sum(heights)+2),"white")
    draw=ImageDraw.Draw(image)
    yy=1
    for r,height in enumerate(heights,1):
        xx=1
        for c,width in enumerate(widths,1):
            cell=sheet.cell(r,c)
            text="" if cell.value is None else str(cell.value)
            if isinstance(cell.value,(float,int)):
                decimals=cell.number_format.count("0")-1 if "." in cell.number_format else 0
                text=f"{cell.value:.{max(0,decimals)}f}"
            draw.rectangle((xx,yy,xx+width,yy+height),outline="#ABB5BF",fill="white")
            lines=wrap_pixels(text,font,width-12,draw)
            for k,line in enumerate(lines):
                y=yy+(height-len(lines)*21)/2+k*21
                assert y+21<=yy+height+1,(sheet.title,r,c,text,"clipped")
                x=xx+6 if c==4 and r>1 else xx+(width-draw.textlength(line,font=font))/2
                draw.text((x,y),line,fill="#172D3C",font=font)
            xx+=width
        yy+=height
    image.save(path)


def workbook(plan):
    previews=WORK/"表格预览"
    previews.mkdir(exist_ok=True)
    original=ROOT/"D题/结果提交模板.xlsx"
    w=openpyxl.load_workbook(original)
    for sheet in w:
        sheet.row_dimensions[1].height=48
        render_sheet(sheet,previews/f"模板_{sheet.title}.png",last_row=1)
    # Reload: preview adjustments never modify the source template.
    w=openpyxl.load_workbook(original)
    sheet=w["Q1_单点组批"]
    row_style=[copy(c._style) for c in sheet[2]]
    for r,p in enumerate(plan,2):
        values=[p["sortie"],p["zone"],p["vehicle"],"; ".join(p["boxes"]),
                p["mass"],p["volume"],p["fly_s"],p["energy"],100*p["return_soc"]]
        for c,value in enumerate(values,1):
            cell=sheet.cell(r,c,value)
            cell._style=copy(row_style[c-1])
            align=copy(cell.alignment)
            align.vertical="center"
            align.horizontal="left" if c==4 else "center" if c<=3 else "right"
            cell.alignment=align
        for col,fmt in {5:"0.0",6:"0.000",7:"0.000000",8:"0.000000000",9:"0.000000"}.items():
            sheet.cell(r,col).number_format=fmt
        sheet.row_dimensions[r].height=max(32,20*math.ceil(len(p["boxes"])/2)+8)
    sheet.row_dimensions[1].height=42
    for column,width in {"G":19,"H":20,"I":18}.items():
        sheet.column_dimensions[column].width=width
    sheet.freeze_panes="E2"
    sheet.auto_filter.ref=f"A1:I{len(plan)+1}"
    sheet.print_area=f"A1:I{len(plan)+1}"
    sheet.page_setup.orientation="landscape"
    sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth=1
    sheet.page_setup.fitToHeight=0
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.print_title_rows="1:1"
    sheet["G1"].comment=Comment("往返时间=附录2去程与回程飞行时间之和，含爬升、巡航、下降。累计作业时间还包括准备、装载和交接，见问题一求解结果.md。","User")
    sheet["H1"].comment=Comment("仅填写最小最大遗憾最终方案。E_cal=可用电量；爬升附加能耗=mgh/(0.72×3.6e6)；返航安全余量20%。","User")
    sheet["I1"].comment=Comment("数值单位为百分数，例如23.088493表示23.088493%。公式为100×(1−架次能耗/该机型电池可用能量)。","User")
    output=WORK/"问题一_结果提交.xlsx"
    w.save(output)
    saved=openpyxl.load_workbook(output,data_only=True)
    render_sheet(saved["Q1_单点组批"],previews/"最终Q1表.png")
    for sh in list(saved)[1:]:
        assert not any(c.value is not None for row in sh.iter_rows(min_row=2) for c in row)
    return output


def figures(summary,plan):
    cap=pd.read_csv(S1/"最大安全载荷.csv")
    matrix=cap.pivot(index="zone",columns="vehicle",values="safe_load_kg")[["A","B","C"]]
    fig,ax=plt.subplots(figsize=(7.2,8))
    im=ax.imshow(matrix.to_numpy(),cmap="YlGnBu",vmin=0,vmax=80,aspect="auto")
    for i,row in enumerate(matrix.to_numpy()):
        for j,q in enumerate(row):
            ax.text(j,i,f"{q:.3f}",ha="center",va="center",color="white" if q>55 else NAVY)
    ax.set_xticks(range(3),["A型","B型","C型"])
    ax.set_yticks(range(len(matrix)),matrix.index)
    ax.set_title("最大安全载荷（kg）｜返航安全余量20%")
    fig.colorbar(im,ax=ax,label="kg",fraction=.045,pad=.04)
    fig.tight_layout()
    save_figure(fig,S1/"最大安全载荷")
    frame=pd.DataFrame(plan)
    counts=frame.groupby(["zone","vehicle"]).size().unstack(fill_value=0).reindex(columns=["B","C"],fill_value=0)
    fig,axs=plt.subplots(1,2,figsize=(11,6))
    yy=np.arange(len(counts))
    axs[0].barh(yy,counts.B,color=TEAL,label="B型")
    axs[0].barh(yy,counts.C,left=counts.B,color=NAVY,label="C型")
    axs[0].set_yticks(yy,counts.index)
    axs[0].invert_yaxis()
    axs[0].set_xticks([0,1,2])
    axs[0].set_xlabel("架次数")
    axs[0].set_title("最终方案：各服务区机型与架次")
    axs[0].legend(frameon=False)
    energies=frame.groupby("zone").energy.sum().reindex(counts.index)
    axs[1].barh(yy,energies,color=ORANGE)
    axs[1].invert_yaxis()
    axs[1].set_yticks(yy,counts.index)
    axs[1].set_xlabel("运输能耗（kWh）")
    axs[1].set_title("各服务区合计能耗")
    fig.tight_layout()
    save_figure(fig,S2/"最终组批方案")
    front=pd.read_csv(S2/"非支配方案.csv")
    fig,axs=plt.subplots(1,2,figsize=(10.8,4))
    for i,r in front.iterrows():
        color=TEAL if r["selected"] else ORANGE
        axs[0].scatter(r.E_kWh,r.T_s/60,s=120,c=color,zorder=3)
        axs[0].annotate(f"{int(r.N)}架次",(r.E_kWh,r.T_s/60),
                        xytext=(0,12 if i==0 else -23),textcoords="offset points",ha="center")
    axs[0].set_xlim(front.E_kWh.min()-.025,front.E_kWh.max()+.025)
    axs[0].set_ylim(front.T_s.min()/60-10,front.T_s.max()/60+10)
    axs[0].set_xlabel("总运输能耗（kWh）")
    axs[0].set_ylabel("累计作业时间（min）")
    axs[0].set_title("三目标非支配集：2组目标值")
    for i,r in front.iterrows():
        axs[1].bar(np.arange(3)+(i-.5)*.34,[r.normalized_N,r.normalized_E,r.normalized_T],
                   width=.32,color=TEAL if r["selected"] else ORANGE,label=f"{int(r.N)}架次")
    axs[1].set_xticks(range(3),["架次数","能耗","累计时间"])
    axs[1].set_ylabel("基准Pareto范围归一化值")
    axs[1].set_title("数值越小越优")
    axs[1].legend(frameon=False)
    fig.tight_layout()
    save_figure(fig,S2/"非支配方案与权衡")
    fig,axs=plt.subplots(1,2,figsize=(10.8,4))
    w=np.linspace(0,1,501)
    axs[0].plot(w,w,color=TEAL,label="18架次：评价 = w_E")
    axs[0].plot(w,1-w,color=ORANGE,label="19架次：评价 = 1−w_E")
    axs[0].axvspan(1/6,.5,color=TEAL,alpha=.1,label="基准偏好区间")
    axs[0].axvline(.5,color="gray",linestyle=":")
    axs[0].set_xlabel("能耗权重 w_E")
    axs[0].set_ylabel("归一化加权评价")
    axs[0].legend(frameon=False,fontsize=8)
    axs[0].set_title("选择切换阈值为 w_E=0.5")
    delta=np.linspace(0,4/3,401)
    axs[1].plot(delta,np.maximum(0,2*np.minimum(1,1/3+delta/2)-1),color=TEAL,label="18架次")
    axs[1].plot(delta,np.maximum(0,1-2*np.maximum(0,1/3-delta/2)),color=ORANGE,label="19架次")
    axs[1].axvline(1/3,color="gray",linestyle=":",label="基准 δ=1/3")
    axs[1].set_xlabel("等权中心下的L1偏好半径 δ")
    axs[1].set_ylabel("最坏遗憾")
    axs[1].legend(frameon=False)
    fig.tight_layout()
    save_figure(fig,S2/"偏好与最小最大遗憾")
    dense=np.load(S3/"双GPU载荷响应曲面.npz")
    rho,keys,q=dense["rho"],dense["keys"],dense["safe_load_kg"][:,:,2]
    fig,axs=plt.subplots(1,3,figsize=(12,4),sharex=True)
    highlight={"S008":RED,"S003":NAVY,"S012":TEAL}
    for k,(zone,vehicle) in enumerate(keys):
        ax=axs[["A","B","C"].index(vehicle)]
        ax.plot(rho*100,q[k],color=highlight.get(zone,"#BBBBBB"),
                lw=1.7 if zone in highlight else .8,alpha=1 if zone in highlight else .6,
                label=zone if zone in highlight else None,zorder=3 if zone in highlight else 1)
    for ax,v in zip(axs,["A","B","C"]):
        ax.axvline(20,color="black",linestyle=":",lw=1)
        ax.set_title(f"{v}型")
        ax.set_xlabel("返航安全余量（%）")
        ax.set_xlim(0,60)
        ax.set_ylabel("最大安全载荷（kg）")
    axs[2].legend(frameon=False,fontsize=8)
    fig.suptitle("双GPU载荷响应｜η=0.72；灰线为其余服务区；不可达处断线")
    fig.tight_layout()
    save_figure(fig,S3/"安全余量与最大载荷")
    response=pd.read_csv(S3/"安全余量方案响应.csv")
    valid=response[response.feasible]
    fig,axs=plt.subplots(1,3,figsize=(12,4))
    for ax,col,label,div in zip(axs,["N","E_kWh","T_s"],["架次数","总能耗（kWh）","累计作业时间（h）"],[1,1,3600]):
        ax.plot(valid.rho*100,valid[col]/div,"o-",color=TEAL,ms=4)
        ax.axvline(summary["baseline_plan_feasible_rho"]*100,color=ORANGE,ls=":",label="原方案失效阈值")
        ax.axvline(summary["global_feasible_rho"]*100,color=RED,ls="--",label="任务可行性阈值")
        ax.set_xlabel("返航安全余量（%）")
        ax.set_ylabel(label)
        ax.set_xlim(0,40)
        ax.axvspan(summary["global_feasible_rho"]*100,40,color=RED,alpha=.07)
    axs[0].legend(fontsize=8,frameon=False)
    fig.suptitle("固定基准归一化尺度与偏好规则后重新优化")
    fig.tight_layout()
    save_figure(fig,S3/"安全余量与最终方案")


def results_report(s,plan):
    front=pd.read_csv(S2/"非支配方案.csv")
    cap=pd.read_csv(S1/"最大安全载荷.csv").pivot(index="zone",columns="vehicle",values="safe_load_kg")
    counts=pd.read_csv(S1/"候选规模统计.csv").groupby("zone").sum(numeric_only=True)
    lps=pd.read_csv(S2/"全局连续松弛下界.csv")
    response=pd.read_csv(S3/"安全余量方案响应.csv")
    preferences=pd.read_csv(S2/"偏好敏感性.csv")
    eta=pd.read_csv(S3/"爬升效率敏感性.csv")
    critical=read_json(S3/"安全余量临界值.json")
    alt=read_json(S2/"非支配方案明细.json")[1]["plan"]
    b8=[p for p in alt if p["zone"]=="S008"]
    de=front.E_kWh.iloc[0]-front.E_kWh.iloc[1]
    dt=front.T_s.iloc[1]-front.T_s.iloc[0]
    text=f"""# 问题一求解结果与核验说明

## 1. 最终提交方案

在正文假设1–4、返航安全余量20%、爬升效率0.72、等权偏好中心和L1偏好半径δ=1/3下，最终选择**18架次方案**，完成15个服务区全部80箱、758 kg、2.011 m³物资交付。B型和C型各执行9架次，A型未被选用；这些是架次数，不是实体无人机数量。

总运输能耗为**{s['final_objectives'][1]:.9f} kWh**，累计作业时间为**{s['final_objectives'][2]:.6f} s（{s['final_objectives'][2]/3600:.4f}累计作业小时）**。累计纯飞行时间为{s['cumulative_flight_s']:.6f} s，最低返航SOC为**{s['min_return_soc']*100:.6f}%**。累计作业时间是各架次准备、装载、飞行和交接时间之和，不是全部任务实际完成时刻。

[正式提交表](问题一_结果提交.xlsx)的`Q1_单点组批`只填写这18条最终架次。候选库、另一个非支配方案和敏感性情景另存分析文件，不写入第一问结果表。

### 1.1 原题与模板口径

第一问明确不考虑实体无人机和共享电池调度；医疗与首批保障时限作为调度约束在问题二引入。因此不分配U01等实体编号、不编造电池编号或出发时刻，不将本问方案宣称为已经满足问题二时限与资源约束。

原模板第一问包含服务区、机型、货箱编号列表、质量、体积、往返时间、能耗及SOC。用户所列八列表头属于原模板`Q2_运输架次`；已另存[八列架次视图](子问题二/最终方案_八列架次视图.csv)，不适用的实体无人机、电池、开始时刻和返回绝对时刻四列留空。

第一问表中的“往返时间”采用附录2去程与回程飞行时间之和，含爬升、巡航、下降。准备、装载和交接另计入累计作业时间，两者在[最终方案明细](子问题二/最终方案明细.csv)均列出。交接参数来自机型参数表，因此程序按机型g索引，而非服务区i。

## 2. 子问题一：最大安全载荷与候选组批

直接读取原始节点、机型、逐箱清单和DEM MAT文件，以调度中心为原点建立局部WGS84米制坐标。遍历航段相交的全部DEM像元，以最高高程加50 m确定巡航海拔；去程带载、返程空载。源文件SHA-256见[输入校验](输入文件校验.json)。

### 2.1 三机型、15服务区的最大安全载荷

{table(["服务区","A型/kg","B型/kg","C型/kg"],[[z,*[f"{cap.loc[z,v]:.6f}" for v in ["A","B","C"]]] for z in cap.index])}

45组均可达，39组达到额定载荷；B–S008及C–S002、S003、S004、S008、S012共6组受能量约束。连续最大载荷不保证存在同质量的不可拆箱组合，实际组批仍需满足体积和货箱组成要求。

![最大安全载荷](子问题一/最大安全载荷.png)

### 2.2 完整候选库与等价压缩

逐箱可行候选架次共**19525个**。同一服务区内，质量和体积相同的货箱对本问能耗、时间目标可交换。因此枚举“服务区—机型—各规格箱数”得到**732个模式**，删除相同箱数组合下能耗、时间被支配的机型后保留**463个模式**。

模式p对规格t的用箱数记为a_tp，可用数为n_t。实现采用整数次数y_p：

\\[
\\sum_p a_{{tp}}y_p=n_t,\\quad
y_p\\in\\mathbb Z_{{\\ge0}},\\quad
y_p\\le\\min_{{t:a_{{tp}}>0}}\\left\\lfloor n_t/a_{{tp}}\\right\\rfloor.
\\]

整数解按规格从编号池无重复取箱展开；任何逐箱可行方案也能汇总为模式次数，因此与正文逐箱0–1精确覆盖等价。第一问不涉及的时限和实体资源不能被这种可交换性自动处理。

{table(["服务区","逐箱候选数","压缩模式数","筛选后模式数"],[[z,int(r.labelled_candidates),int(r.candidate_patterns),int(r.retained_patterns)] for z,r in counts.iterrows()])}

## 3. 子问题二：ε搜索、下界与最终择优

### 3.1 ε约束与自适应搜索

先解三个单目标整数规划，再以能耗为主目标，限制架次数和累计时间。每取得新点，就把时间上界收紧到所得时间减0.001 s；不采用固定规则网格。

最低能耗方案为19架次。任何架次数大于19的非支配方案，其累计时间必须短于该方案，否则将被支配。在该时间上界内最大化架次数仍得19，所以架次上界只需搜索18与19。

{table(["方案","架次数","能耗/kWh","累计作业时间/s","最坏遗憾","提交"],[[r.solution,int(r.N),f"{r.E_kWh:.9f}",f"{r.T_s:.6f}",f"{r.max_regret:.9f}","是" if r.selected else "否"] for r in front.itertuples()])}

5个区域中3次返回整数最优点（含1个重复目标点），2次证实整数不可行，活动区域全部穷尽。本例各区域LP均可行，**未发生LP不可行剪枝**；LP可行不代表整数可行。[搜索日志](子问题二/自适应epsilon搜索日志.csv)

独立动态规划进一步确认全局非支配目标集同样只有2个点。这里统计的是目标向量；同规格货箱编号互换可能形成多个等价的最优编号方案。

![非支配方案](子问题二/非支配方案与权衡.png)

### 3.2 连续松弛下界

{table(["目标","LP下界","整数最优值","相对差距"],[[r.objective,f"{r.LP_lower_bound:.6f}",f"{r.IP_optimum:.6f}",f"{r.relative_gap*100:.4f}%"] for r in lps.itertuples()])}

此处差距是整数模型与LP松弛的结构性差距，不是求解器未收敛的MIP gap。整数子问题均检查最优状态、整数性和原始约束；日志另存求解器对偶界、MIP gap及耗时。

### 3.3 偏好不确定性与最小最大遗憾

采用已核实非支配集的理想点及最差分量进行归一化：

\\[
\\mathbf f^{{min}}=(18,\\ {s['normalization_ideal'][1]:.9f},\\ {s['normalization_ideal'][2]:.6f}),\\\\
\\mathbf f^{{max}}=(19,\\ {s['normalization_nadir'][1]:.9f},\\ {s['normalization_nadir'][2]:.6f}).
\\]

18与19架次方案分别归一化为(0,1,0)、(1,0,1)。基准w⁰=(1/3,1/3,1/3)、δ=1/3是本次声明的决策设置，不是题目参数。偏好集合有6个极点，1/6≤w_E≤1/2；两方案评价分别为w_E和1−w_E。所以18架次方案在基准权重集合内均不劣，最大遗憾为0，而19架次为2/3。

程序还在完整整数可行域求解极小极大模型，避免仅在已发现档案中择优：

\\[
\\min_{{y,r}}r,\\qquad
(\\mathbf w^v)^\\mathsf T\\bar{{\\mathbf f}}(y)-\\Phi^*(\\mathbf w^v)\\le r,\\quad\\forall v.
\\]

极点最优值Φ*来自完整加权整数规划；最小最大遗憾模型的下界和最终遗憾均为0。同遗憾时，用三个归一化目标的正权和破同优。全单纯形情景下两方案最大遗憾均为1，按此规则选18架次。

{table(["偏好中心","δ","架次数","能耗/kWh","累计时间/s","最大遗憾"],[[r.center,f"{r.delta:.6f}",int(r.N),f"{r.E_kWh:.6f}",f"{r.T_s:.3f}",f"{r.max_regret:.6f}"] for r in preferences.itertuples()])}

偏重能耗w⁰=(0.2,0.6,0.2)、δ=1/3时会选择19架次。这表明“最终最优”以已声明偏好集合为条件，并非三项目标同时更优。

![偏好与遗憾](子问题二/偏好与最小最大遗憾.png)

### 3.4 权衡来源与最终方案

两组目标值的区别全部来自S008：18架次方案使用1个C型架次运45 kg；19架次方案改用2个B型架次，分别运{'、'.join(str(int(p['mass'])) for p in b8)} kg。增加1架次节省{de:.9f} kWh（{de/front.E_kWh.iloc[0]*100:.4f}%），同时增加{dt:.6f} s（{dt/60:.4f} min）累计作业时间。

{table(["架次","服务区","机型","箱数","质量/kg","体积/m³","飞行往返/s","累计作业/s","能耗/kWh","返航SOC"],[[p["sortie"],p["zone"],p["vehicle"],len(p["boxes"]),f'{p["mass"]:.0f}',f'{p["volume"]:.3f}',f'{p["fly_s"]:.3f}',f'{p["operation_s"]:.3f}',f'{p["energy"]:.6f}',f'{p["return_soc"]*100:.4f}%'] for p in plan])}

完整货箱编号见正式Excel与CSV。全部货箱恰好交付一次，各架次仅服务一个服务区，质量、体积和20%返航SOC约束逐架次复核通过。

![最终组批](子问题二/最终组批方案.png)

## 4. 子问题三：安全余量与效率敏感性

### 4.1 连续载荷与可行性阈值

双GPU计算45组服务区—机型、1201个安全余量、5个爬升效率，共270225组双精度载荷点。安全载荷随安全余量单调不增，不可达单独记为NaN，不与可行但零载荷混同。

![载荷响应](子问题三/安全余量与最大载荷.png)

原18架次方案的最大统一安全余量为**{s['baseline_plan_feasible_rho']*100:.9f}%**，瓶颈是S004的59 kg C型架次；超过此值需要重新组批。

允许重新组批时，任务最大可行安全余量为**{s['global_feasible_rho']*100:.9f}%**。第一问不限制架次及实体资源，每个货箱单独可飞是可行的充要条件：必要性来自能耗随载荷不减，充分性由每箱单独成架次构造。因此

\\[
\\rho^{{feas}}=\\min_b\\max_{{g:\\,m_b\\le Q_g,\\,v_b\\le V_g}}
\\left(1-\\frac{{F_{{g,i(b)}}(m_b)}}{{E_g^{{use}}}}\\right).
\\]

瓶颈是S008的14 kg饮用水货箱，最佳机型为{critical['group_thresholds'][critical['limiting_groups'][0]]['vehicle']}型。临界点仍有27架次可行方案，临界点上方0.000001由整数模型证实不可行。

### 4.2 重新优化的组批响应

各情景保持基准归一化尺度、等权中心和δ不变，重新生成可行模式并求解遗憾模型。只改变电量预算，固定能耗校准电量；归一化值允许超出[0,1]，不截断、不逐情景重标度。

{table(["安全余量","可行","架次数","B型/C型","能耗/kWh","累计时间/s"],[[f'{r.rho*100:.6f}%',"是" if r.feasible else "否",int(r.N) if r.feasible else "—",f'{int(r.n_B)}/{int(r.n_C)}' if r.feasible else "—",f'{r.E_kWh:.6f}' if r.feasible else "—",f'{r.T_s:.3f}' if r.feasible else "—"] for r in response.itertuples()])}

![方案响应](子问题三/安全余量与最终方案.png)

安全余量为25%、30%、34%、35%时，选中方案分别为19、21、24、26架次。重新优化后的每个目标分量不必单调：例如34%至35%能耗略降、架次和累计时间上升，是最小最大遗憾重新权衡的结果，不与可行域收缩矛盾。

### 4.3 爬升效率假设检验

{table(["η","架次数","能耗/kWh","累计时间/s","最低SOC"],[[f'{r.eta:.2f}',int(r.N),f'{r.E_kWh:.6f}',f'{r.T_s:.3f}',f'{r.min_return_soc*100:.4f}%'] for r in eta.itertuples()])}

η=0.60至0.84仍选择18架次，累计时间不变。该检验只反映等效参数敏感性，不证明mgh/η等同于实测旋翼爬升全过程能耗。

## 5. 独立核验、硬件与复现

核验程序不调用HiGHS：逐箱枚举确认19525个可行候选，以另一种线段—像元相交算法复核15条航段，再用337个动态规划状态得到相同的完整非支配目标集。13个安全余量情景的逐箱方案也全部通过物理约束核验。[核验结果](独立核验结果.json)

8个CPU工作进程并行完成29个情景，每个HiGHS实例1线程。两张RTX 3080 Ti分别实际计算138115与132110组载荷；180次CPU/GPU抽样比较最大误差为{s['resources']['max_gpu_cpu_difference_kg']:.3e} kg。候选库、稀疏矩阵和批次缓存于内存；80箱任务所需存储远低于64GB，不为占满硬件创建无关数据。详见[资源记录](计算资源记录.json)。

计算核心约{s['elapsed_s']:.2f} s，不含报告、图表和独立核验。GPU设备计算约0.1 s/卡，不含进程及CUDA初始化，不能据此宣称端到端GPU加速比。HiGHS在CPU运行。使用现有PyTorch 2.8.0+cu128和MKL顺序线程层，未安装依赖，也未启用允许重复OpenMP运行库的不安全选项。

在本目录完整重算：

```powershell
& 'E:\\tool\\anaconda3\\python.exe' '.\\solve_q1.py'
```

该命令重算三个子问题，生成正式表、图表、报告并执行核验。仅重导出已有结果运行`report_q1.py`；独立核验运行`verify_q1.py`。

代码：`q1_core.py`负责物理模型与优化；`solve_q1.py`负责主流程；`gpu_q1.py`负责双GPU；`parallel_q1.py`负责CPU多进程；`verify_q1.py`负责独立核验；`report_q1.py`负责导出。三个子目录分别保留对应原始计算结果、PNG与PDF图表。正式表保留原始精度，显示数值按列格式四舍五入。
"""
    (WORK/"问题一求解结果.md").write_text(text,encoding="utf-8")


def make_outputs():
    summary=read_json(WORK/"求解摘要.json")
    plan=read_json(S2/"最终方案.json")
    output=workbook(plan)
    figures(summary,plan)
    from verify_q1 import verify
    verify()
    results_report(summary,plan)
    readme=r"""# 问题一代码与结果

先阅读 `问题一求解结果.md`。正式文件 `问题一_结果提交.xlsx` 的 `Q1_单点组批` 只包含最终18架次，其他问题未填写。

全部重算（双GPU、8个CPU工作进程）：

```powershell
& 'E:\tool\anaconda3\python.exe' '.\solve_q1.py'
```

仅计算加参数 `--compute-only`。仅从已有数值导出Excel、图表、报告并核验：运行 `report_q1.py`。单独核验：运行 `verify_q1.py`。

前提：原始文件仍在同级D题目录；使用指定Anaconda Python及已有numpy、pandas、scipy、openpyxl、matplotlib、Pillow、psutil、CUDA版PyTorch；两张CUDA设备可用。程序不安装包、不改原始附件。GPU只负责批量物理计算，LP/MILP由CPU上的HiGHS完成。

设置：E_cal=E_use；η=0.72；基准安全余量20%；目标按完整基准Pareto范围归一化；等权中心、L1半径1/3。第一问不含实体机、电池时序、通信和时限约束。

模板往返时间=去回程飞行时间，累计作业时间还含准备、装载、交接。八列表头属于原模板Q2，另存八列视图中的实体编号与绝对时刻留空。详细参数和适用范围见报告。
"""
    (WORK/"README.md").write_text(readme,encoding="utf-8")
    print(f"EXPORTED {output}",flush=True)


if __name__=="__main__":
    make_outputs()
